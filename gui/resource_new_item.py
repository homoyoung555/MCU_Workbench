"""Append complete analog resource units through the existing draft transaction."""
import re

from openpyxl.utils import get_column_letter, column_index_from_string


def _iomux_choices(context):
    sheet = context['editor'].by_name.get('IOMUX')
    if not sheet:
        return []
    columns = sheet['columns']
    pin_col = next((i for i,c in enumerate(columns) if c['label'].strip().lower()=='pin'),None)
    adc_col = next((i for i,c in enumerate(columns) if c['label'].strip().lower()=='analog / adc'),None)
    if pin_col is None or adc_col is None:
        return []
    edits = {c['cell']:c['value'] for c in context['payload']['changes'] if c['sheet']=='IOMUX'}
    choices = []
    for row,cells in enumerate(sheet['rows'],1):
        if row<=sheet['headerRows'] or not cells[adc_col]['editable']:
            continue
        address = f'{get_column_letter(adc_col+1)}{row}'
        pin = edits.get(f'{get_column_letter(pin_col+1)}{row}',cells[pin_col]['value'])
        if pin is None or str(pin).strip() in ('','/'):
            continue
        choices.append({'row':row,'pin':str(pin),'cell':address,
                        'analog':str(edits.get(address,cells[adc_col]['value']) or '')})
    return choices


def _with_iomux(context, identity, selected):
    if not isinstance(selected,list) or len(selected)!=22:
        raise ValueError('请选择完整的 22 个通道引脚')
    choices = {p['row']:p for p in _iomux_choices(context)}
    if any(type(r) is not int or r not in choices for r in selected):
        raise ValueError('请选择当前 IOMUX 中可编辑的引脚')
    # Existing channel assignments are authoritative, including shared ADC12 tokens.
    from program_gen.core.adc_test import _iomux_pin_map, _normalize_pin
    current = _iomux_pin_map(context['workbook'].dataframe('IOMUX'),0).get(identity.upper(),{})
    changes = {(c['sheet'],c['cell']):dict(c) for c in context['payload']['changes']}
    for channel,row in enumerate(selected):
        pin = choices[row]
        if channel in current:
            if current[channel]!=_normalize_pin(pin['pin']):
                raise ValueError(f'{identity} IN{channel} 已映射到 {current[channel]}，请先在 IOMUX 修改原映射')
            continue
        key = ('IOMUX',pin['cell'])
        cell = context['editor'].by_name['IOMUX']['rows'][row-1][column_index_from_string(re.match(r'[A-Z]+',pin['cell'])[0])-1]
        change = changes.setdefault(key,{'sheet':'IOMUX','cell':pin['cell'],'old':cell['value'],'value':pin['analog']})
        previous = change['value'].strip()
        change['value'] = (previous+'/' if previous not in ('','/') else '')+f'{identity.upper()}_IN{channel}'
    return {**context['payload'],'changes':list(changes.values())}


def resource_new_item(session, payload):
    from gui.flow_new_item import _context, _text
    context = _context(session, payload)
    name = payload.get('sheet')
    if name not in ('ADC', 'OPA', 'DAC', 'PVD'):
        raise ValueError('请选择 ADC、OPA、DAC 或 PVD 专表')
    base = session.editor.by_name[name]
    frame = context['workbook'].dataframe(name)
    columns = [c['label'] for c in base['columns']]
    edit_row = payload.get('editRow')
    editing = edit_row is not None
    if editing and (type(edit_row) is not int or edit_row<2):
        raise ValueError('请选择专表数据行')
    additions = [op for op in context['payload'].get('appendRows',[]) if op['sheet']==name]
    def original_cells(row):
        if row<=base['maxRow']:
            return base['rows'][row-1]
        op = additions[row-base['maxRow']-1]
        return [dict(c,value=c['value'] if op['copy'] else None,editable=True) for c in base['rows'][op['sourceRow']-1]]
    groups = []
    for i, row in frame.iloc[:base['maxRow']+len(additions)-1 if editing else base['maxRow']-1].iterrows():
        record = {k: _text(row.get(k)) for k in columns}
        if name == 'ADC':
            if record['ADC'] and (not groups or groups[-1]['label']!=record['ADC']):
                groups.append({'label':record['ADC'], 'rows':[], 'records':[]})
            if not groups: continue
            record['ADC'] = groups[-1]['label']
            if not record['step']: record['step'] = '/'
        elif name == 'PVD':
            if not groups or groups[-1]['label'] != record['case']:
                groups.append({'label':record['case'], 'rows':[], 'records':[]})
        else:
            groups.append({'label':f"第 {int(i)+2} 行 · {record['mode']}", 'rows':[], 'records':[]})
        groups[-1]['rows'].append(int(i)+2)
        groups[-1]['records'].append(record)
    groups = [g for g in groups if all(r in base['resourceRows'] or editing and r>base['maxRow'] for r in g['rows']) and
              (len(g['rows']) == (22 if name=='ADC' else 2 if name=='PVD' else 1))]
    if editing:
        groups = [g for g in groups if edit_row in g['rows']]
        if not groups:
            raise ValueError('当前行不属于可编辑的完整配置组')
        for g in groups:
            g['editable'] = [[c['editable'] for c in original_cells(r)] for r in g['rows']]
    if payload['action'] == 'resource-options':
        return {'sheet':name, 'columns':columns, 'templates':groups,
                'iomuxPins':_iomux_choices(context) if name=='ADC' else [],
                'notice':('编辑当前配置组；保留原有行、合并单元格和关联标识。' if editing else {'ADC':'新增完整 CH0–CH21；新 ADC 名称须在 IOMUX 配齐对应引脚。',
                          'PVD':'同一 case 一次新增 up/down 两行，共用 step；阈值及 expect 分别填写。',
                          'OPA':'新增一个 OPA 用例；所有引用 OPA 专表的测试项都会使用它。',
                          'DAC':'新增一个 DAC 用例；所有引用 DAC 专表的测试项都会使用它。'}[name])}
    template = next((g for g in groups if g['rows'] == payload.get('sourceRows')), None)
    records = payload.get('records')
    if template is None or not isinstance(records,list) or len(records)!=len(template['rows']):
        raise ValueError('请选择完整的专表参考配置')
    if any(not isinstance(r,dict) or set(r)!=set(columns) or
           any(not isinstance(v,str) or len(v)>32767 for v in r.values()) for r in records):
        raise ValueError('专表字段必须完整且为文本')
    if editing:
        locked = {'ADC':('ADC','Chanel'),'PVD':('case','up/down')}.get(name,())
        if any(r[k]!=old[k] for r,old in zip(records,template['records']) for k in locked):
            raise ValueError('关联标识不可在参数编辑中修改')
        if 'iomuxRows' in payload:
            raise ValueError('已有配置的引脚映射请在 IOMUX 表修改')
    if name in ('OPA','DAC'):
        required = ('mode','down_limit','up_limit','target') if name=='OPA' else ('mode','code','down_limit','up_limit')
        if any(records[0][k].strip() in ('','/') for k in required):
            raise ValueError('新增配置必须填写：'+', '.join(required))
    if name in ('ADC','PVD'):
        key = 'ADC' if name=='ADC' else 'case'
        identity = records[0][key].strip()
        if not identity or any(r[key].strip()!=identity for r in records):
            raise ValueError('配置组名称必须一致且不能为空')
        if not editing and any(_text(v).casefold()==identity.casefold() for v in frame[key]):
            raise ValueError('专表已存在该配置组名称')
        if name=='ADC':
            if not re.fullmatch(r'ADC\d+',identity,re.I): raise ValueError('ADC名称应为 ADC 加编号')
            if 'iomuxRows' in payload:
                context = _context(session,_with_iomux(context,identity,payload['iomuxRows']))
            from program_gen.core.adc_test import _adc_sheet_rows, _iomux_pin_map
            import pandas as pd
            channels = _adc_sheet_rows(pd.DataFrame(records),0)[identity.upper()]
            pins = _iomux_pin_map(context['workbook'].dataframe('IOMUX'),0).get(identity.upper(),{})
            if set(channels)!=set(range(22)) or set(pins)!=set(range(22)):
                raise ValueError('ADC 配置及 IOMUX 都必须包含对应 CH0–CH21 / IN0–IN21')
        else:
            from program_gen.core.pvd_test import _compile_cases
            import pandas as pd
            if not _compile_cases(pd.DataFrame(records),frozenset()) and not editing:
                raise ValueError('新增 PVD 配置不能全部停用')
    draft = {k:context['payload'].get(k,[]) for k in ('changes','appendRows','flowOrder','trimStorageMappings','verifyStorageMappings','patternEdits','patternLinks')}
    draft = {k:list(v) for k,v in draft.items()}
    first = template['rows'][0] if editing else base['maxRow']+len(additions)+1
    updates = {(c['sheet'],c['cell']):c for c in draft['changes']}
    for offset,(source,record) in enumerate(zip(template['rows'],records)):
        if editing:
            for i,k in enumerate(columns,1):
                if record[k]==template['records'][offset][k]:
                    continue
                cell = original_cells(source)[i-1]
                if not cell['editable']:
                    raise ValueError(f'{name} 第{source}行 {k} 为只读或合并单元格')
                address = f'{get_column_letter(i)}{source}'
                if record[k]==('' if cell['value'] is None else str(cell['value'])):
                    updates.pop((name,address),None)
                else:
                    updates[name,address] = {'sheet':name,'cell':address,'old':cell['value'],'value':record[k]}
            continue
        draft['appendRows'].append({'sheet':name,'sourceRow':source,'copy':False})
        draft['changes'].extend({'sheet':name,'cell':f'{get_column_letter(i)}{first+offset}',
                                 'old':None,'value':record[k]} for i,k in enumerate(columns,1))
    if editing:
        draft['changes'] = list(updates.values())
    raw = session.editor.apply(session.editor.sha256,draft['changes'],draft['appendRows'],draft['flowOrder'],
                               draft['trimStorageMappings'],draft['verifyStorageMappings'],draft['patternEdits'], draft.get('patternLinks'))
    _, compiled = session.compile(raw)
    session.editor.ensure_source_current()
    if not compiled.is_valid:
        return {'ready':False,'issues':[{'message':d.message} for d in compiled.diagnostics.errors]}
    return {'ready':True,'draft':draft,'row':first,'count':len(records)}
