"""SWD Excel 往返与两平台 Pattern 生成的公开边界。"""
from copy import copy
from datetime import date, datetime, time
from io import BytesIO
import posixpath
import re
import xml.etree.ElementTree as ET
from zipfile import ZipFile

import openpyxl

from gui import swd_legacy

GENERATOR_REVISION = '3'

_REQUIRED = ("pattern", "addr", "data", "type", "continue")
_ALIASES = {"PCD009": "PT104"}
_DEFAULT_PINS = {"PT104": {"swc": "PA14", "swd": "PA13", "reset": "NRST", "trig": "TRIGPIN"}}


def pin_defaults(chip="PT104"):
    """Return a copy so callers can safely override SWD/SWC/reset/trigger pins."""
    canonical = _ALIASES.get(str(chip).upper(), str(chip).upper())
    try: return dict(_DEFAULT_PINS[canonical])
    except KeyError: raise ValueError(f"没有 {chip} 的SWD默认引脚") from None


def _json_value(value):
    if isinstance(value, (datetime, date, time)): return value.isoformat()
    return value


def _sheet(book, sheet):
    if sheet is None:
        sheet = "Sheet1" if "Sheet1" in book.sheetnames else book.sheetnames[0]
    if sheet not in book.sheetnames: raise ValueError(f"工作表不存在：{sheet}")
    return book[sheet]


def _headers(ws):
    labels = [str(cell.value).strip() if cell.value is not None else "" for cell in ws[1]]
    headers, used = [], set()
    for column, label in enumerate(labels, 1):
        # Real pattern.xlsx has two blank trailing columns.  A JSON object needs
        # stable unique keys, while the physical-column order remains unchanged.
        name = label or f"__blank_{column}"
        if name in used: name = f"__column_{column}"
        headers.append(name)
        used.add(name)
    if not all(name in headers for name in _REQUIRED):
        missing = ", ".join(name for name in _REQUIRED if name not in headers)
        raise ValueError(f"{ws.title}缺少SWD表头：{missing}")
    return headers


def load_excel(raw, sheet=None):
    """Parse the selected SWD sheet without discarding other sheets or columns."""
    if not isinstance(raw, (bytes, bytearray)): raise TypeError("raw必须是Excel字节")
    book = openpyxl.load_workbook(BytesIO(raw), data_only=False, read_only=True)
    try:
        ws = _sheet(book, sheet)
        headers = _headers(ws)
        rows = []
        for row_number, cells in enumerate(ws.iter_rows(min_row=2, max_col=len(headers), values_only=True), 2):
            values = {header: _json_value(value) for header, value in zip(headers, cells)}
            if any(value is not None for value in values.values()):
                values["row"] = row_number
                rows.append(values)
        patterns = list(dict.fromkeys(str(row["pattern"]) for row in rows if row.get("pattern") not in (None, "")))
        return {"headers": headers, "rows": rows, "patterns": patterns, "sheets": list(book.sheetnames)}, ws.title
    finally:
        book.close()


def _worksheet_paths(package):
    ns = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
    rel_ns = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
    relations = {item.attrib['Id']: item.attrib['Target']
                 for item in ET.fromstring(package.read('xl/_rels/workbook.xml.rels'))
                 if item.attrib['Type'].endswith('/worksheet')}
    sheets = ET.fromstring(package.read('xl/workbook.xml')).find('{%s}sheets' % ns)
    paths = {}
    for item in sheets:
        target = relations.get(item.attrib['{%s}id' % rel_ns])
        if target:
            paths[item.attrib['name']] = (target.lstrip('/') if target.startswith('/')
                                          else posixpath.normpath(posixpath.join('xl', target)))
    return paths


def _editable_package(raw, sheet):
    """Skip auxiliary sheet parsing when linked parts can stay unchanged."""
    with ZipFile(BytesIO(raw)) as package:
        paths = _worksheet_paths(package)
        selected = paths.get(sheet)
        if selected is None or 'xl/styles.xml' not in package.namelist():
            return raw, None
        relationships = posixpath.join(posixpath.dirname(selected), '_rels', posixpath.basename(selected) + '.rels')
        # Drawings, tables and comments may renumber relationships on save. Keep
        # their existing full-workbook path rather than splice mismatched parts.
        if relationships in package.namelist() and any(not item.attrib['Type'].endswith('/printerSettings')
                for item in ET.fromstring(package.read(relationships))):
            return raw, None
        if 'xl/calcChain.xml' in package.namelist() and not re.search(rb'<calcPr\b[^>]*/>', package.read('xl/workbook.xml')):
            return raw, None
        omitted = set(paths.values()) - {selected}
        shadow = BytesIO()
        with ZipFile(shadow, 'w') as target:
            for entry in package.infolist():
                data = (b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData/></worksheet>'
                        if entry.filename in omitted else package.read(entry.filename))
                target.writestr(entry, data)
        return shadow.getvalue(), selected


def _restore_package(raw, edited, sheet, selected):
    """Replace the edited sheet; preserve every unrelated package member verbatim."""
    result = BytesIO()
    with ZipFile(BytesIO(raw)) as original, ZipFile(BytesIO(edited)) as saved, ZipFile(result, 'w') as output:
        replacements = {selected: saved.read(_worksheet_paths(saved)[sheet])}
        # Existing style IDs are reused. Copy a new style table only if openpyxl
        # appended styles (for example, a changed merged-cell border).
        tag = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}cellXfs'
        if len(ET.fromstring(saved.read('xl/styles.xml')).find(tag)) > len(ET.fromstring(original.read('xl/styles.xml')).find(tag)):
            replacements['xl/styles.xml'] = saved.read('xl/styles.xml')
        # Calculation chains are a cache of formula positions, not user data.
        # Remove stale chains and request Excel recalculate after row edits.
        if 'xl/calcChain.xml' in original.namelist():
            for name, attribute, suffix in [('xl/_rels/workbook.xml.rels', 'Type', '/calcChain'),
                                             ('[Content_Types].xml', 'PartName', '/calcChain.xml')]:
                tree = ET.fromstring(original.read(name))
                for item in list(tree):
                    if item.attrib.get(attribute, '').endswith(suffix): tree.remove(item)
                replacements[name] = ET.tostring(tree, encoding='utf-8', xml_declaration=True)
            def force_recalculation(match):
                element = re.sub(rb'\s(?:fullCalcOnLoad|forceFullCalc)="[^"]*"', b'', match.group())
                return element[:-2] + b' fullCalcOnLoad="1" forceFullCalc="1"/>'
            replacements['xl/workbook.xml'] = re.sub(rb'<calcPr\b[^>]*/>', force_recalculation, original.read('xl/workbook.xml'), count=1)
        for entry in original.infolist():
            if entry.filename == 'xl/calcChain.xml': continue
            output.writestr(entry, replacements.get(entry.filename, original.read(entry.filename)))
    return result.getvalue()


def _same_rows(raw, sheet, rows):
    model, selected = load_excel(raw, sheet)
    return selected == sheet and model["rows"] == list(rows)


def _clone_row(ws, source, destination, headers, values):
    for column in range(1, ws.max_column + 1):
        original, target = ws.cell(source, column), ws.cell(destination, column)
        target._style = copy(original._style)
        target.value = None
    ws.row_dimensions[destination].height = ws.row_dimensions[source].height
    for column, header in enumerate(headers, 1):
        if header in values: ws.cell(destination, column).value = values[header]


def export_excel(raw, sheet, rows):
    """Apply ordered physical rows while preserving unselected sheets and extra columns."""
    rows = list(rows)
    if _same_rows(raw, sheet, rows): return bytes(raw)
    editable, selected_part = _editable_package(raw, sheet)
    book = openpyxl.load_workbook(BytesIO(editable), data_only=False)
    ws = _sheet(book, sheet)
    headers = _headers(ws)
    added_headers = []
    for item in rows:
        if not isinstance(item, dict): raise ValueError("rows必须是对象列表")
        for name, value in item.items():
            if name != "row" and name not in headers and name not in added_headers and value not in (None, ""):
                added_headers.append(name)
    for name in added_headers:
        column = ws.max_column + 1
        ws.cell(1, column).value = name
        headers.append(name)
    originals = {number: {header: ws.cell(number, col).value for col, header in enumerate(headers, 1)} for number in range(2, ws.max_row + 1)}
    seen = set()
    for item in rows:
        if not isinstance(item, dict): raise ValueError("rows必须是对象列表")
        number = item.get("row")
        if number is not None:
            if not isinstance(number, int) or number not in originals or number in seen: raise ValueError("row必须是唯一的原物理行号或None")
            seen.add(number)
            for header, original in originals[number].items():
                if isinstance(original, str) and original.startswith("=") and item.get(header, original) != original:
                    raise ValueError(f"第{number}行公式不可在此编辑")
    original_max = ws.max_row
    active_rows = [number for number, values in originals.items() if any(value is not None for value in values.values())]
    append_only = ([item.get('row') for item in rows[:len(active_rows)]] == active_rows
                   and all(item.get('row') is None for item in rows[len(active_rows):]))
    positions = (active_rows + list(range(original_max + 1, original_max + 1 + len(rows) - len(active_rows)))
                 if append_only else list(range(2, len(rows) + 2)))
    merged_ranges = [copy(area) for area in ws.merged_cells.ranges]
    snapshots = {number: [copy(ws.cell(number, column)) for column in range(1, ws.max_column + 1)] for number in seen}
    heights = {number: ws.row_dimensions[number].height for number in seen}
    # MergedCell placeholders cannot receive values. Restore only merges whose
    # rows remain contiguous and whose covered cells have no distinct edits.
    for area in merged_ranges:
        ws.unmerge_cells(str(area))
    template = min(seen) if seen else (2 if original_max >= 2 else 1)
    for destination, item in zip(positions, rows):
        source = item.get("row")
        if source is not None:
            cells = snapshots[source]
            for column, original in enumerate(cells, 1):
                target = ws.cell(destination, column)
                target._style, target.value = copy(original._style), original.value
            ws.row_dimensions[destination].height = heights[source]
            for column, header in enumerate(headers, 1):
                if header in item: ws.cell(destination, column).value = item[header]
        else:
            _clone_row(ws, template, destination, headers, item)
    final_row = max([original_max] + positions) if append_only else max(positions, default=1)
    if original_max >= final_row + 1: ws.delete_rows(final_row + 1, original_max - final_row)
    destinations = ({number: number for number in range(1, original_max + 1)} if append_only else
                    {item['row']: destination for destination, item in zip(positions, rows) if item.get('row') is not None})
    destinations[1] = 1
    for area in merged_ranges:
        mapped = [destinations.get(number) for number in range(area.min_row, area.max_row + 1)]
        if any(number is None for number in mapped) or mapped != list(range(mapped[0], mapped[0] + len(mapped))):
            continue
        anchor = ws.cell(mapped[0], area.min_col).value
        covered = [ws.cell(number, column) for number in mapped for column in range(area.min_col, area.max_col + 1)
                   if (number, column) != (mapped[0], area.min_col)]
        if all(cell.value in (None, '', anchor) for cell in covered):
            ws.merge_cells(start_row=mapped[0], end_row=mapped[-1], start_column=area.min_col, end_column=area.max_col)
    output = BytesIO()
    book.save(output)
    edited = output.getvalue()
    return _restore_package(raw, edited, sheet, selected_part) if selected_part else edited


def _text(value):
    return "" if value is None else str(value).strip()


def _hex(value):
    text = _text(value).upper()
    return text[2:] if text.startswith("0X") else text


def _effective_value(row):
    data, steps, reset = _text(row.get("data")), _text(row.get("step")), _text(row.get("reset value"))
    if steps in ("", "/"): return _hex(data)
    if not reset: raise ValueError("位域写/读缺少reset value")
    result = _hex(reset)
    values, shifts = data.split(","), steps.split(",")
    if len(values) == len(shifts):
        for value, shift in zip(values, shifts): result = swd_legacy.bit_operation(value, shift, result)
    else:
        for shift in shifts: result = swd_legacy.bit_operation("1", shift, result)
    return result


def _operations(pattern_name, rows):
    current = None
    selected = []
    for row in rows:
        if row.get("pattern") not in (None, ""): current = _text(row["pattern"])
        if current == pattern_name: selected.append(row)
    if not selected: raise ValueError(f"未找到Pattern：{pattern_name}")
    return selected


def generate_pattern(pattern_name, rows, platform, pins=None, chip="PT104"):
    """Generate legacy-compatible Chroma/S100 text, encoded as gb18030 bytes."""
    platform = _text(platform).lower()
    canonical = _ALIASES.get(_text(chip).upper(), _text(chip).upper())
    if canonical != "PT104": raise ValueError(f"当前最小生成器仅支持PT104/PCD009，收到：{chip}")
    settings = pin_defaults(canonical)
    settings.update(pins or {})
    if set(settings) < {"swd", "swc", "reset"}: raise ValueError("pins必须包含swd、swc、reset")
    constants, core = swd_legacy.constants(platform)
    rows = _operations(pattern_name, rows)
    # Old pin.dec order is SWC, SWD, reset: PT104 therefore remains PA14,PA13,NRST.
    pin_list = ",".join(_text(settings[name]) for name in ("swc", "swd", "reset"))
    if platform == "chroma": output = f'SET_DEC_FILE "pin.dec"\n\nHEADER\n\n{pin_list};\n\nSPM_PATTERN({pattern_name}) {{\n\n'
    elif platform == "s100": output = f"HEAD[PIN_NAME]:\n{pin_list},{_text(settings.get('trig', 'TRIGPIN'))};\nEND_HEAD;\n@@PATTERN_DEFINE\n"
    else: raise ValueError(f"不支持的平台：{platform}")
    if _text(rows[0].get("continue")).upper() == "Y":
        output += core
        output += constants["tar"] + swd_legacy.data_bits("E000EDF0", platform, "write")
        output += constants["csw"] + swd_legacy.data_bits("22000042", platform, "write")
        output += constants["drw"] + swd_legacy.data_bits("A05F0003", platform, "write")
        output += " *101* RPT 53;\n\n" if platform == "chroma" else " 1011;RPT 53;\n\n"
        output += constants["csw"] + swd_legacy.data_bits("23000012", platform, "write")
    else: output += "START:*XXX* TS1;\n" if platform == "chroma" else "START:XXXX;\n"
    for row in rows:
        operation, address, source = _text(row.get("type")).upper(), _text(row.get("addr")), row.get("row", "?")
        suffix = " *101* RPT 53;\n\n" if platform == "chroma" else " 1011;RPT 53;\n\n"
        if not operation and not address:
            continue  # Pattern title/blank line: it only establishes the fill-down name.
        if "WAIT" in address.upper():
            if operation == "US": output += (f" *XXX* RPT {int(row['data']) - 156};//wait {row['data']}us\n\n" if platform == "chroma" else f" XXXX;RPT {int(row['data']) - 156};//wait {row['data']}us\n\n")
            elif operation == "MS": output += (f" *XXX* RPT {int(row['data']) * 1000};//wait {row['data']}ms\n\n" if platform == "chroma" else f" XXXX;RPT {int(row['data']) * 1000};//wait {row['data']}ms\n\n")
            else: raise ValueError(f"第{source}行未知WAIT类型：{operation}")
        elif address.upper().startswith("0X"):
            value = _effective_value(row)
            if operation == "W": output += constants["tar"] + swd_legacy.data_bits(_hex(address), platform, "write") + constants["drw"] + swd_legacy.data_bits(value, platform, "write") + suffix
            elif operation in {"R", "CAP", "C_H", "C_L"}:
                mode = {"R": "read", "CAP": "cap", "C_H": "high", "C_L": "low"}[operation]
                output += "//read" + address + "\n" + constants["tar"] + swd_legacy.data_bits(_hex(address), platform, "write") + constants["read"] + swd_legacy.data_bits(value, platform, mode) + suffix
            elif operation == "SRAM":
                sram_mode = "read" if platform == "s100" else "cap"
                for value in range(int(_hex(address), 16), int(_hex(row["data"]), 16) + 4, 4): output += "//read " + hex(value).upper() + "\n" + constants["tar"] + swd_legacy.data_bits(f"{value:08X}", platform, "write") + constants["read"] + swd_legacy.data_bits("FFFFFFFF", platform, sram_mode) + suffix
            else: raise ValueError(f"第{source}行未知SWD操作：{operation}")
        elif "_" in address and operation in {"W", "R"}:
            register, command = address.split("_", 1)
            output += swd_legacy.command(command, register.upper(), operation, platform) + swd_legacy.data_bits(_hex(row["data"]), platform, "write" if operation == "W" else "read") + suffix
        else: raise ValueError(f"第{source}行未知SWD操作：{operation}")
    output += " *101* RPT 53;\n *101*;\n END:*011* STOP;\n}\n" if platform == "chroma" else " 1011;RPT 53;\n 1011;\n END:0111;\n@@END_PATTERN_DEFINE\n"
    return output.encode("gb18030")
