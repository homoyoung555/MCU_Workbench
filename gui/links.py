"""Read-only GUI links from compiled plans to Pattern and storage sources."""


from openpyxl.utils import get_column_letter
from gui.specialized_links import specialized_links, _resource_slot
from gui.memory_links import module_patterns, module_details, test_entry, generated_names


def _step_references(steps):
    actions, unexpanded = [], []
    for step in steps:
        module, plan = step.module, step.module_plan
        if module in {'ENTER', 'RUN'} and plan:
            actions.append((module, plan.pattern.name))
        elif module == 'JUDGE' and plan:
            if plan.pe_window:
                actions.extend(('JUDGE_WINDOW', run.pattern.name)
                               for run in (plan.pe_window.before, plan.pe_window.after))
        elif module == 'DAC_WRITE' and plan:
            actions.append((module, plan.target.pattern.name))
        elif module == 'COMMIT' and plan:
            for target in (*plan.targets, *plan.field_values):
                actions.extend((module, name) for name in (target.program_pattern, target.read_pattern))
        elif module == 'SETVPOSP' and plan:
            if plan.updates or plan.should_run:
                actions.append((module, plan.run_pattern.pattern.name))
        elif module == 'DAC_ADJ' and plan:
            pass  # DAC_ADJ changes a value; its SETVPOSP step owns the Pattern reference.
        elif plan and (names := module_patterns(module,plan)) is not None:
            actions.extend((module,name) for name in names)
        elif module != 'EXPORT':
            unexpanded.append(module)
    return actions, unexpanded


def _flow_location(editor, row, label, column=None):
    sheet = editor.by_name.get('FLOW', {})
    if not row or not 1 <= row <= sheet.get('maxRow', 0):
        return []
    col = next((i + 1 for i, c in enumerate(sheet.get('columns', [])) if c.get('label') == column), None)
    coordinate = f'{get_column_letter(col)}{row}' if col else f'row {row}'
    return [{'sheet': 'FLOW', 'coordinate': coordinate, 'label': label}]


def _module_details(spec, steps, catalogs, editor):
    details = []
    for step in steps:
        plan = step.module_plan
        if plan is None or step.module not in {'DAC_ADJ', 'SETVPOSP'}:
            continue
        locations = _flow_location(editor, spec.excel_row, '查看本项Recipe', 'Recipe')
        if step.module == 'DAC_ADJ':
            condition, source = plan.condition, plan.target_source
            text = (f'RAW条件：{condition.source_symbol} {condition.operator} {condition.threshold:g} {condition.unit}；'
                    f'满足时将 {plan.target_symbol} 沿DAC码表移动 {plan.delta:+d} 个位置，两端限幅。')
            if source:
                if source.produced_at_excel_row:
                    text += f' 起始值来自本CP的TRIM结果（FLOW第{source.produced_at_excel_row}行）。'
                    locations += _flow_location(editor, source.produced_at_excel_row, '查看TRIM来源')
                elif source.restored_at_excel_row:
                    text += (f' 起始值来自寄存器恢复值：{source.register_symbol}'
                             f'[{source.bit_lsb + source.bit_width - 1}:{source.bit_lsb}]，页 {source.page}。')
                    locations += _flow_location(editor, source.restored_at_excel_row, '查看读取来源')
            bindings = getattr(catalogs, 'result_bindings', None)
            if bindings:
                try:
                    binding = bindings.for_symbol(spec.cp, condition.source_symbol)
                except KeyError:
                    pass
                else:
                    locations += _flow_location(editor, binding.producer_row + 2, '查看RAW结果来源')
        else:
            text = ('执行已准备的Pattern。' if plan.should_run else
                    '改写Pattern动态字段，本步骤不执行Pattern。' if plan.updates else
                    '配置已准备，本步骤无需改写或执行Pattern。')
            text += ' ' + '；'.join(dict.fromkeys(
                f'{u.pump.item}：{u.pump.symbol} → {u.label.field_label_base}'
                f' [{u.label.bit_hi}:{u.label.bit_lo}]（寄存器 {u.label.addr_bits}）'
                for u in plan.updates))
        details.append({'title': step.module, 'text': text.strip(), 'locations': locations})
    order=list(dict.fromkeys(s.module for s in steps))
    return sorted(details + module_details(spec,steps,editor),key=lambda d:order.index(d['title']))


def selected_storage_fields(catalog, kind, cp, item):
    """The same effective physical fields for the picker and association preview."""
    selected = catalog.selected_storage_fields(kind, item)
    if selected is not None:
        return selected
    matches = catalog.matches_trim if kind == 'TRIM' else catalog.matches_verify
    legacy = [f for f in catalog.fields if matches(cp, item, f)]
    return tuple(f for f in catalog.storage_fields if any(
        (f.page, f.register, f.lsb, f.width) == (old.page, old.register, old.lsb, old.width)
        for old in legacy))


def _configured_storage_details(spec, catalogs, editor, kind, enabled=None, consumed=()):
    catalog = getattr(catalogs, 'addr_map', None)
    if not callable(getattr(catalog, 'selected_storage_fields', None)):
        return []
    sheet = editor.by_name.get('ADDR_MAP', {})
    details = []
    for field in selected_storage_fields(catalog, kind, spec.cp, spec.item):
        destination = 'UID' if field.page.casefold() == 'uid' else 'OTP'
        used = (field.page, field.register, field.lsb, field.width) in consumed
        if enabled is not None and destination not in enabled and not used:
            continue
        row = field.register_layout.row_index + 2
        name = 'description' if field.source == 'description' else 'CP addr map'
        column = next((i + 1 for i, c in enumerate(sheet.get('columns', [])) if c.get('label') == name), None)
        coordinate = f'{get_column_letter(column)}{row}' if column else f'row {row}'
        locations = ([{'sheet': 'ADDR_MAP', 'coordinate': coordinate, 'label': f'查看 ADDR_MAP!{coordinate}'}]
                     if 1 <= row <= sheet.get('maxRow', 0) else [])
        details.append({'title': 'ADDR_MAP 存储映射',
                        'summary': f'已配置 · {field.page} / {field.name}',
                        'text': f'已配置 → 页 {field.page} · {field.register}[{field.lsb + field.width - 1}:{field.lsb}] · 字段 {field.name} · {field.width} bit。字段位置与存储映射页一致；后续使用步骤单独展示，不表示已经执行烧写。',
                        'storageMapping': {'page': field.page, 'register': field.register, 'field': field.name},
                        'locations': locations})
    return details


def _storage_details(excel_row, spec, catalogs, editor, kind='TRIM', symbol=None, enabled=None):
    symbol = symbol or getattr(spec, 'dac_final_symbol', None)
    writes = []
    for plan in getattr(catalogs, 'flash_window_plans', {}).values():
        if plan.cp != spec.cp or plan.excel_row <= excel_row:
            continue
        for step in plan.steps:
            if step.module != 'BUILDMAP' or step.module_plan is None:
                continue
            for register in step.module_plan.registers:
                for write in register.writes:
                    if symbol and write.symbol == symbol and any(
                            term.lsb == write.lsb and term.symbol == symbol for term in register.terms):
                        writes.append((register.page, register.register, write.field, write.lsb,
                                       write.width, plan.excel_row, 'BUILDMAP'))
    persistence = getattr(catalogs, 'persistence', None)
    for plan in getattr(persistence, 'rows', {}).values():
        if plan.cp != spec.cp or plan.excel_row <= excel_row:
            continue
        for operation in plan.operations:
            if operation.module == 'COMMIT':
                for write in operation.field_writes:
                    if symbol and write.value.kind == 'VARIABLE' and write.value.token == symbol:
                        writes.append((write.page, write.register, write.field, write.lsb,
                                       write.width, plan.excel_row, 'COMMIT'))
    sheet = editor.by_name.get('ADDR_MAP', {})
    column = next((i + 1 for i, c in enumerate(sheet.get('columns', []))
                   if c.get('label') == 'CP addr map'), None)
    registers = getattr(getattr(catalogs, 'addr_map', None), 'registers', ())
    details = []
    for page, register, field, lsb, width, consumer, module in dict.fromkeys(writes):
        locations = []
        for source in registers:
            if source.page != page or source.register != register:
                continue
            row = source.row_index + 2
            if 1 <= row <= sheet.get('maxRow', 0):
                physical = next((f for f in getattr(source, 'storage_fields', ()) if f.lsb == lsb and f.width == width), None)
                source_column = next((i + 1 for i, c in enumerate(sheet.get('columns', [])) if c.get('label') == 'description'), None) if physical and physical.source == 'description' else column
                coordinate = f'{get_column_letter(source_column)}{row}' if source_column else f'row {row}'
                locations.append({'sheet': 'ADDR_MAP', 'coordinate': coordinate,
                                  'label': f'查看 ADDR_MAP!{coordinate}'})
        locations += _flow_location(editor, consumer, f'查看 {module} 使用步骤', 'Recipe')
        text = (f'{spec.cp} · {symbol}（FLOW第{excel_row}行产生） → 页 {page} · '
                f'{register}[{lsb + width - 1}:{lsb}] · 字段 {field} · {width} bit。'
                f' FLOW第{consumer}行的 {module} 使用此结果；位段为寄存器有效值内的位置。'
                ' 此处展示数据组装关系，不代表已经执行烧写。')
        if not any(location['sheet'] == 'ADDR_MAP' for location in locations):
            text += ' 存储绑定已编译，但 ADDR_MAP 来源坐标未解析。'
        details.append({'title': 'ADDR_MAP 存储映射', 'summary': f'后续使用 · {module} · FLOW {consumer}',
                        'text': text, 'locations': locations})
    consumed = {(page, register, lsb, width) for page, register, field, lsb, width, consumer, module in writes}
    configured = _configured_storage_details(spec, catalogs, editor, kind, enabled, consumed)
    if enabled == set() and not details:
        return [{'title': 'ADDR_MAP 存储映射', 'summary': '当前行未启用 OTP/UID 存储',
                 'text': '当前行 OTP=no、UID=no，共用字段配置不会作为本行已启用的存储映射展示。', 'locations': []}]
    return configured + (details or [{'title': 'ADDR_MAP 存储映射', 'summary': '后续使用步骤尚未确认',
                        'text': f'{spec.cp} · {symbol or "结果来源未解析"}：未解析到后续 BUILDMAP/COMMIT 使用步骤；字段配置状态见已配置映射，不能据此判定没有映射。',
                        'locations': []}])


def pattern_links(compilation, editor, flow=None, workbook=None):
    if not compilation.is_valid:
        return []
    catalogs = compilation.plan.catalogs

    def enabled_storage(row):
        if flow is None:
            return None
        record = flow.iloc[row - 2]
        return {name for name in ('OTP', 'UID') if str(record.get(name, '')).strip().casefold() == 'yes'}
    pattern_plan = catalogs.fbist_pattern_plan
    output_patterns = list(pattern_plan.patterns) if pattern_plan else []
    aggregate = getattr(pattern_plan, 'pump_aggregate', None)
    if aggregate:
        output_patterns.append(aggregate)
    sources = {p.pattern_name: list(dict.fromkeys(r for w in p.writes for r in w.source_rows))
               for p in output_patterns}
    generated = set(sources) | generated_names(catalogs)
    previews = {p.pattern_name: [
        {'address': w.addr_bits, 'data': w.data_bits,
         'fields': [{'name': f.field_name, 'high': f.bit_hi, 'low': f.bit_lo,
                     'value': f.value.rstrip('*'), 'dynamic': f.dynamic} for f in w.fields]}
        for w in p.writes] for p in output_patterns}
    pattern_sheet = editor.by_name.get('pattern', {})
    description_col = next((i + 1 for i, c in enumerate(pattern_sheet.get('columns', []))
                            if str(c.get('label', '')).strip() == 'description'), None)
    entries, sequences = [], {}
    for row, spec in sorted(catalogs.trim_specs.items()):
        actions = [(op.kind, op.reference) for op in (*spec.operations, *spec.search.per_point_actions)]
        actions += [('PROTOCOL_ENTER', spec.protocol.enter_pattern), ('PROTOCOL_EXIT', spec.protocol.exit_pattern)]
        sequence=[('PROTOCOL_ENTER',[spec.protocol.enter_pattern])] if spec.protocol.enter_pattern else []
        for op in spec.operations:
            names=[op.reference]
            if op.kind=='TRIM': names.extend(a.reference for a in spec.search.per_point_actions)
            sequence.append((op.kind,[n for n in names if n]))
        if spec.protocol.exit_pattern: sequence.append(('PROTOCOL_EXIT',[spec.protocol.exit_pattern]))
        sequences[row+2]=sequence
        entries.append((row + 2, spec.cp, spec.item, 'TRIM', actions, [],
                        _storage_details(row + 2, spec, catalogs, editor, enabled=enabled_storage(row + 2))))
    for spec in catalogs.standard_verify_plans.values():
        rows = {}
        points = spec.shmoo.points if spec.shmoo else (spec,)
        for point in points:
            sequences[point.excel_row]=[(s.module,[n for _,n in _step_references((s,))[0]]) for s in point.steps]
            if spec.protocol_enter: sequences[point.excel_row].insert(0,('PROTOCOL_ENTER',[spec.protocol_enter.pattern.name]))
            if spec.protocol_exit: sequences[point.excel_row].append(('PROTOCOL_EXIT',[spec.protocol_exit.pattern.name]))
            actions, unexpanded = rows.setdefault(point.excel_row, ([], []))
            references, pending = _step_references(point.steps)
            actions.extend(references)
            unexpanded.extend(pending)
        for row, (actions, unexpanded) in rows.items():
            for use, run in (('PROTOCOL_ENTER', spec.protocol_enter), ('PROTOCOL_EXIT', spec.protocol_exit)):
                if run:
                    actions.append((use, run.pattern.name))
            symbol = next((getattr(step.module_plan, 'symbol', None) for point in points if point.excel_row == row
                           for step in point.steps if step.module == 'EXPORT' and step.module_plan), None)
            entries.append((row, spec.cp, spec.item, 'VERIFY', actions, list(dict.fromkeys(unexpanded)),
                            _storage_details(row, spec, catalogs, editor, 'VERIFY', symbol, enabled_storage(row))
                            + _module_details(spec, spec.steps, catalogs, editor)))
    for catalog_name, kind in (('dac_adjust_plans', 'DAC_ADJUST'), ('flash_window_plans', 'FLASH_WINDOW'),
                               ('trimdata_read_plans', 'TRIMDATA_READ'), ('main_plans', 'MAIN'),
                               ('trimdata_write_plans', 'TRIMDATA_WRITE'), ('uid_plans', 'UID'),
                               ('id_module_plans', 'ID')):
        for spec in getattr(catalogs, catalog_name, {}).values():
            actions, unexpanded = _step_references(spec.steps)
            sequences[spec.excel_row]=[(s.module,[n for _,n in _step_references((s,))[0]]) for s in spec.steps]
            if not actions and not any(s.module in {'DAC_ADJ', 'SETVPOSP'} for s in spec.steps):
                continue
            entries.append((spec.excel_row, spec.cp, spec.item, kind, actions, list(dict.fromkeys(unexpanded)),
                            _module_details(spec, spec.steps, catalogs, editor)))
    for catalog_name,kind in (('margin_plans','MARGIN'),('vt_plans','VT')):
        for spec in getattr(catalogs,catalog_name,{}).values():
            entry=test_entry(spec,kind,editor)
            entries.append(entry)
            sequences[spec.excel_row]=[(use,[name]) for use,name in entry[4]]
    result = []
    for row, cp, item, kind, actions, unexpanded, details in sorted(entries):
        refs = {}
        for use, name in actions:
            if not name:
                continue
            refs.setdefault(name, [])
            if use not in refs[name]:
                refs[name].append(use)
        patterns = []
        for name, uses in refs.items():
            locations = []
            for source_row in sources.get(name, []):
                if 1 <= source_row <= pattern_sheet.get('maxRow', 0):
                    if description_col:
                        coordinate = f'{get_column_letter(description_col)}{source_row}'
                    else:
                        coordinate = f'row {source_row}'
                    locations.append({'sheet': 'pattern', 'coordinate': coordinate})
            patterns.append({'name': name, 'uses': uses, 'locations': locations,
                             'preview': previews.get(name, []),
                             'note': '' if locations else '执行计划引用了此资源，当前尚无可定位的 pattern 表来源；不代表缺失。'})
        slots = []
        for name, uses in refs.items():
            slot = _resource_slot(' / '.join(uses), 'pattern', '执行 Pattern', {'shared':[name]})
            # Lack of a worksheet coordinate alone cannot prove an external stub.
            slot['status'] = 'generated' if name in generated else 'referenced'
            slots.append(slot)
        result.append({'flowRow': row, 'cp': cp, 'item': item, 'kind': kind,
                       'patterns': patterns, 'unexpanded': unexpanded, 'details': details,
                       'resourceSlots': slots,
                       'stages': [{'label':use,'patterns':list(dict.fromkeys(names)),
                                   'details':[i for i,d in enumerate(details) if d['title']==use]}
                                  for use,names in sequences.get(row,[])]})
    return sorted(result + specialized_links(catalogs,editor,workbook),key=lambda item:item['flowRow'])
