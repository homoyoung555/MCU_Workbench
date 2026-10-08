"""Project one ordinary Pattern field onto the authoritative FB MAP position."""
import re

from openpyxl.utils import get_column_letter

import fb_map
from program_gen.core.fbist_pattern_plan import (
    _FIELD_RE, _blank, _description_lines, _field_width, _resolve_output,
    _row_valid,
)
from program_gen.workbook import WorkbookReader


_ADDR_RE = re.compile(r"[01]{7}\Z")


def fb_mapping(session, payload):
    """Return, or edit, one Pattern row without requiring a whole-workbook compile."""
    editor, row = session.editor, payload.get('row')
    if type(row) is not int or row < 2:
        raise ValueError('请选择普通生成 Pattern 的有效行')
    raw = editor.apply(payload.get('baseSha256'), payload.get('changes'),
                       payload.get('appendRows', []), payload.get('flowOrder', []),
                       payload.get('trimStorageMappings'), payload.get('verifyStorageMappings'), payload.get('patternEdits'), payload.get('patternLinks'))
    workbook = WorkbookReader(byte_loader=lambda _: raw).read('gui-fb-mapping.xlsx')
    view = _view(workbook, row)
    if 'target' not in payload:
        editor.ensure_source_current()
        return view

    line, target_id = payload.get('line'), payload.get('target')
    field = next((field for field in view['fields'] if field['line'] == line), None)
    target = next((target for target in view['targets'] if target['id'] == target_id), None)
    if type(line) is not int or field is None or not isinstance(target_id, str) or target is None:
        raise ValueError('请选择当前字段和有效 FB MAP 物理字段')
    if field['width'] != target['width']:
        raise ValueError('FB MAP 目标位宽与当前字段值位宽不一致，不能自动截断或扩展')
    if len(view['fields']) > 1 and target['address'] != view['address']:
        raise ValueError('同一 Pattern 行含多个字段，只能映射到同一地址；跨地址会移动其他字段')

    replacement = f"{target['high']},{target['name']}[{field['width'] - 1}:0]={field['value']}"
    fields = list(view['fields'])
    fields[field['line']].update(name=target['name'], start=target['high'], high=target['high'],
                                 low=target['low'], current=target['id'], status='mapped', text=replacement)
    _validate_layout(fields)
    source_lines = str(_row_description(workbook, row)).splitlines()
    nonempty = [index for index, text in enumerate(source_lines) if text.strip()]
    source_lines[nonempty[field['line']]] = replacement
    description = '\n'.join(source_lines)
    sheet = editor.by_name['pattern']
    address_col, description_col = _columns(sheet)
    original_addr = sheet['rows'][row - 1][address_col]['value']
    original_description = sheet['rows'][row - 1][description_col]['value']
    address_cell = f'{get_column_letter(address_col + 1)}{row}'
    description_cell = f'{get_column_letter(description_col + 1)}{row}'
    desired_address = target['address'] if len(fields) == 1 else view['address']
    edits = [change for change in payload['changes'] if (change['sheet'], change['cell']) not in {
        ('pattern', address_cell), ('pattern', description_cell)
    }]
    if str(original_addr).strip() != desired_address:
        edits.append({'sheet': 'pattern', 'cell': address_cell, 'old': original_addr,
                      'value': desired_address})
    edits.append({'sheet': 'pattern', 'cell': description_cell, 'old': original_description,
                  'value': description})
    candidate = editor.apply(payload.get('baseSha256'), edits,
                             payload.get('appendRows', []), payload.get('flowOrder', []),
                             payload.get('trimStorageMappings'), payload.get('verifyStorageMappings'), payload.get('patternEdits'), payload.get('patternLinks'))
    candidate_view = _view(WorkbookReader(byte_loader=lambda _: candidate).read('gui-fb-mapping.xlsx'), row)
    editor.ensure_source_current()
    return {**candidate_view, 'changes': edits,
            'message': '已按 FB MAP 更新 Pattern 地址/字段；尚未执行整本检查'}


def _view(workbook, row):
    try:
        sheet, frame = workbook['pattern'], workbook.dataframe('pattern')
    except KeyError as exc:
        raise ValueError('当前工作簿没有 pattern 表') from exc
    if row > sheet.max_row or row - 2 not in frame.index:
        raise ValueError('Pattern 行不存在')
    if not _row_valid(sheet, row):
        raise ValueError('该 Pattern 行不是普通生成有效行')
    item, role = _inherited_context(sheet, frame, row)
    record = frame.loc[row - 2]
    description = record.get('description')
    if _blank(description):
        raise ValueError('该 Pattern 行没有可映射字段')
    output_pattern, excluded = _resolve_output(item, role, description, row)
    if excluded or output_pattern is None:
        raise ValueError('该 Pattern 行不生成普通输出，不能进行 FB MAP 映射')
    address = str(record.get('addr') or '').strip()
    if not _ADDR_RE.fullmatch(address):
        raise ValueError('Pattern 地址必须为 7 位二进制，才能映射物理字段')
    index = fb_map.build_index_from_workbook(workbook)
    if not index:
        raise ValueError('当前工作簿没有可用 FB MAP')
    targets = _targets(index)
    fields = _fields(description, int(address, 2), index, targets)
    return {'row': row, 'item': item, 'role': role, 'output_pattern': output_pattern,
            'address': address, 'fields': fields, 'targets': targets}


def _row_description(workbook, row):
    return workbook.dataframe('pattern').loc[row - 2].get('description')


def _inherited_context(sheet, frame, selected_row):
    item = role = ''
    for row in range(2, selected_row + 1):
        if not _row_valid(sheet, row) or row - 2 not in frame.index:
            continue
        record = frame.loc[row - 2]
        if not _blank(record.get('item')):
            item = str(record['item']).strip()
        if not _blank(record.get('pattern')):
            role = str(record['pattern']).strip()
    if not item or not role:
        raise ValueError('该 Pattern 行缺少可继承的 item 或 pattern role')
    return item, role


def _targets(index):
    targets = []
    for name, hits in index.items():
        for address, high, low, register in hits:
            if not 0 <= address < 128:
                continue
            target = {'name': name, 'register': '' if register is None else str(register),
                      'address': format(address, '07b'), 'high': high, 'low': low,
                      'width': high - low + 1}
            # Name remains part of the id: aliases at one physical span are explicit choices.
            target['id'] = f"{target['address']}:{high}:{low}:{name}"
            targets.append(target)
    return sorted(targets, key=lambda target: (target['address'], -target['high'], target['low'], target['name']))


def _fields(description, address, index, targets):
    fields = []
    physical = {}
    for target in targets:
        physical.setdefault((target['address'], target['high'], target['low']), []).append(target)
    for line, text in enumerate(_description_lines(description)):
        parsed = _parse_field(text, address, index)
        high = parsed['start']
        low = high - parsed['width'] + 1 if high is not None else None
        current = physical.get((format(address, '07b'), high, low), [])
        name = fb_map.normalize_field_name(parsed['name'])
        matching = next((target for target in current if target['name'] == name), None)
        status = 'mapped' if matching else ('name_mismatch' if current else 'unmapped')
        fields.append({**parsed, 'line': line, 'high': high, 'low': low,
                       'current': matching['id'] if matching else None,
                       'candidates': [target['id'] for target in current], 'status': status, 'text': text})
    if not fields:
        raise ValueError('该 Pattern 行没有可解析字段')
    return fields


def _parse_field(text, address, index):
    match = _FIELD_RE.fullmatch(text)
    if match:
        name, value = match.group(2).strip(), match.group(3) + ('*' if match.group(4) else '')
        start = int(match.group(1))
    else:
        parsed = fb_map.parse_pattern_description(text)
        if len(parsed) != 1 or parsed[0][1] is not None or '=' not in text:
            raise ValueError(f'无法解析 Pattern 字段行 {text!r}')
        name, value = parsed[0][0], text.split('=', 1)[1].strip()
        if not re.fullmatch(r'[01xX]+\*?', value):
            raise ValueError(f'无法解析 Pattern 字段行 {text!r}')
        try:
            start = fb_map.resolve_field_start_bit(name, address, index)
        except ValueError:
            start = None
    width = _field_width(name, len(value.rstrip('*')))
    if width != len(value.rstrip('*')):
        raise ValueError(f'字段 {name} 位宽与值不一致')
    return {'name': name, 'value': value, 'dynamic': '*' in value or 'x' in value.lower(),
            'start': start, 'width': width}


def _validate_layout(fields):
    occupied = set()
    for field in fields:
        if field['start'] is None:
            raise ValueError('其他字段位段无法确认，不能安全修改本行映射')
        low = field['start'] - field['width'] + 1
        if low < 0 or field['start'] >= 16:
            raise ValueError('目标字段超出 16 位寄存器范围')
        overlap = occupied.intersection(range(low, field['start'] + 1))
        if overlap:
            raise ValueError('目标字段与同一 Pattern 行已有字段位段重叠')
        occupied.update(range(low, field['start'] + 1))


def _columns(sheet):
    labels = {column['label']: index for index, column in enumerate(sheet['columns'])}
    try:
        return labels['addr'], labels['description']
    except KeyError as exc:
        raise ValueError('pattern 表缺少 addr 或 description 列') from exc
