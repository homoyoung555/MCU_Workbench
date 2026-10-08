"""Apply a confirmed source name to the workbook draft, before saving a Link."""
from copy import deepcopy
from dataclasses import replace
import re

from openpyxl.utils import get_column_letter


def adopt_source_name(workspace, target, name, draft, preview_links):
    if not isinstance(name, str) or not re.fullmatch(r'[A-Za-z_][A-Za-z0-9_]*', name):
        raise ValueError('来源 Pattern 名称不是有效标识符')
    if not isinstance(draft, dict):
        raise ValueError('更新 Pattern 名称需要当前工作簿草稿')
    result = deepcopy(draft)
    editor = workspace.session.editor
    if target['role'] == 'download':
        # A firmware phase owns one FLOW reference cell. Never replace equal text
        # elsewhere in the workbook, or mutate the loaded mother workbook.
        row = int(target['phase'])
        table = editor.by_name['FLOW']
        order = draft.get('flowOrder') or []
        original_row = order[row - 1] if order else row
        if original_row <= table['maxRow']:
            original = table['rows'][original_row - 1]
        else:
            additions = [op for op in draft.get('appendRows', []) if op['sheet'] == 'FLOW']
            operation = additions[original_row - table['maxRow'] - 1]
            original = table['rows'][operation['sourceRow'] - 1] if operation['copy'] else [dict(value=None) for _ in table['columns']]
        edits_by_cell = {c['cell']: c['value'] for c in draft.get('changes', []) if c['sheet'] == 'FLOW'}
        names = {n for values in target['names'].values() for n in values}
        columns = [i + 1 for i, column in enumerate(table['columns'])
                   if column['label'].lower() in ('hex', 'step')
                   and str(edits_by_cell.get(f'{get_column_letter(i + 1)}{original_row}', original[i]['value'])).strip() in names]
        if len(columns) != 1:
            raise ValueError('无法唯一定位固件名称单元格，请检查 FLOW 的 hex/step 引用')
        column = columns[0]
        cell = f'{get_column_letter(column)}{original_row}'
        old = original[column - 1]['value']
        edits = [c for c in draft.get('changes', []) if (c['sheet'], c['cell']) != ('FLOW', cell)]
        if str(old) != name:
            edits.append({'sheet': 'FLOW', 'cell': cell, 'old': old, 'value': name})
        result['changes'] = edits
    elif target['role'] == 'pattern':
        if any(values != [name] for values in target['names'].values()):
            raise ValueError('此资源名称由 Recipe 或内置协议决定，请编辑相应引用；未保存 Link')
    else:
        from program_gen.core.pattern_links import validate_pattern_links
        records = deepcopy(draft.get('patternLinks', editor.pattern_links))
        for platform, names in target['names'].items():
            for current in names:
                eligible = [r for r in records if r['cp'] == target['cp']
                            and (platform == 'shared' or r['platform'] in (platform, 'shared'))]
                matches = [r for r in eligible if r['pattern'] == current]
                if not matches:
                    matches = [r for r in eligible if r['original'] == current]
                for p in ('chroma', 's100'):
                    if sum(r['platform'] in (p, 'shared') for r in matches) > 1:
                        raise ValueError('多个原资源已使用同名 Pattern，请先解除冲突关联')
                if not matches:
                    record = {key: target[key] for key in ('cp', 'item', 'phase', 'role')}
                    record.update(platform=platform, original=current, pattern=name)
                    records.append(record)
                else:
                    for record in matches:
                        record['pattern'] = name
        result['patternLinks'] = validate_pattern_links(records)
    # Validate the exact output draft using the normal cell/metadata checks, but
    # do not parse or compile it just to change an already known resource name.
    raw = workspace._draft_bytes(result)
    links, firmware_plans = renamed_preview(workspace, preview_links, target, name, result.get('patternLinks', editor.pattern_links))
    targets = workspace._target_state(links)['targets']
    found = [t for t in targets if all(t[k] == target[k] for k in ('cp', 'item', 'phase', 'role'))
             and all(values == [name] for values in t['names'].values())]
    if len(found) != 1:
        raise ValueError('所选资源无法应用来源名称，未保存 Link')
    workspace._target_links_cache = (raw, links)
    workspace._target_firmware_plans = firmware_plans
    return result, found[0], targets


def renamed_preview(workspace, preview_links, target, name, records):
    """Project only a confirmed name edit; other draft changes use full compilation."""
    from gui.specialized_links import _project_slot_names, _firmware
    from program_gen.core.pattern_link_names import PatternLinkNameResolver
    links = deepcopy(preview_links)
    firmware_plans = workspace._target_firmware_plans
    if target['role'] == 'download':
        firmware_plans = dict(firmware_plans)
        replacements = {}
        for key, plan in firmware_plans.items():
            if plan.cp != target['cp'] or plan.item != target['item']:
                continue
            def renamed(phase):
                if getattr(phase, 'row', -2) + 2 == int(target['phase']) and hasattr(phase, 'pattern'):
                    return replace(phase, pattern=replace(phase.pattern, name=name, logical_identity=f'firmware-download:{name}'))
                return phase
            updated = replace(plan, phases=tuple(map(renamed, plan.phases)),
                              execution_steps=tuple(map(renamed, plan.execution_steps)))
            firmware_plans[key] = updated
            replacements.update((e['flowRow'], e) for e in _firmware(updated, workspace._target_editor))
        if not replacements:
            raise ValueError('固件阶段引用已变化，请重新打开关联')
        # Preserve existing SWD metadata projection on the other firmware phases.
        for i, entry in enumerate(links):
            if entry.get('specialized') and entry['flowRow'] in replacements:
                links[i] = _project_slot_names(replacements[entry['flowRow']], PatternLinkNameResolver(records))
    elif target['role'] != 'pattern':
        records = [dict(cp=target['cp'], item=target['item'], phase=target['phase'], role=target['role'],
                        platform=platform, original=original, pattern=name)
                   for platform, names in target['names'].items() for original in names]
        resolver = PatternLinkNameResolver(records)
        links = [_project_slot_names(entry, resolver) if entry.get('specialized') else entry for entry in links]
    return links, firmware_plans


def rebase_bindings(bindings, old_targets, new_targets):
    """Refresh names for shared references without changing their selected source."""
    old_by_id = {t['id']: t for t in old_targets}
    if len(old_targets) != len(new_targets):
        raise ValueError('名称修改改变了资源数量，未保存 Link')
    replacements = {old['id']: new for old, new in zip(old_targets, new_targets)
                    if all(old[k] == new[k] for k in ('cp', 'item', 'phase', 'role'))}
    result = deepcopy(bindings)
    for binding in result:
        if binding['target']['id'] in old_by_id:
            new = replacements.get(binding['target']['id'])
            if new is None:
                raise ValueError('资源次序已变化，未保存 Link')
            binding['target'] = new
    return result
