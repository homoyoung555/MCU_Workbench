"""GUI workbook snapshots and narrow OOXML cell edits; formula caches stay intact."""
from datetime import date, datetime, time, timedelta
from dataclasses import asdict
import json
from io import BytesIO
import hashlib
import math
from pathlib import Path
import re
from xml.sax.saxutils import escape
import xml.etree.ElementTree as ET
from zipfile import ZipFile

import openpyxl
from openpyxl.formula import Tokenizer
from openpyxl.utils import get_column_letter, coordinate_to_tuple, range_boundaries
from preflight.schema import TYPED, ENUM
from program_gen.core.project_mappings import MAPPINGS_PART, parse_storage_mappings, read_storage_mappings, read_verify_storage_mappings
from program_gen.core.pattern_links import PATTERN_LINKS_PART, patch_pattern_links, read_pattern_links, validate_pattern_links

# UI-only choices for literal keys consumed by flash_seq / addr_map / pvd_test.
# Hardware values, resource names and expressions remain open text fields.
CELL_CHOICES = {key: [''] + spec['enum'] + (['/'] if spec.get('slash_ok') else [])
                for key, spec in TYPED.items() if spec['type'] == ENUM}
CELL_CHOICES.update({
    ('PIN DEC', 'CP NO'): CELL_CHOICES[('FLOW', 'CP NO')],
    ('PIN DEC SINGEL', 'TESTER'): ['', 'Chroma', 'S100'],
    ('ADDR_MAP', 'data type'): ['', 'Reserved', 'fixed', 'lable'],
    ('pattern', 'pattern'): ['', 'trim mode', 'trim dac', 'verify mode', 'pump dac', 'reuse dac', 'read_status_mode', 'set_all_final_vposp'],
    ('FLASH_EPR', 'Description'): ['', 'default', 'clear page latches', 'normal mode', 'write page latches', 'PEP', 'erase page', 'bulk erase', 'program page', 'bulk program even/odd pages', 'bulk program all pages'],
    ('FLASH_EPR', 'space'): ['', 'NM', 'INFO', 'PAGE_BYTES', 'boot loader'],
    ('FLASH_EPR', 'opr_mode'): ['', 'pre', 'erase', 'program'],
    ('FLASH_EPR', 'chessboard_mode'): ['', 'PCKB', 'NCKB'],
    ('FLASH_EPR', 'fbist_mode'): ['', 'enter', 'exit'],
    ('FLASH_EPR', 'comp_flag'): ['', 'no', 'yes'],
    ('PVD', 'up/down'): ['', 'up', 'down'],
})

NS = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
REL = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'
NUMBER = re.compile(r'[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?\Z')


def sha(raw):
    return hashlib.sha256(raw).hexdigest().upper()


def formula_dependencies(workbook):
    """Conservatively freeze formula sheets/dependencies; never pretend to recalculate."""
    locked = set()
    for sheet in workbook:
        for row in sheet:
            for cell in row:
                if cell.data_type != 'f':
                    continue
                locked.add(sheet.title)
                try:
                    tokens = Tokenizer(cell.value).items
                    for token in tokens:
                        if token.type == 'FUNC' and token.value.upper() in ('INDIRECT(', 'OFFSET('):
                            return set(workbook.sheetnames)
                        if token.type != 'OPERAND' or token.subtype != 'RANGE':
                            continue
                        value = token.value
                        if '!' in value:
                            title, value = value.rsplit('!', 1)
                            title = title.strip("'").replace("''", "'")
                            if title not in workbook.sheetnames:
                                return set(workbook.sheetnames)
                            locked.add(title)
                        if not re.fullmatch(r'\$?[A-Z]+\$?\d+(?::\$?[A-Z]+\$?\d+)?', value):
                            return set(workbook.sheetnames)  # Named/dynamic/3D references need Excel.
                except Exception:
                    return set(workbook.sheetnames)
    return locked


class WorkbookEditor:
    def __init__(self, raw, name, source_path=None):
        self.raw, self.name, self.sha256 = raw, Path(name).name, sha(raw)
        self.source_path = Path(source_path) if source_path else None
        with ZipFile(BytesIO(raw)) as archive:
            if len(archive.infolist()) > 5000 or sum(p.file_size for p in archive.infolist()) > 256 * 1024**2:
                raise ValueError('工作簿展开后过大，当前GUI不支持')
            if 'xl/vbaProject.bin' in archive.namelist():
                raise ValueError('当前GUI只接受不含宏的xlsx母本')
            relationships = {e.attrib['Id']: e.attrib['Target'] for e in ET.fromstring(archive.read('xl/_rels/workbook.xml.rels'))}
            self.sheet_paths = {}
            for sheet in ET.fromstring(archive.read('xl/workbook.xml')).find(NS + 'sheets'):
                target = relationships[sheet.attrib[REL + 'id']]
                self.sheet_paths[sheet.attrib['name']] = target.lstrip('/') if target.startswith('/') else 'xl/' + target
        self.trim_storage_mappings = json.loads(json.dumps([asdict(m) for m in read_storage_mappings(raw)]))
        self.verify_storage_mappings = json.loads(json.dumps([asdict(m) for m in read_verify_storage_mappings(raw)]))
        self.pattern_links = read_pattern_links(raw)
        formulas = openpyxl.load_workbook(BytesIO(raw), data_only=False)
        cached = openpyxl.load_workbook(BytesIO(raw), data_only=True)
        try:
            if sum(s.max_row * s.max_column for s in formulas) > 500000:
                raise ValueError('工作簿有效矩形超过50万格，当前GUI不支持')
            locked = formula_dependencies(formulas)
            self.sheets = []
            self.by_name = {}
            for sheet in formulas:
                head = 2 if sheet.title == 'IOMUX' else 1
                note = ('含公式或为公式依赖，当前整表只读；请在Excel修改、重算并保存后重新载入。'
                        if sheet.title in locked else '可修改现有数据格；FLOW支持同CP测试定义拖动，保留格式及颜色。表头和合并从格只读。')
                merged_children = set()
                for area in sheet.merged_cells.ranges:
                    merged_children.update((r, c) for r in range(area.min_row, area.max_row + 1)
                                           for c in range(area.min_col, area.max_col + 1)
                                           if (r, c) != (area.min_row, area.min_col))
                rows = []
                for row in sheet:
                    cells = []
                    for c in row:
                        v = cached[sheet.title].cell(c.row, c.column).value
                        reason = ''
                        if sheet.title in locked: reason = note
                        elif c.row <= head: reason = '表头只读'
                        elif (c.row, c.column) in merged_children: reason = '合并单元格的从格，请编辑主格'
                        elif sheet.protection.sheet and c.protection.locked: reason = '受保护单元格，请在Excel中修改'
                        elif isinstance(v, (date, datetime, time, timedelta)): reason = '日期/时间格式请在Excel编辑'
                        if isinstance(v, (date, datetime, time, timedelta)): v = str(v)
                        cells.append({'value': v, 'formula': c.value if c.data_type == 'f' else None,
                                      'editable': not reason, 'reason': reason})
                    rows.append(cells)
                columns = []
                for c in range(1, sheet.max_column + 1):
                    labels = [str(sheet.cell(r, c).value) for r in range(1, head + 1) if sheet.cell(r, c).value is not None]
                    column = {'letter': get_column_letter(c), 'label': ' / '.join(labels)}
                    choices = CELL_CHOICES.get((sheet.title, column['label']))
                    if choices is not None:
                        column['choices'] = list(choices)
                    columns.append(column)
                info = {'name': sheet.title, 'maxRow': sheet.max_row, 'maxCol': sheet.max_column,
                        'headerRows': head, 'note': note, 'columns': columns, 'rows': rows,
                        'merges': [{'minRow': a.min_row, 'maxRow': a.max_row,
                                    'minCol': a.min_col, 'maxCol': a.max_col}
                                   for a in sheet.merged_cells.ranges]}
                # Append only: existing coordinates, formula caches and cross-sheet references never shift.
                merged_rows = {r for area in sheet.merged_cells.ranges for r in range(area.min_row, area.max_row+1)}
                info['appendableRows'] = ([r for r in range(head+1, sheet.max_row+1)
                                          if r not in merged_rows and all(c['editable'] for c in rows[r-1])]
                                         if sheet.title == 'FLOW' and sheet.title not in locked and not sheet.tables else [])
                info['resourceRows'] = ([r for r in range(head+1,sheet.max_row+1)
                    if all(c['editable'] or c['reason'].startswith('合并单元格的从格') for c in rows[r-1])]
                    if sheet.title in ('ADC','OPA','DAC','PVD') and sheet.title not in locked and not sheet.tables else [])
                self.sheets.append(info)
                self.by_name[sheet.title] = info
        finally:
            formulas.close()
            cached.close()

    def ensure_source_current(self):
        if self.source_path and self.source_path.read_bytes() != self.raw:
            raise ValueError('磁盘母本已变化，请重新载入后操作')

    def apply(self, base_sha, changes, append_rows=None, flow_order=None, trim_storage_mappings=None, verify_storage_mappings=None, pattern_edits=None, pattern_links=None):
        self.ensure_source_current()
        if base_sha != self.sha256:
            raise ValueError('母本版本不匹配，请重新载入')
        from gui.pattern_editor import validate_edits, pattern_items, patch_pattern_xml
        pattern_edits = validate_edits(self, [] if pattern_edits is None else pattern_edits)
        edited_names = {entry['item'] for entry in pattern_edits}
        edited_ranges = [(item['start'], item['end']) for item in pattern_items(self) if item['item'] in edited_names]
        mappings = self.trim_storage_mappings if trim_storage_mappings is None else [
            asdict(m) for m in parse_storage_mappings(trim_storage_mappings, "trimStorage")]
        mappings = json.loads(json.dumps(mappings))
        verify_mappings = self.verify_storage_mappings if verify_storage_mappings is None else json.loads(json.dumps([
            asdict(m) for m in parse_storage_mappings(verify_storage_mappings, "verifyStorage")]))
        mappings_changed = mappings != self.trim_storage_mappings or verify_mappings != self.verify_storage_mappings
        links = self.pattern_links if pattern_links is None else validate_pattern_links(pattern_links)
        links_changed = links != self.pattern_links
        if not isinstance(changes, list) or len(changes) > 10000:
            raise ValueError('单次改动最多10000格')
        append_rows = [] if append_rows is None else append_rows
        if not isinstance(append_rows, list) or len(append_rows) > 100:
            raise ValueError('单次最多追加100行数据')
        flow_order = [] if flow_order is None else flow_order
        if not isinstance(flow_order, list):
            raise ValueError('FLOW排序清单必须是数组')
        additions = {}
        for operation in append_rows:
            if not isinstance(operation, dict) or set(operation) != {'sheet','sourceRow','copy'}:
                raise ValueError('新增行清单格式不正确')
            flow = self.by_name.get(operation['sheet'], {})
            source = operation['sourceRow']
            if (operation['sheet'] not in ('FLOW','ADC','OPA','DAC','PVD') or type(source) is not int or
                    source not in flow.get('appendableRows' if operation['sheet']=='FLOW' else 'resourceRows', []) or type(operation['copy']) is not bool):
                raise ValueError('参考行不支持新增；请选择当前工作表允许的参考配置')
            if operation['sheet']=='ADC' and operation['copy']:
                raise ValueError('ADC合并模板仅允许填写完整的新通道组')
            additions.setdefault(operation['sheet'], []).append([dict(c, editable=True, reason='', value=c['value'] if operation['copy'] else None) for c in flow['rows'][source-1]])
        added = additions.get('FLOW', [])
        if any(self.by_name[n]['maxRow']+len(a)>1048576 for n,a in additions.items()) or sum(s['maxRow']*s['maxCol'] for s in self.sheets)+sum(len(a)*self.by_name[n]['maxCol'] for n,a in additions.items())>500000:
            raise ValueError('新增行后超过工作簿大小限制')
        updates, seen = {}, set()
        for change in changes:
            if not isinstance(change, dict) or set(change) != {'sheet', 'cell', 'old', 'value'}:
                raise ValueError('修改清单格式不正确')
            name, address = change['sheet'], change['cell']
            if not isinstance(name, str) or name not in self.by_name or not isinstance(address, str) or not re.fullmatch(r'[A-Z]+[1-9]\d*', address):
                raise ValueError('修改坐标不存在')
            if (name, address) in seen:
                raise ValueError(f'{name}!{address} 重复修改')
            seen.add((name, address))
            row, col = coordinate_to_tuple(address)
            if pattern_edits and name == 'pattern' and (col in (1, 4) or any(lo <= row <= hi for lo, hi in edited_ranges)):
                raise ValueError('Pattern结构编辑与原表单元格修改冲突，请在同一编辑方式中完成该item')
            info = self.by_name[name]
            extra = additions.get(name, [])
            if row > info['maxRow'] + len(extra) or col > info['maxCol']:
                raise ValueError(f'{name}!{address} 超出当前表范围')
            cell = (info['rows'][row-1] if row <= info['maxRow'] else extra[row-info['maxRow']-1])[col-1]
            if not cell['editable']:
                raise ValueError(f'{name}!{address}: {cell["reason"]}')
            if change['old'] != cell['value'] or isinstance(change['old'], bool) != isinstance(cell['value'], bool):
                raise ValueError(f'{name}!{address} 原值冲突')
            value = change['value']
            if not isinstance(value, str) or len(value) > 32767 or re.search(r'[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff]', value):
                raise ValueError(f'{name}!{address} 内容过长或包含无效字符')
            if value.lstrip().startswith('='):
                raise ValueError(f'{name}!{address} 暂不允许新增公式，请在Excel中修改并重算')
            kind = 'inlineStr'
            if value == '':
                kind = 'empty'
            elif isinstance(cell['value'], bool):
                if value.lower() not in ('true', 'false', '1', '0'):
                    raise ValueError(f'{name}!{address} 需要true/false')
                kind, value = 'b', '1' if value.lower() in ('true', '1') else '0'
            elif (cell['value'] is None or isinstance(cell['value'], (int, float))) and NUMBER.fullmatch(value):
                if not math.isfinite(float(value)):
                    raise ValueError(f'{name}!{address} 数值溢出')
                # Keep explicit '+' and leading zeros as domain text (limits/IDs).
                if not value.startswith('+') and not re.match(r'-?0\d', value) and len(re.sub(r'\D', '', value)) <= 15:
                    kind = 'n'
            updates.setdefault(name, {})[address] = (kind, value)
        flow = self.by_name.get('FLOW')
        total_rows = flow['maxRow'] + len(added) if flow else 0
        if flow_order:
            if flow is None:
                raise ValueError('当前工作簿没有FLOW表，不能调整执行顺序')
            if (any(type(row) is not int or not 1 <= row <= total_rows for row in flow_order)
                    or len(set(flow_order)) != len(flow_order)):
                raise ValueError('FLOW顺序必须是有效原行号且不能重复；省略的行表示删除')
            if flow_order[0] != 1:
                raise ValueError('FLOW表头不可移动')
            if flow_order != list(range(1, total_rows + 1)):
                _validate_flow_order(flow_order, flow, _flow_values(flow, added, updates),
                                     set(flow['appendableRows']) | set(range(flow['maxRow'] + 1, total_rows + 1)))
        if not pattern_edits and not mappings_changed and not links_changed and not updates and not additions and (not flow_order or flow_order == list(range(1, total_rows + 1))):
            return self.raw
        result = BytesIO()
        with ZipFile(BytesIO(self.raw)) as source, ZipFile(result, 'w') as output:
            output.comment = source.comment
            paths = {self.sheet_paths[name]: values for name, values in updates.items()}
            for part in source.infolist():
                if mappings_changed and part.filename == MAPPINGS_PART:
                    continue
                if links_changed and part.filename == PATTERN_LINKS_PART:
                    continue
                raw = source.read(part)
                for name in additions:
                    if part.filename == self.sheet_paths[name]:
                        info = self.by_name[name]
                        raw = _append_flow_rows(raw.decode('utf-8-sig'), [op for op in append_rows if op['sheet']==name],
                                                info['maxRow'], info['maxCol']).encode('utf-8')
                if part.filename in paths:
                    raw = _patch_sheet(raw.decode('utf-8-sig'), paths[part.filename]).encode('utf-8')
                if pattern_edits and part.filename == self.sheet_paths['pattern']:
                    raw = patch_pattern_xml(raw.decode('utf-8-sig'), self, pattern_edits).encode('utf-8')
                if flow_order and flow_order != list(range(1, total_rows + 1)) and part.filename == self.sheet_paths['FLOW']:
                    raw = _reorder_flow_rows(raw.decode('utf-8-sig'), flow_order).encode('utf-8')
                output.writestr(part, raw)
            if mappings_changed and (mappings or verify_mappings):
                output.writestr(MAPPINGS_PART, json.dumps({'version': 3, 'trimStorage': mappings, 'verifyStorage': verify_mappings},
                    ensure_ascii=False, sort_keys=True).encode('utf-8'))
        return patch_pattern_links(result.getvalue(), links) if links_changed else result.getvalue()


def _flow_values(flow, added, updates):
    values = [[cell['value'] for cell in row] for row in flow['rows']]
    values.extend([[cell['value'] for cell in row] for row in added])
    for address, (kind, value) in updates.get('FLOW', {}).items():
        row, col = coordinate_to_tuple(address)
        values[row - 1][col - 1] = None if kind == 'empty' else value
    return values


def _validate_flow_order(order, flow, rows, allowed):
    labels = {column['label']: index for index, column in enumerate(flow['columns'])}
    try:
        cp, item, flow_type, recipe = (labels[name] for name in ('CP NO', 'item', 'type', 'Recipe'))
    except KeyError:
        raise ValueError('FLOW缺少CP NO、item、type或Recipe列，不能调整执行顺序') from None
    if any(source != destination and (source not in allowed or destination not in allowed)
           for destination, source in enumerate(order, 1)):
        raise ValueError('FLOW排序只能移动可追加的非合并、非公式、未保护行')
    retained = set(order)
    removed = set(range(2, len(rows) + 1)) - retained
    if removed - allowed:
        raise ValueError('不能删除只读、公式或合并行')
    cp_of = lambda source: str(rows[source - 1][cp] or '').strip().casefold()
    original_cps = [cp_of(source) for source in range(2, flow['maxRow'] + 1) if source in retained]
    ordered_cps = [cp_of(source) for source in order[1:] if source <= flow['maxRow']]
    if original_cps != ordered_cps:
        raise ValueError('FLOW排序不能跨CP移动测试定义')
    for index, source in enumerate(order[1:], 1):
        if source <= flow['maxRow'] or not cp_of(source):
            continue
        neighbors = [cp_of(order[i]) for i in (index-1, index+1) if 0 < i < len(order)]
        if neighbors and cp_of(source) not in neighbors:
            raise ValueError('新增行必须插入相应CP内')
    groups, source_group, source_key = [], {}, {}
    previous = None
    for source in range(2, len(rows) + 1):
        value = rows[source - 1]
        key = _flow_group_key(value[cp], value[item], value[flow_type], value[recipe])
        if key is None:
            continue
        if key != previous:
            groups.append([])
            previous = key
        groups[-1].append(source)
        source_group[source] = len(groups) - 1
        source_key[source] = key
    # Only native ADC/firmware phase groups own an editable internal order.
    native = {i: _native_phase_family(flow, rows, group) for i, group in enumerate(groups)}
    for i, group in enumerate(groups):
        if not native[i] or not group or any(r <= flow['maxRow'] for r in group):
            continue
        matches = [j for j, other in enumerate(groups[:i]) if other and
                   source_key[other[0]] == source_key[group[0]] and native[j] == native[i]]
        if len(matches) == 1:
            target = matches[0]
            for r in group: source_group[r] = target
            groups[target].extend(group)
            groups[i] = []
    seen, last_group, last_key, output_groups = set(), None, None, {index: [] for index in range(len(groups))}
    for source in order[1:]:
        group = source_group.get(source)
        if group is None:
            continue
        key = source_key[source]
        if group != last_group:
            if group in seen:
                raise ValueError('FLOW排序不能拆分同一测试定义的连续行')
            if last_key == key:
                raise ValueError('FLOW排序不能合并原本独立的同类测试定义')
            seen.add(group)
        output_groups[group].append(source)
        last_group, last_key = group, key
    if any(not native[index] and output_groups[index] != [r for r in rows_in_group if r in retained] for index, rows_in_group in enumerate(groups)):
        raise ValueError('FLOW排序必须保持同一测试定义内的行顺序')


def _native_phase_family(flow, rows, group):
    import pandas as pd
    from program_gen.flow_kinds import compile_flow_rows
    records = [{c['label']: value for c, value in zip(flow['columns'], rows[r - 1])} for r in group]
    try:
        kinds = {op.typed_kind for op in compile_flow_rows(pd.DataFrame(records),
                 workbook_identity='phase-order', source_sha256='phase-order')}
    except ValueError:
        return None
    if kinds == {'ADC_TEST'}: return 'adc'
    if kinds and kinds <= {'VERIFY_HEX', 'VERIFY_KEIL'}: return 'firmware'
    return None


def _flow_group_key(cp, item, flow_type, recipe):
    cp, item, flow_type = (str(value or '').strip() for value in (cp, item, flow_type))
    if not cp or not item or not flow_type:
        return None
    if recipe in (None, ''):
        return cp.upper(), item.casefold(), flow_type.casefold(), 'BASE'
    from flow_recipe import parse_recipe
    try:
        shmoo = any(call.named_parameter('SHMOO') == 'STEP' for call in parse_recipe(recipe))
    except (TypeError, ValueError) as error:
        raise ValueError('FLOW Recipe无法解析，不能调整执行顺序：' + str(error)) from None
    return cp.upper(), item.casefold(), flow_type.casefold(), 'SHMOO' if shmoo else 'BASE'


def _reorder_flow_rows(xml, order):
    _validate_positional_dependencies(xml, order)
    pattern = re.compile(r'<(?P<tag>(?:[\w.-]+:)?row)\b(?=[^>]*\br=["\'](\d+)["\'])[^>]*(?:/>|>.*?</(?P=tag)>)', re.S)
    matches = list(pattern.finditer(xml))
    by_row = {int(match[2]): match for match in matches}
    if set(by_row) != set(range(1, max(by_row) + 1)) or not set(order) <= set(by_row):
        raise ValueError('FLOW XML行不连续，当前GUI不能安全调整执行顺序')
    result = [xml[:matches[0].start()]]
    for destination, source in enumerate(order, 1):
        row = by_row[source][0]
        row = re.sub(r'\br=(["\'])' + str(source) + r'\1', f'r="{destination}"', row, count=1)
        row = re.sub(r'(<(?:[\w.-]+:)?c\b[^>]*\br=["\'])([A-Z]+)\d+(["\'])',
                     lambda match: match[1] + match[2] + str(destination) + match[3], row)
        result.append(row)
    result.append(xml[matches[-1].end():])
    reordered = ''.join(result)
    reordered = re.sub(r'(<(?:[\w.-]+:)?dimension\b[^>]*\bref=["\'])([A-Z]+\d+:[A-Z]+)\d+(["\'])',
                       lambda m: m[1] + m[2] + str(len(order)) + m[3], reordered)
    reordered = _remap_flow_validations(reordered, order, max(by_row))
    ET.fromstring(reordered)
    return reordered


def _validate_positional_dependencies(xml, order):
    max_row = max(order)
    unsupported = ('sheetProtection', 'mergeCells', 'hyperlinks', 'conditionalFormatting', 'autoFilter', 'tableParts',
                   'legacyDrawing', 'drawing', 'rowBreaks', 'colBreaks', 'extLst')
    if any(re.search(r'<(?:[\w.-]+:)?' + tag + r'\b', xml) for tag in unsupported):
        raise ValueError('FLOW含有依赖物理行号的对象，当前GUI不能安全调整执行顺序')
    root = ET.fromstring(xml)
    validations = root.find(NS + 'dataValidations')
    if validations is None:
        return
    for validation in validations:
        if any(re.search(r'(?<![A-Z0-9_])\$?[A-Z]{1,3}\$?\d+', child.text or '')
               and not (child.text or '').startswith('"') for child in validation):
            raise ValueError('FLOW数据验证引用了单元格，当前GUI不能安全调整执行顺序')


def _remap_flow_validations(xml, order, source_count):
    root = ET.fromstring(xml)
    validations = root.find(NS + 'dataValidations')
    if validations is None:
        return xml
    for validation in list(validations):
        references = []
        for reference in validation.attrib.get('sqref', '').split():
            low_col, low_row, high_col, high_row = range_boundaries(reference)
            if None in (low_col, low_row, high_col, high_row):
                raise ValueError('FLOW数据验证需要明确的单元格范围')
            mapped = [i for i, source in enumerate(order, 1) if low_row <= source <= high_row]
            spans = []
            for row in mapped:
                if spans and spans[-1][1] == row - 1:
                    spans[-1][1] = row
                else:
                    spans.append([row, row])
            if high_row > source_count:
                spans.append([max(low_row, source_count + 1) + len(order) - source_count,
                              min(1048576, high_row + len(order) - source_count)])
            for start, end in spans:
                references.append(f'{get_column_letter(low_col)}{start}:{get_column_letter(high_col)}{end}')
        if references:
            validation.set('sqref', ' '.join(references))
        else:
            validations.remove(validation)
    validations.set('count', str(len(validations)))
    replacement = ET.tostring(validations, encoding='unicode') if len(validations) else ''
    return re.sub(r'<(?P<tag>(?:[\w.-]+:)?dataValidations)\b[^>]*(?:/>|>.*?</(?P=tag)>)',
                  lambda _: replacement, xml, flags=re.S)


def _append_flow_rows(xml, operations, max_row, max_col):
    prefix = re.search(r'<((?:[\w.-]+:)?)worksheet\b', xml)[1]
    original = xml
    additions = []
    for index, operation in enumerate(operations, max_row+1):
        source = operation['sourceRow']
        match = re.search(r'<(?P<tag>(?:[\w.-]+:)?row)\b(?=[^>]*\br=["\']'+str(source)+
                          r'["\'])[^>]*(?:/>|>.*?</(?P=tag)>)', original, re.S)
        if not match:
            raise ValueError('参考行没有可复制的XML内容')
        row = re.sub(r'\br=(["\'])'+str(source)+r'\1', f'r="{index}"', match[0], count=1)
        row = re.sub(r'(<(?:[\w.-]+:)?c\b[^>]*\br=["\'])([A-Z]+)\d+(["\'])',
                     lambda m: m[1]+m[2]+str(index)+m[3], row)
        if not operation['copy']:
            # Empty the cloned cells, keeping the original style IDs and row formatting.
            row = re.sub(r'<(?P<tag>(?:[\w.-]+:)?c)\b[^>]*(?:/>|>.*?</(?P=tag)>)',
                         lambda m: re.sub(r'\s+t=(["\']).*?\1','',m[0].split('>',1)[0].rstrip('/'))+'/>',row,flags=re.S)
        additions.append(row)
    closing = f'</{prefix}sheetData>'
    xml = xml.replace(closing, ''.join(additions)+closing, 1)
    xml = re.sub(r'(<(?:[\w.-]+:)?dimension\b[^>]*\bref=["\'])[^"\']+(["\'])',
                 lambda m: m[1]+f'A1:{get_column_letter(max_col)}{max_row+len(operations)}'+m[2],xml,count=1)
    ET.fromstring(xml)
    return xml


def _patch_sheet(xml, updates):
    prefix = re.search(r'<((?:[\w.-]+:)?)worksheet\b', xml)[1]

    def pattern(tag, coordinate):
        return re.compile(r'<(?P<tag>(?:[\w.-]+:)?' + tag + r')\b(?=[^>]*\br=["\']' +
                          re.escape(coordinate) + r'["\'])[^>]*(?:/>|>.*?</(?P=tag)>)', re.S)

    rows = {}
    for address, value in updates.items():
        rows.setdefault(coordinate_to_tuple(address)[0], {})[address] = value
    for row_number, values in sorted(rows.items()):
        match = pattern('row', str(row_number)).search(xml)
        row = match[0] if match else f'<{prefix}row r="{row_number}"></{prefix}row>'
        if row.endswith('/>'): row = row[:-2] + f'></{prefix}row>'
        for address, (kind, value) in sorted(values.items(), key=lambda item: coordinate_to_tuple(item[0])[1]):
            existing = pattern('c', address).search(row)
            tag = existing['tag'] if existing else prefix + 'c'
            opening = existing[0].split('>', 1)[0].rstrip('/') if existing else f'<{tag} r="{address}"'
            opening = re.sub(r'\s+t=(["\']).*?\1', '', opening)
            if kind == 'empty':
                replacement = opening + '/>'
            elif kind == 'inlineStr':
                replacement = opening + f' t="inlineStr"><{prefix}is><{prefix}t xml:space="preserve">' + escape(value) + f'</{prefix}t></{prefix}is></{tag}>'
            else:
                replacement = opening + f' t="{kind}"><{prefix}v>' + value + f'</{prefix}v></{tag}>'
            if existing:
                row = row[:existing.start()] + replacement + row[existing.end():]
            else:
                following = next((c for c in re.finditer(r'<(?:[\w.-]+:)?c\b[^>]*\br=["\']([A-Z]+\d+)["\']', row)
                                  if coordinate_to_tuple(c[1])[1] > coordinate_to_tuple(address)[1]), None)
                position = following.start() if following else row.rfind('</')
                row = row[:position] + replacement + row[position:]
        if match:
            xml = xml[:match.start()] + row + xml[match.end():]
        else:
            following = next((r for r in re.finditer(r'<(?:[\w.-]+:)?row\b[^>]*\br=["\'](\d+)["\']', xml)
                              if int(r[1]) > row_number), None)
            position = following.start() if following else xml.index(f'</{prefix}sheetData>')
            xml = xml[:position] + row + xml[position:]
    ET.fromstring(xml)  # Reject malformed output before any filesystem write.
    return xml
