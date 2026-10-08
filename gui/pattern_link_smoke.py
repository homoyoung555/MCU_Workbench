"""Program-side tool Link acceptance in the real packaged WebView2 window."""
import json
import base64
from time import perf_counter
from io import BytesIO

import openpyxl

from gui.desktop_smoke import until
from program_gen.core.pattern_links import read_pattern_links


def check_pattern_links(app, js, check):
    session = app.service.session
    def draft():
        return js('window.MCUProgram.draftPayload()')

    def swd_targets():
        return session.swd.targets(draft())

    def target_named(cp, item, role, name=None):
        for target in swd_targets()['targets']:
            if target['cp'] != cp or target['item'] != item or target['role'] != role:
                continue
            if name is None or name in sum(target['names'].values(), []):
                return target
        raise AssertionError(f'找不到 {cp}/{item}/{role} Pattern目标')

    def target_name(target):
        return next(iter(sum(target['names'].values(), [])))

    # The broader desktop smoke leaves deliberate edits in its isolated session.
    js('(()=>{while(window.MCUProgram.status().dirty&&!document.getElementById("undo-btn").disabled)document.getElementById("undo-btn").click();return true;})()')
    adc_rows = [item for item in session.links if item['cp'] == 'CP3' and item['item'] == 'Adc_Test']
    for adc in adc_rows:
        js('window.MCUProgram.navigateFlow(' + str(adc['flowRow']) + ')')
        names = list(dict.fromkeys(name for slot in adc['resourceSlots'] if slot['status'] != 'generated' for values in slot['names'].values() for name in values))
        for name in names:
            selector = json.dumps('[data-resource-link="' + name + '"]')
            until(lambda: js('!!document.querySelector('+selector+')'), 'ADC resource card')
            js('document.querySelector('+selector+').click()')
            until(lambda: js('!!document.querySelector("#swd-link-dialog[open]")'), 'ADC shared target at FLOW '+str(adc['flowRow'])+' '+name)
            check('ADC FLOW '+str(adc['flowRow'])+' 卡片可关联 '+name, js('document.querySelector("#swd-link-dialog [data-target-label]").textContent.includes('+json.dumps(name)+')'))
            js('document.querySelector("#swd-link-dialog [data-cancel]").click()')
    table_item = next(item for item in session.links if any(p.get('locations') and any(slot['status'] == 'generated' and p['name'] in sum(slot['names'].values(), []) for slot in item.get('resourceSlots', [])) for p in item['patterns']))
    js('window.MCUProgram.navigateFlow(' + str(table_item['flowRow']) + ')')
    until(lambda: js('!!document.querySelector(".pattern-card-main[data-fbist-link]")'), 'FBIST Link card')
    check('存储映射独立显示且没有重复顶部Link', js('!!document.querySelector("section[data-link-column=storage]") && !document.querySelector("[data-pattern-link]")'))
    js('document.querySelector(".pattern-card-main[data-fbist-link]").click()')
    until(lambda: js('!!document.querySelector("#fbist-link-dialog[open] [data-fbist-choice]")'), 'FBIST candidates')
    check('FBIST 卡片打开Link候选且不跳源表', js('document.getElementById("row-selection").textContent.includes("FLOW") && !window.MCUProgram.status().dirty'))
    js('document.querySelector("#fbist-link-dialog [data-cancel]").click()')
    check('取消FBIST选择不改变草稿', not js('window.MCUProgram.status().dirty'))
    js('document.querySelector(".pattern-card-main[data-fbist-link]").click()')
    until(lambda: js('!!document.querySelector("#fbist-link-dialog[open] [data-fbist-choice]")'), 'FBIST reopen')
    js('Array.from(document.querySelectorAll("#fbist-link-dialog [data-fbist-choice]")).find(b=>!b.textContent.includes("当前关联")).click()')
    check('选择FBIST候选尚未写入草稿', not js('window.MCUProgram.status().dirty'))
    js('document.querySelector("#fbist-link-dialog [data-apply]").click()')
    until(lambda: js('!document.querySelector("#fbist-link-dialog[open]")'), 'FBIST explicit confirmation')
    check('确认FBIST关联写入可撤销Recipe草稿', js('window.MCUProgram.status().dirty'))
    js('document.getElementById("undo-btn").click()')
    check('FBIST关联一次撤销恢复原稿', not js('window.MCUProgram.status().dirty'))
    js('document.querySelector(".pattern-card-more [data-link-pattern]").click()')
    check('编辑源表次要入口仍定位配置表', js('document.getElementById("row-selection").textContent.includes("pattern")'))
    js('document.getElementById("link-return").click()')
    target = target_named('CP1', 'LSI_IWDG_Test', 'download')
    resource = next(iter(target['names'].values()))[0]
    selector = json.dumps('[data-resource-link="' + resource + '"]')
    raw = b':020000040800F2\n:0400000001020304F2\n:00000001FF\n'
    task = app.service.hex.add_input(raw, 'link_acceptance.hex')
    app.service.hex.start({'taskIds': [task['id']], 'profileId': 'fbist-pt104-flash',
                          'pins': ['PA11', 'PA13', 'NRST'], 'outputPlatform': 'both'})
    until(lambda: not app.service.hex.state()['running'], 'HEX conversion for Link')
    task = app.service.hex.tasks[task['id']]
    check('Link 验收 HEX 双平台转换完成', task['status'] == 'completed')
    js('window.MCUProgram.navigateFlow(' + str(target['row']) + ')')
    until(lambda: js('!!document.querySelector(' + selector + ')'), 'download Link button')
    check('HEX 下载行不再显示普通FBIST映射入口', js('!document.querySelector("#linked-body [data-trim-pattern]")'))
    js('document.querySelector(' + selector + ').click()')
    until(lambda: js('!!document.querySelector("#swd-link-dialog[open] [data-generation]")'), 'Pattern Link modal')
    check('LSI_IWDG 卡片直接打开固件选择器且目标正确', js('!document.querySelector("#trim-pattern-dialog[open]") && document.querySelector("#swd-link-dialog [data-target]").textContent.includes("LSI_IWDG_V01_00")'))
    check('固件下载弹窗同时提供 SWD 与 HEX 工具', js('!!document.querySelector("#swd-link-dialog [data-open=hex]") && !!document.querySelector("#swd-link-dialog [data-open=swd]")'))
    js('document.querySelector("#swd-link-dialog [data-search]").value="no_match";document.querySelector("#swd-link-dialog [data-search]").dispatchEvent(new Event("input"))')
    check('关联结果搜索显示空状态', js('!!document.querySelector(".pattern-picker-empty")'))
    js('document.querySelector("#swd-link-dialog [data-search]").value="";document.querySelector("#swd-link-dialog [data-search]").dispatchEvent(new Event("input"))')
    js('document.querySelector(' + json.dumps('[data-choice="hex:' + task['id'] + '"]') + ').click()')
    link_started = perf_counter()
    js('document.querySelector("#swd-link-dialog [data-apply]").click()')
    until(lambda: js('!document.querySelector("#swd-link-dialog[open]")'), 'HEX Link saved')
    link_elapsed = perf_counter() - link_started
    check(f'HEX确认关联响应 {link_elapsed:.2f}秒（小于3秒）', link_elapsed < 3)
    target = target_named('CP1', 'LSI_IWDG_Test', 'download')
    resource = target_name(target)
    selector = json.dumps('[data-resource-link="' + resource + '"]')
    frozen = session.swd.freeze(draft())
    check('真实点击冻结 HEX 双平台结果', any(b['sourceKind'] == 'hex' and set(b['contents']) == {'chroma', 's100'} for b in frozen))
    check('HEX 卡片直接显示已关联文件来源', js('document.querySelector('+selector+').textContent.includes("link_acceptance.hex")'))
    app.service.hex.clear()
    check('清空 HEX 历史不影响程序已关联版本', session.swd.freeze(draft()) == frozen)
    js('document.querySelector(' + selector + ').click()')
    until(lambda: js('!!document.querySelector("#swd-link-dialog[open]")'), 'saved HEX Link')
    check('再次打开可查看已保存的 HEX 来源', js('document.querySelector("#swd-link-dialog [data-current]").textContent.includes("link_acceptance.hex")'))
    js('document.querySelector("#swd-link-dialog [data-open=hex]").click()')
    top = app.window.evaluate_js
    check('打开生成器保留返回目标', top('document.getElementById("wb-pattern-return").textContent.includes('+json.dumps(resource)+')'))
    top('document.querySelector("#wb-pattern-return button").click()')
    until(lambda: js('!!document.querySelector("#swd-link-dialog[open]")'), 'return from HEX tool')
    check('返回生成器未自动更换已保存关联', session.swd.freeze(draft()) == frozen)
    check('返回后仍锁定原资源', js('document.querySelector("#swd-link-dialog [data-target]").selectedOptions[0].textContent.includes('+json.dumps(resource)+')'))
    js('(()=>{const original=window.confirm;window.confirm=()=>true;document.querySelector("#swd-link-dialog [data-unlink]").click();window.confirm=original;})()')
    until(lambda: js('!document.querySelector("#swd-link-dialog[open]")'), 'unlink HEX')
    target = target_named('CP1', 'LSI_IWDG_Test', 'download', resource)
    check('解除 HEX Link 生效', not any(b['sourceKind'] == 'hex' for b in session.swd.freeze(draft())))
    # Link a saved definition with no generated artifact, then render on demand.
    steps = openpyxl.Workbook()
    steps.active.append(['pattern', 'addr', 'data', 'step', 'reset value', 'type', 'continue', 'description'])
    steps.active.append(['PT104_demo', '0X40021000', '0X00000001', '/', '/', 'W', 'N', 'demo'])
    steps_raw = BytesIO(); steps.save(steps_raw); steps.close()
    uploaded_steps = json.dumps(base64.b64encode(steps_raw.getvalue()).decode())
    generations_before = len(session.swd.data['generations'])
    js('document.querySelector(' + selector + ').click()')
    until(lambda: js('!!document.querySelector("#swd-link-dialog[open]")'), 'SWD source chooser')
    js('(()=>{const input=document.querySelector("#swd-link-dialog [data-swd-file]"),dt=new DataTransfer();dt.items.add(new File([new Uint8Array([1,2,3])],"bad_steps.xlsx"));input.files=dt.files;input.dispatchEvent(new Event("change",{bubbles:true}));return true;})()')
    until(lambda: js('document.querySelector("#swd-link-dialog [data-error]").textContent.includes("SWD导入失败")'), 'bad SWD import error')
    js('(()=>{const input=document.querySelector("#swd-link-dialog [data-swd-file]"),bytes=Uint8Array.from(atob(' + uploaded_steps + '),c=>c.charCodeAt(0)),dt=new DataTransfer();dt.items.add(new File([bytes],"program_steps.xlsx"));input.files=dt.files;input.dispatchEvent(new Event("change",{bubbles:true}));return true;})()')
    until(lambda: js('document.querySelector("#swd-link-dialog [data-document-filter]").textContent.includes("program_steps.xlsx")'), 'good SWD import')
    doc = next((item for item in session.swd.pattern_state()['sources'] if item['sourceLabel'] == 'program_steps.xlsx'), None)
    if not doc:
        raise AssertionError('SWD 新导入文件未出现在 Pattern 来源')
    js('document.querySelector("#swd-link-dialog [data-kind=swd]").click()')
    choice = doc
    choice_selector = json.dumps('#swd-link-dialog [data-choice="' + choice['id'] + '"]')
    check('SWD弹窗提供已载入文档筛选', js('!!document.querySelector("#swd-link-dialog [data-document-filter]") && document.querySelector("#swd-link-dialog [data-document-filter]").textContent.includes("program_steps.xlsx")'))
    js('document.querySelector("#swd-link-dialog [data-document-filter]").value=' + json.dumps(choice['documentId']) + ';document.querySelector("#swd-link-dialog [data-document-filter]").dispatchEvent(new Event("change"))')
    check('SWD文档筛选只显示所选文档来源', js('Array.from(document.querySelectorAll("#swd-link-dialog [data-choice]")).every(x=>x.querySelector("small").textContent.includes("program_steps.xlsx"))'))
    check('SWD导入后不自动确认Link', js('document.querySelector("#swd-link-dialog [data-generation]").value==="" && document.querySelector("#swd-link-dialog [data-apply]").disabled'))
    js('document.querySelector("#swd-link-dialog [data-kind=hex]").click()')
    check('SWD文件筛选不隐藏HEX来源', js('document.querySelectorAll("#swd-link-dialog [data-choice]").length>0'))
    js('document.querySelector("#swd-link-dialog [data-kind=swd]").click()')

    check('SWD候选标题为Pattern名称而非文档名称', js('document.querySelector('+choice_selector+').querySelector("strong").textContent==="PT104_demo"'))
    js('document.querySelector('+choice_selector+').click()')
    before_cancel = draft()
    js('document.querySelector("#swd-link-dialog [data-cancel]").click()')
    check('取消SWD来源选择不改变程序草稿', draft() == before_cancel)
    js('document.querySelector(' + selector + ').click()')
    until(lambda: js('!!document.querySelector("#swd-link-dialog[open]")'), 'reopen SWD source chooser')
    js('document.querySelector("#swd-link-dialog [data-kind=swd]").click()')
    js('document.querySelector("#swd-link-dialog [data-document-filter]").value=' + json.dumps(choice['documentId']) + ';document.querySelector("#swd-link-dialog [data-document-filter]").dispatchEvent(new Event("change"))')
    js('document.querySelector('+choice_selector+').click()')
    js('document.querySelector("#swd-link-dialog [data-edit-source]").click()')
    until(lambda: top('!!document.getElementById("wb-swd-frame").contentWindow.document.querySelector("#swd-pattern-dialog[open]")'), 'targeted SWD step editor')
    check('从程序卡片准确打开所选步骤编辑器', top('document.getElementById("wb-swd-frame").contentWindow.MCUSWD.state().pattern==="PT104_demo"'))
    top('document.getElementById("wb-swd-frame").contentWindow.document.getElementById("swd-pattern-dialog").close()')
    top('document.querySelector("#wb-pattern-return button").click()')
    until(lambda: js('!!document.querySelector("#swd-link-dialog[open]")'), 'return to source selection')
    js('document.querySelector("#swd-link-dialog [data-kind=swd]").click()')
    js('document.querySelector('+choice_selector+').click()')
    link_started = perf_counter()
    js('document.querySelector("#swd-link-dialog [data-apply]").click()')
    until(lambda: js('!document.querySelector("#swd-link-dialog[open]")'), 'SWD definition linked')
    link_elapsed = perf_counter() - link_started
    check(f'SWD确认关联响应 {link_elapsed:.2f}秒（小于3秒）', link_elapsed < 3)
    target = target_named('CP1', 'LSI_IWDG_Test', 'download')
    resource = target_name(target)
    selector = json.dumps('[data-resource-link="' + resource + '"]')
    binding = next(b for b in swd_targets()['links'] if b['target']['id'] == target['id'])
    check('SWD确认仅保存步骤关联，不先生成Pattern', 'swdSource' in binding and len(session.swd.data['generations']) == generations_before)
    check('程序资源卡显示自动生成状态', js('document.querySelector('+selector+').textContent.includes("生成程序时自动生成")'))
    rendered = session.swd.freeze(draft(), cps={target['cp']})
    check('生成时可取得SWD双平台波形且不创建历史产物', any(b['sourceKind'] == 'swd' and set(b['contents']) == {'chroma', 's100'} for b in rendered) and len(session.swd.data['generations']) == generations_before)
    session.swd.link({'targetId': target['id'], 'generationId': None})
    # A real non-download ADC card must use the same source-name adoption path.
    adc_candidates = [t for t in swd_targets()['targets']
                      if t['cp'] == 'CP3' and t['item'] == 'Adc_Test' and t['role'] != 'download']
    if not adc_candidates:
        raise AssertionError('找不到 CP3/Adc_Test 非下载 Pattern目标')
    adc_target = next((t for t in adc_candidates if any('ADC1' in name for name in sum(t['names'].values(), []))),
                      adc_candidates[0])
    adc_resource = target_name(adc_target)
    adc_selector = json.dumps('[data-resource-link="' + adc_resource + '"]')
    js('window.MCUProgram.navigateFlow(' + str(adc_target['row']) + ')')
    until(lambda: js('!!document.querySelector(' + adc_selector + ')'), 'ADC1 non-download Link card')
    before_adc_draft = draft()
    js('document.querySelector(' + adc_selector + ').click()')
    until(lambda: js('!!document.querySelector("#swd-link-dialog[open]")'), 'ADC1 source chooser')
    js('document.querySelector("#swd-link-dialog [data-kind=swd]").click()')
    adc_choice_selector = json.dumps('#swd-link-dialog [data-choice="' + choice['id'] + '"]')
    js('document.querySelector(' + adc_choice_selector + ').click()')
    js('document.querySelector("#swd-link-dialog [data-cancel]").click()')
    check('ADC1取消来源选择不改变草稿', draft() == before_adc_draft)
    js('document.querySelector(' + adc_selector + ').click()')
    until(lambda: js('!!document.querySelector("#swd-link-dialog[open]")'), 'ADC1 source confirm chooser')
    js('document.querySelector("#swd-link-dialog [data-kind=swd]").click()')
    js('document.querySelector(' + adc_choice_selector + ').click()')
    link_started = perf_counter()
    js('document.querySelector("#swd-link-dialog [data-apply]").click()')
    until(lambda: js('!document.querySelector("#swd-link-dialog[open]")'), 'ADC1 source Link saved')
    link_elapsed = perf_counter() - link_started
    check(f'ADC SWD确认关联响应 {link_elapsed:.2f}秒（小于3秒）', link_elapsed < 3)
    adc_draft = draft()
    adc_patterns = [entry for entry in adc_draft['patternLinks']
                    if entry['cp'] == adc_target['cp'] and entry['item'] == adc_target['item']]
    check('ADC1确认更新草稿 Pattern Link 来源名称', any(entry['pattern'] == 'PT104_demo' for entry in adc_patterns))
    candidate_raw = session.scan.candidate(adc_draft)[1]
    check('ADC1候选工作簿写入内部 Pattern Link JSON', any(entry['pattern'] == 'PT104_demo' for entry in read_pattern_links(candidate_raw)
                                                         if entry['cp'] == adc_target['cp'] and entry['item'] == adc_target['item']))
    # Exercise the real JS-to-Python save bridge with deterministic picker choices;
    # the platform's native dialog implementation itself is not automated here.
    from pathlib import Path
    from gui.workbook_editor import sha
    chosen = session.output / 'chosen folder' / '关联草稿.xlsx'
    chosen.parent.mkdir(parents=True, exist_ok=True)
    choices = [None, (str(chosen),)]
    calls = []
    native_dialog = app.window.create_file_dialog
    def pick_path(*args, **kwargs):
        calls.append((args, kwargs))
        return choices.pop(0)
    app.window.create_file_dialog = pick_path
    try:
        before_save = draft()
        js('document.getElementById("draft-file-btn").click()')
        until(lambda: len(calls) == 1 and js('!window.MCUProgram.status().busy'), 'native Save As cancellation')
        check('桌面另存取消保留Link草稿且不写目标文件', draft() == before_save and not chosen.exists())
        js('document.getElementById("draft-file-btn").click()')
        until(lambda: chosen.exists() and js('!window.MCUProgram.status().busy'), 'native Save As chosen path')
        saved_raw = chosen.read_bytes()
        check('桌面另存使用指定目录与文件名', len(calls) == 2 and calls[-1][1]['file_types'] == ('Excel workbook (*.xlsx)',))
        check('指定位置Excel保留内嵌Link名称映射', read_pattern_links(saved_raw) == adc_draft['patternLinks'])
        check('指定位置Excel指纹仍关联原SWD来源', any(b.get('swdSource', {}).get('pattern') == 'PT104_demo'
              for b in session.swd.data['links'][sha(saved_raw)]))
    finally:
        app.window.create_file_dialog = native_dialog
    adc_target = target_named('CP3', 'Adc_Test', adc_target['role'])
    adc_resource = target_name(adc_target)
    adc_selector = json.dumps('[data-resource-link="' + adc_resource + '"]')
    check('ADC1卡片显示来源 Pattern 名称', js('document.querySelector(' + adc_selector + ').textContent.includes("PT104_demo")'))
    js('document.querySelector(' + adc_selector + ').click()')
    until(lambda: js('!!document.querySelector("#swd-link-dialog[open]")'), 'ADC1 saved Link')
    js('(()=>{const original=window.confirm;window.confirm=()=>true;document.querySelector("#swd-link-dialog [data-unlink]").click();window.confirm=original;})()')
    until(lambda: js('!document.querySelector("#swd-link-dialog[open]")'), 'unlink ADC1')
    check('解除ADC1绑定后仍保留草稿来源名称', any(entry['pattern'] == 'PT104_demo' for entry in draft()['patternLinks']) and
         any(change['value'] == 'PT104_demo' for change in draft()['changes'] if change['sheet'] == 'FLOW'))
    # The reference workbook has no SCAN row. Add one only to this test session.
    book = openpyxl.load_workbook(BytesIO(session.editor.raw))
    sheet = book['FLOW']
    columns = {cell.value: cell.column for cell in sheet[1]}
    row = sheet.max_row + 1
    for name, value in {'CP NO': 'CP3', 'item': 'SCAN_Link_Acceptance', 'type': 'verify', 'mode': 'scan'}.items():
        sheet.cell(row, columns[name], value)
    raw = BytesIO(); book.save(raw); book.close()
    encoded = json.dumps(base64.b64encode(raw.getvalue()).decode())
    js('(()=>{const bytes=Uint8Array.from(atob(' + encoded + '),c=>c.charCodeAt(0));'
       'const original=window.confirm;window.confirm=()=>true;window.MCUProgram.loadFile(new File([bytes],"scan_link_acceptance.xlsx"));window.confirm=original;return true;})()')
    until(lambda: js('window.MCUProgram.status().name==="scan_link_acceptance.xlsx"'), 'SCAN workbook reload')
    scan_target = session.scan.targets()['targets'][0]
    js('window.MCUProgram.navigateFlow(' + str(scan_target['row']) + ')')
    until(lambda: js('!!document.querySelector("[data-scan-link]")'), 'program SCAN entry')
    js('document.querySelector("[data-scan-link]").click()')
    until(lambda: js('!!document.querySelector("#scan-link-dialog[open]")'), 'SCAN Link modal')
    check('程序页 SCAN 入口展示三件套关联', js('document.querySelector("#scan-link-dialog").textContent.includes("时序定义")'))
    js('document.querySelector("#scan-link-dialog [data-cancel]").click()')
