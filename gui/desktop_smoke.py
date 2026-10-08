"""Opt-in integration check inside this application's own WebView2 windows."""
import base64
import hashlib
import json
from pathlib import Path
import sys
import time
import traceback
from gui.paths import WORKSPACE_ROOT


def until(condition, label, timeout=90):
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        result = condition()
        if result:
            return result
        time.sleep(.2)
    raise AssertionError('超时：' + label)


def run_smoke(app, report_path):
    report = {'passed': False, 'frozen': bool(getattr(sys, 'frozen', False)), 'checks': []}
    master = WORKSPACE_ROOT / 'input/workbook/program_PCD009_V1.0.xlsx'
    before = hashlib.sha256(master.read_bytes()).hexdigest()
    main = app.window
    top_js = main.evaluate_js
    def js(script):
        return top_js('document.getElementById("wb-program-frame").contentWindow.eval(' + json.dumps(script) + ')')
    def swd_js(script):
        return top_js('document.getElementById("wb-swd-frame").contentWindow.eval(' + json.dumps(script) + ')')
    def check(label, condition):
        if not condition:
            raise AssertionError(label)
        report['checks'].append(label)
    def edit(address, value):
        # Child may own native focus: dispatch the complete edit lifecycle explicitly.
        js('(()=>{const e=document.querySelector(' + json.dumps('[data-edit="' + address + '"]') + ');'
           'if(!e)throw Error("cell missing");e.focus();e.dispatchEvent(new FocusEvent("focusin",{bubbles:true}));e.value=' + json.dumps(value) + ';'
           'e.dispatchEvent(new Event("input",{bubbles:true}));e.dispatchEvent(new FocusEvent("blur"));e.blur();})()')
    try:
        if app.error:
            raise AssertionError(app.error)
        until(lambda: top_js('!!window.MCUWorkbench'), 'workbench entry')
        check('工作台启动不依赖默认母本', app.service.session.editor is None)
        app.service.session.load(master.read_bytes(), master.name, master)
        top_js('window.MCUWorkbench.open("program")')
        until(lambda: js('!!document.querySelector("[data-edit=S4]") && !document.getElementById("item-detach").hidden'), 'main workbook and bridge')
        check('真实主窗口载入207行参考母本', app.service.session.editor.name == master.name)
        report['inputSha256'] = app.service.session.editor.sha256
        check('固定type列显示契约下拉，普通item列保持输入框', js("document.querySelector('[data-edit=E4]').tagName==='SELECT' && [...document.querySelector('[data-edit=E4]').options].some(o=>o.value==='shmoo') && document.querySelector('[data-edit=B4]').tagName==='INPUT'"))
        js("document.querySelector('[data-flow-drag=\"4\"]').dispatchEvent(new KeyboardEvent('keydown',{key:'ArrowDown',altKey:true,bubbles:true}))")
        check('FLOW排序手势改变真实行序', js("document.querySelector('[data-edit=B4]').value==='Iref1u' && document.querySelector('[data-edit=B5]').value==='Vref0p8_ABS'"))
        payload = js('JSON.parse(localStorage.getItem(' + json.dumps('mcu-gui-draft:' + app.service.session.editor.sha256) + '))')
        raw = app.service.session.editor.apply(payload['baseSha256'], payload['changes'], payload['appendRows'], payload['flowOrder'])
        from io import BytesIO
        import openpyxl
        ordered = openpyxl.load_workbook(BytesIO(raw), data_only=True)
        try:
            check('打包后端按界面顺序写回Excel', ordered['FLOW']['B4'].value == 'Iref1u' and ordered['FLOW']['B5'].value == 'Vref0p8_ABS')
        finally:
            ordered.close()
        _, compiled = app.service.session.compile(raw)
        cp = compiled.plan.cp_plans['CP1']
        items = {definition.identity: definition.item for definition in cp.definitions}
        execution = [items[entry.definition_id] for entry in cp.schedule]
        check('真实编译器执行计划采用调整后的FLOW顺序', execution.index('Iref1u') < execution.index('Vref0p8_ABS'))
        js("document.getElementById('undo-btn').click()")
        check('FLOW撤销恢复行序且清除顺序草稿', js("document.querySelector('[data-edit=B4]').value==='Vref0p8_ABS' && !document.getElementById('draft-count').textContent.includes('顺序已调整')"))
        js("document.querySelector('[data-recipe=F4]').click()")
        until(lambda: js("document.querySelectorAll('#recipe-table tbody tr').length===3"), 'recipe default modules')
        check('Recipe跳转逐模块表格，默认流程查看不产生草稿', js("!document.getElementById('recipe-page').hidden && document.getElementById('draft-count').textContent.startsWith('0 ') && document.getElementById('recipe-help-list').textContent.includes('SRAM_CAPTURE')"))
        js("(()=>{const e=document.querySelector('[aria-label=\"模块1参数\"]');e.value='WAIT=1MS';e.dispatchEvent(new Event('input',{bubbles:true}));document.getElementById('recipe-apply').click();})()")
        until(lambda: js("document.getElementById('draft-count').textContent.startsWith('1 ') && document.getElementById('recipe-error').textContent===''") , 'recipe apply')
        js("document.getElementById('recipe-return').click()")
        check('模块参数写回原Recipe坐标', js("document.querySelector('[data-recipe=F4]').textContent.includes('ENTER(WAIT=1MS)-TRIM-JUDGE')"))
        js("document.getElementById('undo-btn').click()")
        check('Recipe一次撤销恢复空白', js("document.getElementById('draft-count').textContent.startsWith('0 ') && document.querySelector('[data-recipe=F4]').textContent.includes('默认 / 未填写')"))
        js("document.querySelector('.app-menu summary').click()")
        check('文件菜单展开且桌面隐藏关闭服务', js("document.querySelector('.app-menu').open && document.getElementById('stop-btn').hidden"))
        js("document.getElementById('workspace-build').click()")
        check('生成工作区独占内容且关闭菜单', js("!document.getElementById('build-page').hidden && document.getElementById('editor-page').hidden && !document.querySelector('.app-menu').open"))
        js("document.getElementById('workspace-review').click()")
        check('检查记录独立显示', js("!document.getElementById('review-page').hidden && document.getElementById('build-page').hidden"))
        js("document.getElementById('workspace-edit').click()")
        js('document.querySelector(\'[data-select-row="4"]\').click()')
        until(lambda: js('document.getElementById("item-code").value.length>100'), 'initial item')
        original = js('document.getElementById("item-code").value')
        js('document.getElementById("item-detach").click()')
        child = until(lambda: app.bridge._preview, 'native child window')
        until(lambda: child.evaluate_js('document.getElementById("code")?.value') == original, 'child initial code')
        check('独立原生窗口显示同一代码，主页面收起代码区', child.uid != main.uid and js('document.getElementById("item-inline").hidden'))
        position = (main.x, main.y)
        child.move(80, 90)
        child.resize(960, 700)
        check('独立窗口可移动和缩放，主窗口位置不变', (main.x, main.y) == position and child.width == 960)
        edit('S4', '0.801')
        top_js('window.MCUWorkbench.open("home")')
        until(lambda: top_js('Math.round(document.querySelector(".wb-sidebar").getBoundingClientRect().width)') == 176, 'home navigation transition')
        check('返回首页恢复完整导航', top_js('Math.round(document.querySelector(".wb-sidebar").getBoundingClientRect().width)') == 176)
        top_js('window.MCUWorkbench.open("scan");window.MCUWorkbench.open("program")')
        until(lambda: top_js('Math.round(document.querySelector(".wb-sidebar").getBoundingClientRect().width)') == 64, 'tool navigation transition')
        check('跨工作台模块保留程序草稿并收窄导航', js("document.querySelector('[data-edit=S4]').value==='0.801'")
              and top_js('Math.round(document.querySelector(".wb-sidebar").getBoundingClientRect().width)') == 64)
        js("document.getElementById('workspace-build').click();document.getElementById('workspace-review').click();document.getElementById('workspace-edit').click()")
        report['editSwitch'] = js("({value:document.querySelector('[data-edit=S4]').value,undoDisabled:document.getElementById('undo-btn').disabled,draft:document.getElementById('draft-count').textContent,hasBlurFix:document.scripts[document.scripts.length-1].textContent.includes('model.activeEditBefore=null;updateButtons()')})")
        check('真实窗口切换保留输入与撤销记录', js("document.querySelector('[data-edit=S4]').value === '0.801' && !document.getElementById('undo-btn').disabled"))
        until(lambda: child.evaluate_js('document.getElementById("code").value') == '', 'clear stale code', 5)
        check('编辑立即清空独立窗口旧代码', True)
        code = until(lambda: js('document.getElementById("item-code").value.includes("0.801") && document.getElementById("item-code").value'), 'edited item')
        until(lambda: child.evaluate_js('document.getElementById("code").value') == code, 'edited child code')
        check('真实FLOW参数0.801经生成器同步到独立窗口', code != original)
        child.destroy()
        until(lambda: app.bridge._preview is None and not js('document.getElementById("item-inline").hidden'), 'close restores inline')
        edit('S4', '0.802')
        code = until(lambda: js('document.getElementById("item-code").value.includes("0.802") && document.getElementById("item-code").value'), 'second edit')
        js('document.getElementById("item-detach").click()')
        reopened = until(lambda: app.bridge._preview, 'reopen')
        until(lambda: reopened.evaluate_js('document.getElementById("code")?.value') == code, 'reopen latest code')
        check('关闭后恢复内联，再次打开显示最新0.802代码', reopened.uid != child.uid)
        js('document.getElementById("item-detach").click()')
        time.sleep(.5)
        check('重复打开复用已有窗口', app.bridge._preview is reopened and len(app.webview.windows) == 2)
        # Exercise both platform plans from the packaged modules, without writing artifacts.
        from gui.generation import available_targets, preview_plan
        cache = app.service.session.item_cache
        plans = {}
        for platform in ('chroma', 's100'):
            target = next(t['id'] for t in available_targets(cache['workbook'], cache['compilation'])
                          if t['id'].startswith(platform + ':') and ':CP1:lab:' in t['id'])
            plan = preview_plan(cache['workbook'], cache['compilation'], target)
            plans[platform] = len(plan.artifacts)
        report['planArtifacts'] = plans
        check('打包模块可构造双平台完整程序计划', all(plans.values()))
        check('预览未签发检查或创建生成任务', app.service.session.checked is None and app.service.session.job is None)
        # Keep smoke fixtures and all SWD products beside the requested report, never in user workspace outputs.
        from gui.swd_workspace import SwdWorkspace
        swd_output = report_path.parent / 'desktop-smoke-swd'
        app.service.session.output = swd_output
        app.service.swd = SwdWorkspace(app.service.session)
        app.service.session.swd = app.service.swd
        from io import BytesIO
        import openpyxl
        fixture_book = openpyxl.Workbook()
        fixture_sheet = fixture_book.active
        fixture_sheet.title = 'SWD'
        fixture_sheet.append(['pattern', 'addr', 'data', 'step', 'reset value', 'type', 'continue', 'description'])
        fixture_sheet.append(['SMOKE_PATTERN', 'WAIT', 1, '/', '/', 'ms', 'N', 'desktop smoke fixture'])
        fixture_sheet.merge_cells('I2:J2')
        fixture_sheet['I2'] = 'merged source note'
        fixture = BytesIO()
        fixture_book.save(fixture)
        fixture_raw = fixture.getvalue()
        fixture_path = report_path.parent / 'desktop-smoke-swd-fixture.xlsx'
        fixture_path.write_bytes(fixture_raw)
        report['swdFixtureSha256'] = hashlib.sha256(fixture_raw).hexdigest()
        report['swdOutputRoot'] = str(swd_output)
        top_js('window.MCUWorkbench.open("swd")')
        until(lambda: top_js('!document.querySelector("[data-page-view=swd]").hidden && !!document.getElementById("wb-swd-frame").src'), 'SWD workbench navigation')
        until(lambda: swd_js('!!window.MCUSWD && !!document.getElementById("swd-excel-picker")'), 'SWD page ready')
        check('工作台导航载入持久 SWD 页面', top_js('!document.querySelector("[data-page-view=swd]").hidden && document.getElementById("wb-swd-frame").src.endsWith("/swd")'))
        encoded_fixture = base64.b64encode(fixture_raw).decode('ascii')
        swd_js('(()=>{const token=JSON.parse(document.getElementById("connection").textContent).token;'
               'const raw=Uint8Array.from(atob(' + json.dumps(encoded_fixture) + '),c=>c.charCodeAt(0));'
               'fetch("/api/swd/import",{method:"POST",headers:{"Content-Type":"application/octet-stream","X-File-Name":encodeURIComponent("desktop-smoke-swd-fixture.xlsx"),"X-GUI-Session":token},body:raw})'
               '.then(r=>{if(!r.ok)throw Error("HTTP "+r.status);return window.MCUSWD.refresh();})'
               '.catch(e=>window.__desktopSmokeSwdError=String(e));})()')
        until(lambda: len(app.service.swd.state()['documents']) == 1, 'SWD fixture import')
        until(lambda: swd_js('!!document.querySelector("[data-row-id=\\"2\\"]") && !document.getElementById("swd-generate").disabled'), 'SWD fixture render')
        check('SWD 页面通过同源 fetch 导入隔离 xlsx 夹具', swd_js('!window.__desktopSmokeSwdError && window.MCUSWD.state().documents[0].name==="desktop-smoke-swd-fixture.xlsx"'))
        swd_js('(()=>{const input=document.querySelector("[data-row-id=\\"2\\"][data-field=\\"data\\"]");input.value="2";input.dispatchEvent(new Event("input",{bubbles:true}));document.getElementById("swd-add-row").click();document.getElementById("swd-save").click();})()')
        until(lambda: len(app.service.swd.state()['documents'][0]['rows']) == 2, 'SWD step save')
        check('SWD 步骤编辑、新增并保存到隔离工作区', app.service.swd.state()['documents'][0]['rows'][0]['data'] == '2')
        swd_js('(()=>{const values={"swd-pin-swc":"PA11","swd-pin-swd":"PA13","swd-pin-reset":"NRST","swd-pin-trig":"TRIGPIN"};Object.entries(values).forEach(([id,value])=>{const input=document.getElementById(id);input.value=value;input.dispatchEvent(new Event("input",{bubbles:true}));});document.getElementById("swd-save-config").click();})()')
        until(lambda: app.service.swd.state()['config']['pins']['swc'] == 'PA11', 'SWD pin configuration')
        check('SWD 页面保存四个管脚配置', app.service.swd.state()['config']['pins'] == {'swc': 'PA11', 'swd': 'PA13', 'reset': 'NRST', 'trig': 'TRIGPIN'})
        swd_js('document.getElementById("swd-generate").click()')
        until(lambda: len(app.service.swd.state()['generations']) == 1, 'SWD dual-platform generation')
        first_generation = app.service.swd.state()['generations'][0]
        check('SWD 页面生成 Chroma 与 S100 两个平台产物', {f['platform'] for f in first_generation['files']} == {'chroma', 's100'} and not first_generation['stale'])
        until(lambda: swd_js('document.getElementById("swd-generation-list").textContent.includes("当前")'), 'current SWD generation marker')
        swd_js('(()=>{const input=document.getElementById("swd-pin-swc");input.value="PA12";input.dispatchEvent(new Event("input",{bubbles:true}));document.getElementById("swd-save-config").click();})()')
        until(lambda: app.service.swd.state()['generations'][0]['stale'], 'stale SWD generation marker')
        until(lambda: swd_js('document.getElementById("swd-generation-list").textContent.includes("已过期")'), 'stale marker render')
        swd_js('document.getElementById("swd-generate").click()')
        until(lambda: len(app.service.swd.state()['generations']) == 2 and not app.service.swd.state()['generations'][-1]['stale'], 'current regenerated SWD output')
        check('SWD 当前与过期生成记录同时可见', swd_js('document.getElementById("swd-generation-list").textContent.includes("当前") && document.getElementById("swd-generation-list").textContent.includes("已过期")'))
        top_js('window.MCUWorkbench.open("home");window.MCUWorkbench.open("swd")')
        until(lambda: swd_js('window.MCUSWD.state().documents.length===1 && window.MCUSWD.state().generations.length===2'), 'SWD page state survives navigation')
        check('切换页面后保留 SWD 文档、步骤与生成修订', swd_js('window.MCUSWD.state().documents[0].rows.length===2 && window.MCUSWD.state().generations.length===2'))
        swd_js('document.querySelector("[data-cache-id=py32f420]").click()')
        until(lambda: swd_js('!window.MCUSWD.state().busy && window.MCUSWD.state().materials.materials.some(x=>x.kind==="svd_cache")'), 'bundled PY32F420 cache')
        sources = [source for register in app.service.swd.state()['materials']['registers'] for source in register['sources']]
        check('实际安装包加载 PY32F420 缓存与位字段', len(sources) == 525 and sum(len(s['fields']) for s in sources) == 3717)
        swd_js('document.getElementById("swd-create-pattern").click();document.getElementById("swd-new-pattern-name").value="SMOKE_CACHED_PATTERN";document.getElementById("swd-create-confirm").click()')
        until(lambda: swd_js('document.getElementById("swd-pattern-dialog").open && !window.MCUSWD.state().busy'), 'Pattern authoring dialog')
        check('创建Pattern只新增未保存草稿', app.service.swd.state()['documents'][0]['patterns'] == ['SMOKE_PATTERN'] and swd_js('window.MCUSWD.state().dirty'))
        swd_js('document.getElementById("swd-pick-register").click()')
        until(lambda: swd_js('document.getElementById("swd-register-dialog").open'), 'register browser')
        swd_js('(()=>{const row=[...document.querySelectorAll("[data-register-select]")].find(x=>x.querySelector("strong").textContent==="TIM2.CR1");row.click();document.getElementById("swd-register-apply").click();document.querySelector("[data-bit-index=\\"31\\"]").click();})()')
        check('缓存寄存器与32位位图接入 Pattern 草稿', swd_js('document.getElementById("swd-step-data").value==="0x80000000" && [...document.querySelectorAll("[data-bit-index]")].some(b=>b.title.includes("CEN"))'))
        swd_js('document.getElementById("swd-draft-apply").click()')
        until(lambda: swd_js('!document.getElementById("swd-pattern-dialog").open && !window.MCUSWD.state().busy'), 'save Pattern draft')
        check('保存步骤未重写原始Excel', Path(app.service.swd.data['documents'][0]['source']).read_bytes() == fixture_raw)
        swd_js('document.getElementById("swd-generate").click()')
        until(lambda: any(g['pattern'] == 'SMOKE_CACHED_PATTERN' for g in app.service.swd.state()['generations']), 'generate authored Pattern')
        check('安装包新建 Pattern 保存并生成双平台', {f['platform'] for f in app.service.swd.state()['generations'][-1]['files']} == {'chroma', 's100'})
        check('SWD整页无滚动且表格为八列', swd_js('document.documentElement.scrollHeight<=innerHeight && document.body.scrollHeight<=innerHeight && document.querySelectorAll(".swd-table th").length===8'))
        link_target = next(iter(app.service.swd.targets()['targets']), None)
        if link_target is None:
            raise AssertionError('参考母本没有可关联的 SWD Pattern 引用')
        top_js('window.MCUWorkbench.open("program")')
        js('document.querySelector(' + json.dumps('[data-select-row="' + str(link_target['row']) + '"]') + ').click()')
        until(lambda: js('!!document.querySelector("[data-resource-link]")'), 'program SWD Link entry')
        check('程序测试项直接显示资源卡片', js('!!document.querySelector("[data-resource-link]") && !document.querySelector("[data-pattern-link]")'))
        js('document.querySelector("[data-resource-link]").click()')
        until(lambda: js('!!document.querySelector("#swd-link-dialog[open] [data-generation]")'), 'SWD Link dialog')
        check('程序 SWD Link 弹窗取得当前生成修订', js('document.querySelector("#swd-link-dialog [data-generation]").textContent.includes("SMOKE_CACHED_PATTERN")'))
        js('document.querySelector("#swd-link-dialog [data-cancel]").click()')
        from gui.pattern_link_smoke import check_pattern_links
        check_pattern_links(app, js, check)
        check('参考母本字节未改变', hashlib.sha256(master.read_bytes()).hexdigest() == before)
        report['passed'] = True
    except Exception:
        report['error'] = traceback.format_exc()
    finally:
        report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
        app.bridge._busy = False
        main.destroy()
