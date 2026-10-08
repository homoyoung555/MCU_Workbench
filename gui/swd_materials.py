"""首期 SWD 资料导入：保存原文、可复核解析结果与待确认差异。"""

from __future__ import annotations

import hashlib
import io
import json
import re
import xml.etree.ElementTree as ET
from copy import deepcopy
from pathlib import Path


class MaterialLibrary:
    """资料库的唯一公开接口；导入绝不修改已有资料或自动裁决冲突。"""

    _KINDS = {"svd": "svd", "svd_cache": "svd_cache", "excel": "excel", "h": "h", "c": "c", "header": "h", "source": "c"}

    def __init__(self, root):
        self.root = Path(root)
        self.root.mkdir(parents=True, exist_ok=True)
        self.sources = self.root / "sources"
        self.sources.mkdir(exist_ok=True)
        self.path = self.root / "materials.json"
        self._data = self._load()

    def import_material(self, raw, name, kind=None):
        """保存一份原始资料，并提取能够直接由该资料证明的事实。"""
        if kind is None:
            suffix = Path(name).suffix.lower()
            if suffix == ".xls":
                raise ValueError("请将旧 .xls 另存为 .xlsx 后导入")
            kind = {".svd": "svd", ".xlsx": "excel", ".h": "h", ".c": "c"}.get(suffix)
        kind = self._KINDS.get(str(kind).lower())
        if not kind:
            raise ValueError("kind 必须是 svd、svd_cache、excel、h 或 c")
        if not isinstance(name, str) or not name:
            raise ValueError("name 必须是非空文件名")
        raw = self._bytes(raw)
        sha256 = hashlib.sha256(raw).hexdigest()
        existing = next((item for item in self._data["materials"] if item["sha256"] == sha256), None)
        if existing is not None:
            return deepcopy(existing)
        analysis, registers = self._parse(raw, kind)
        safe_name = self._safe_name(name)
        stored_as = Path("sources") / f"{sha256[:16]}_{safe_name}"
        target = self.root / stored_as
        if not target.exists():
            target.write_bytes(raw)
        for register in registers:
            register["materialId"] = sha256
            register["id"] = f"{sha256}:{register['id']}"
        material = {
            "id": sha256,
            "kind": kind,
            "originalName": name,
            "storedAs": stored_as.as_posix(),
            "sha256": sha256,
            "analysis": analysis,
        }
        self._data["materials"].append(material)
        self._data["registerSources"].extend(registers)
        self._save()
        return deepcopy(material)

    def state(self):
        """返回资料、按地址并列的事实，以及尚未裁决的冲突。"""
        registers = {}
        for source in self._data["registerSources"]:
            registers.setdefault(source["address"], []).append(deepcopy(source))
        items = [{"address": address, "sources": sources} for address, sources in sorted(registers.items())]
        return {
            "version": hashlib.sha256(json.dumps(self._data, ensure_ascii=False, sort_keys=True).encode()).hexdigest(),
            "materials": deepcopy(self._data["materials"]),
            "registers": items,
            "conflicts": self._conflicts(items),
            "confirmations": deepcopy(self._data["confirmations"]),
        }

    def confirm(self, conflict_id, source_id=None, note=""):
        """记录用户选择或备注；保留所有原始来源，不改写其中任何一项。"""
        conflict = next((item for item in self.state()["conflicts"] if item["id"] == conflict_id), None)
        if conflict is None:
            raise ValueError("未知冲突")
        if source_id is not None and source_id not in conflict["sourceIds"]:
            raise ValueError("选择的来源不属于该冲突")
        self._data["confirmations"][conflict_id] = {"sourceId": source_id, "note": str(note)}
        self._save()
        return self.state()

    def _load(self):
        if not self.path.exists():
            return {"version": 1, "materials": [], "registerSources": [], "confirmations": {}}
        data = json.loads(self.path.read_text(encoding="utf-8"))
        data.setdefault("materials", [])
        data.setdefault("registerSources", [])
        data.setdefault("confirmations", {})
        return data

    def _save(self):
        temporary = self.path.with_suffix(".tmp")
        temporary.write_text(json.dumps(self._data, ensure_ascii=False, indent=2), encoding="utf-8")
        temporary.replace(self.path)

    @staticmethod
    def _bytes(raw):
        if isinstance(raw, bytes):
            return raw
        if isinstance(raw, str):
            return raw.encode("utf-8")
        try:
            return MaterialLibrary._bytes(raw.read())
        except AttributeError as error:
            raise TypeError("raw 必须是 bytes、str 或返回 bytes 的读取对象") from error

    @staticmethod
    def _safe_name(name):
        name = name.replace("\\", "/").rsplit("/", 1)[-1]
        name = re.sub(r"[^A-Za-z0-9._-]+", "_", name).strip("._")
        return name or "material"

    def _parse(self, raw, kind):
        if kind == "svd":
            return self._parse_svd(raw)
        if kind == "svd_cache":
            return self._parse_svd_cache(raw)
        if kind == "excel":
            return self._parse_excel(raw)
        return self._parse_c(raw, kind)

    @staticmethod
    def _number(value):
        if value is None:
            return None
        text = str(value).strip().replace("_", "")
        match = re.search(r"0x([0-9a-fA-F]+)", text)
        if match:
            return int(match.group(1), 16)
        try:
            return int(text, 10)
        except ValueError:
            return None

    @staticmethod
    def _tag(element):
        return element.tag.rsplit("}", 1)[-1]

    @classmethod
    def _child(cls, element, name):
        return next((item for item in element if cls._tag(item) == name), None)

    @classmethod
    def _text(cls, element, name):
        child = cls._child(element, name)
        return child.text.strip() if child is not None and child.text else ""

    def _parse_svd(self, raw):
        root = ET.fromstring(raw)
        container = next((item for item in root.iter() if self._tag(item) == "peripherals"), None)
        raw_peripherals = {}
        for element in [] if container is None else list(container):
            if self._tag(element) != "peripheral":
                continue
            name = self._text(element, "name")
            if name:
                raw_peripherals[name] = element

        resolved = {}

        def registers_for(name, visiting=()):
            if name in resolved:
                return deepcopy(resolved[name])
            if name in visiting or name not in raw_peripherals:
                return []
            element = raw_peripherals[name]
            parent = element.get("derivedFrom", "").split(".")[-1]
            registers = {item["name"]: item for item in registers_for(parent, visiting + (name,))}
            registers_element = self._child(element, "registers")
            if registers_element is not None:
                for register in registers_element:
                    if self._tag(register) != "register":
                        continue
                    register_name = self._text(register, "name") or register.get("derivedFrom", "").split(".")[-1]
                    inherited = deepcopy(registers.get(register.get("derivedFrom", "").split(".")[-1], {}))
                    offset = self._number(self._text(register, "addressOffset"))
                    if offset is None:
                        offset = inherited.get("offset")
                    if not register_name or offset is None:
                        continue
                    registers[register_name] = {
                        "name": register_name,
                        "offset": offset,
                        "resetValue": self._text(register, "resetValue") or inherited.get("resetValue", ""),
                        "access": self._text(register, "access") or inherited.get("access", ""),
                    }
            resolved[name] = list(registers.values())
            return deepcopy(resolved[name])

        records = []
        peripheral_count = 0
        for name, element in raw_peripherals.items():
            base = self._number(self._text(element, "baseAddress"))
            if base is None:
                continue
            peripheral_count += 1
            for register in registers_for(name):
                records.append(self._register_record("svd", name, base, register))
        return {"device": self._text(root, "name"), "peripheralCount": peripheral_count}, records

    @staticmethod
    def _cache_address(value):
        return value if type(value) is int and 0 <= value <= 0xFFFFFFFF else None

    @staticmethod
    def _cache_text(value):
        return value if isinstance(value, str) else ""

    @classmethod
    def _normalize_reset(cls, value):
        if value is None:
            return ""
        if type(value) is int and value >= 0:
            return hex(value)
        if not isinstance(value, str):
            return ""
        text = value.strip()
        match = re.fullmatch(r"(?i)(\d+)'([bodh])([0-9a-f_xz]+)", text)
        if not match:
            return text
        width, base, digits = match.groups()
        valid_digits = {
            "b": "01_",
            "o": "01234567_",
            "d": "0123456789_",
            "h": "0123456789abcdef_",
        }[base.lower()]
        if any(char.lower() not in valid_digits for char in digits):
            return text
        if "x" in digits.lower() or "z" in digits.lower():
            return text
        try:
            return hex(int(digits.replace("_", ""), {"b": 2, "o": 8, "d": 10, "h": 16}[base.lower()]))
        except ValueError:
            return text

    @classmethod
    def _cache_fields(cls, raw_fields):
        if not isinstance(raw_fields, list):
            return []
        fields = []
        for field in raw_fields:
            if not isinstance(field, dict):
                continue
            name = cls._cache_text(field.get("name"))
            bits = cls._cache_text(field.get("bitRange"))
            offset = field.get("bitOffset")
            width = field.get("bitWidth")
            bit_range = re.fullmatch(r"\[(\d+):(\d+)\]", bits)
            if not name or not bit_range or type(offset) is not int or type(width) is not int or offset < 0 or width <= 0 or offset + width > 32 or (int(bit_range.group(1)), int(bit_range.group(2))) != (offset + width - 1, offset):
                continue
            fields.append({
                "name": name,
                "bits": bits,
                "bitOffset": offset,
                "bitWidth": width,
                "access": cls._cache_text(field.get("access")),
                "description": cls._cache_text(field.get("description")),
                "resetValue": cls._normalize_reset(field.get("resetValue")),
            })
        return fields

    def _parse_svd_cache(self, raw):
        try:
            cache = json.loads(raw.decode("utf-8"))
            data = cache["data"]
            device = data["device"]
            device_name = self._cache_text(device.get("name")) if isinstance(device, dict) else ""
            peripherals = data["peripherals"]
        except (UnicodeDecodeError, json.JSONDecodeError, KeyError, TypeError) as error:
            raise ValueError("无效 SVD 缓存 JSON") from error
        if not device_name.upper().startswith("PY32F420"):
            raise ValueError("SVD 缓存设备必须是 PY32F420")
        if not isinstance(peripherals, list):
            raise ValueError("SVD 缓存 peripherals 必须是列表")

        records = []
        peripheral_count = register_count = field_count = 0
        for peripheral in peripherals:
            if not isinstance(peripheral, dict):
                continue
            name = self._cache_text(peripheral.get("name"))
            base = self._cache_address(peripheral.get("baseAddress"))
            registers = peripheral.get("registers")
            if not name or base is None or not isinstance(registers, list):
                continue
            peripheral_count += 1
            for register in registers:
                if not isinstance(register, dict):
                    continue
                register_name = self._cache_text(register.get("name"))
                address = self._cache_address(register.get("address"))
                if not register_name or address is None or address < base:
                    continue
                fields = self._cache_fields(register.get("fields"))
                records.append(self._register_record("svd_cache", name, base, {
                    "name": register_name,
                    "offset": address - base,
                    "description": self._cache_text(register.get("description")),
                    "resetValue": self._normalize_reset(register.get("resetValue")),
                    "access": self._cache_text(register.get("access")),
                    "fields": fields,
                }))
                register_count += 1
                field_count += len(fields)
        return {"device": device_name, "peripheralCount": peripheral_count, "registerCount": register_count, "fieldCount": field_count}, records

    def _parse_excel(self, raw):
        from openpyxl import load_workbook

        book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        records, sheets = {}, []
        required = {"regname", "baseaddress", "offsetaddress"}
        for sheet in book.worksheets:
            rows = sheet.iter_rows(values_only=True)
            headers = next(rows, ())
            columns = {self._header(value): index for index, value in enumerate(headers) if value is not None}
            if not required.issubset(columns):
                continue
            sheets.append(sheet.title)
            previous = {key: None for key in required}
            for row in rows:
                values = {key: row[index] if index < len(row) else None for key, index in columns.items()}
                for key in required:
                    value = values.get(key)
                    if value not in (None, ""):
                        previous[key] = value
                name = self._value_or_previous(values, previous, "regname")
                bases = self._base_entries(self._value_or_previous(values, previous, "baseaddress"), sheet.title)
                offset = self._number(self._value_or_previous(values, previous, "offsetaddress"))
                if not name or not bases or offset is None:
                    continue
                field_name = values.get("fieldname") or values.get("field")
                bits = values.get("bits") or values.get("bitrange") or ""
                reset = values.get("resetvalue") if values.get("resetvalue") not in (None, "") else ""
                for base, peripheral in bases:
                    key = (peripheral, str(name), base, offset)
                    record = records.setdefault(key, self._register_record("excel", peripheral, base, {"name": str(name), "offset": offset, "resetValue": ""}))
                    if field_name:
                        field = {"name": str(field_name), "bits": str(bits), "resetValue": str(reset)}
                        if field not in record["fields"]:
                            record["fields"].append(field)
        result = list(records.values())
        return {"sheets": sheets, "registerCount": len(result), "resetValueScope": "field"}, result

    @staticmethod
    def _header(value):
        return re.sub(r"[\s_-]+", "", str(value).lower())

    @staticmethod
    def _value_or_previous(values, previous, key):
        value = values.get(key)
        return value if value not in (None, "") else previous[key]

    @staticmethod
    def _addresses(value):
        if value is None:
            return []
        return [int(item.replace(" ", "").replace("_", ""), 16) for item in re.findall(r"0x([0-9a-fA-F](?:[0-9a-fA-F_ ]*[0-9a-fA-F])?)", str(value), re.I)]

    @classmethod
    def _base_entries(cls, value, fallback):
        if value is None:
            return []
        pattern = r"(?:(?P<label>[A-Za-z_]\w*)\s*[:：]\s*)?(?P<address>0x[0-9a-fA-F](?:[0-9a-fA-F_ ]*[0-9a-fA-F])?)"
        return [(int(match.group("address").replace(" ", "").replace("_", ""), 16), match.group("label") or fallback) for match in re.finditer(pattern, str(value), re.I)]

    def _parse_c(self, raw, kind):
        text = raw.decode("utf-8", errors="replace")
        macros, unresolved, direct = [], [], {}
        for match in re.finditer(r"^\s*#\s*define\s+([A-Za-z_]\w*)\s+(.+?)\s*$", text, re.M):
            name, expression = match.groups()
            expression = expression.split("//", 1)[0].strip()
            numeric = self._number(expression) if re.fullmatch(r"[()\s0-9a-fA-FxX_uUlL]+", expression) else None
            macros.append({"name": name, "expression": expression, "value": numeric})
            dependent = re.findall(r"\b[A-Z][A-Z0-9_]*\b", expression)
            dependent = [item for item in dependent if item not in {"U", "L", "UL"}]
            addresses = self._addresses(expression)
            pointer_address = self._pointer_literal_address(expression)
            if pointer_address is not None and not dependent:
                direct[name] = {"address": pointer_address, "resetValue": ""}
            elif dependent or self._looks_like_pointer(expression) or len(addresses) > 1:
                unresolved.append({"name": name, "expression": expression})
        for macro in macros:
            base = re.sub(r"_(?:RESET|RESET_VALUE)$", "", macro["name"])
            if base in direct and macro["value"] is not None:
                direct[base]["resetValue"] = hex(macro["value"])
        records = [self._register_record(kind, "", value["address"], {"name": name, "offset": 0, "resetValue": value["resetValue"]}) for name, value in direct.items()]
        return {"macros": macros, "unresolvedMacros": unresolved, "parseNote": "仅提取直接地址与字面宏；不把 C/H 当作完整寄存器解析。"}, records

    @staticmethod
    def _strip_outer_parentheses(text):
        text = text.strip()
        while text.startswith("(") and text.endswith(")"):
            depth = 0
            encloses_all = True
            for index, char in enumerate(text):
                if char == "(":
                    depth += 1
                elif char == ")":
                    depth -= 1
                    if depth == 0 and index != len(text) - 1:
                        encloses_all = False
                        break
            if not encloses_all or depth != 0:
                break
            text = text[1:-1].strip()
        return text

    @classmethod
    def _pointer_literal_address(cls, expression):
        text = cls._strip_outer_parentheses(expression)
        if text.startswith("*"):
            text = cls._strip_outer_parentheses(text[1:])
        cast = re.match(r"^\(\s*(?:volatile\s+)?(?:[A-Za-z_]\w*\s+)*[A-Za-z_]\w*\s*\*\s*\)", text)
        if not cast:
            return None
        literal = cls._strip_outer_parentheses(text[cast.end():])
        if not re.fullmatch(r"(?:0x[0-9a-fA-F]+|[0-9]+)[uUlL]*", literal):
            return None
        return cls._number(literal)

    @staticmethod
    def _looks_like_pointer(expression):
        return bool(re.search(r"\(\s*(?:volatile\s+)?(?:[A-Za-z_]\w*\s+)*[A-Za-z_]\w*\s*\*\s*\)", expression))

    @staticmethod
    def _register_record(kind, peripheral, base, register):
        address = base + register["offset"]
        source_id = f"{kind}:{peripheral}:{register['name']}:{address:08x}:{register.get('resetValue', '')}"
        return {
            "id": source_id,
            "kind": kind,
            "peripheral": peripheral,
            "name": register["name"],
            "baseAddress": f"0x{base:08x}",
            "offsetAddress": f"0x{register['offset']:x}",
            "address": f"0x{address:08x}",
            "resetValue": register.get("resetValue", ""),
            "access": register.get("access", ""),
            "description": register.get("description", ""),
            "fields": deepcopy(register.get("fields", [])),
        }

    @staticmethod
    def _identity(source):
        peripheral = source["peripheral"].upper()
        name = source["name"].upper()
        if peripheral and name.startswith(f"{peripheral}_"):
            name = name[len(peripheral) + 1:]
        return f"{peripheral}.{name}" if peripheral else name

    @classmethod
    def _conflicts(cls, registers):
        conflicts = []
        by_name = {}
        by_field = {}
        for item in registers:
            sources = item["sources"]
            for source in sources:
                if source["name"]:
                    by_name.setdefault(cls._identity(source), []).append(source)
                for field in source.get("fields", []):
                    if field["resetValue"]:
                        key = (item["address"], cls._identity(source), field["name"].upper(), field["bits"])
                        by_field.setdefault(key, []).append((field["resetValue"], source["id"]))
            names = {cls._identity(source) for source in sources if source["name"]}
            resets = {source["resetValue"] for source in sources if source["resetValue"]}
            for kind, values in (("addressIdentity", names), ("resetValue", resets)):
                if len(values) > 1:
                    relevant = [source["id"] for source in sources if cls._identity(source) in values] if kind == "addressIdentity" else [source["id"] for source in sources if source["resetValue"] in values]
                    conflicts.append({"id": f"{item['address']}:{kind}", "address": item["address"], "kind": kind, "values": sorted(values), "sourceIds": relevant})
        for identity, sources in by_name.items():
            addresses = {source["address"] for source in sources}
            if len(addresses) > 1:
                conflicts.append({"id": f"{identity}:registerAddress", "name": identity, "kind": "registerAddress", "values": sorted(addresses), "sourceIds": [source["id"] for source in sources]})
        for (address, identity, name, bits), values in by_field.items():
            reset_values = {value for value, _ in values}
            if len(reset_values) > 1:
                conflicts.append({"id": f"{address}:{identity}:{name}:{bits}:fieldReset", "address": address, "kind": "fieldReset", "field": {"name": name, "bits": bits}, "values": sorted(reset_values), "sourceIds": [source_id for _, source_id in values]})
        return conflicts
