from io import BytesIO
import unittest
from zipfile import ZipFile

import openpyxl

from gui.pattern_editor import pattern_items
from gui.workbook_editor import WorkbookEditor


class PatternEditorTests(unittest.TestCase):
    def setUp(self):
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = 'pattern'
        sheet.append(['item', 'addr', 'description', 'pattern'])
        sheet.append(['First', '0100010', '2,trg=0', 'trim mode'])
        sheet.append([None, '0100010', '2,trg=1', None])
        sheet.append([None, None, None, None])
        sheet.append(['Second', '0000000', '3,dac[3:0]=0000*', 'trim dac'])
        sheet.merge_cells('A2:A3')
        sheet.merge_cells('D2:D3')
        sheet['B2'].number_format = '@'
        other = book.create_sheet('notes')
        other['A1'] = '=1+2'
        buffer = BytesIO()
        book.save(buffer)
        book.close()
        self.editor = WorkbookEditor(buffer.getvalue(), 'fixture.xlsx')

    def test_ordered_repeat_address_roundtrip_preserves_other_zip_parts_and_merges(self):
        edits = [{'item': 'First', 'groups': [{'role': 'trim mode', 'steps': [
            {'sourceRow': 3, 'addr': '0100010', 'description': '2,trg=1'},
            {'sourceRow': 2, 'addr': '0100010', 'description': '2,trg=0'},
            {'sourceRow': None, 'addr': '0100010', 'description': '2,trg=1'}]}]}]
        raw = self.editor.apply(self.editor.sha256, [], pattern_edits=edits)
        reopened = WorkbookEditor(raw, 'result.xlsx')
        items = pattern_items(reopened)
        self.assertEqual([i['item'] for i in items], ['First', 'Second'])
        self.assertEqual([s['description'] for s in items[0]['groups'][0]['steps']], ['2,trg=1', '2,trg=0', '2,trg=1'])
        book = openpyxl.load_workbook(BytesIO(raw))
        self.assertEqual({str(a) for a in book['pattern'].merged_cells.ranges}, {'A2:A4', 'D2:D4'})
        self.assertEqual(book['pattern']['B3'].number_format, '@')
        self.assertEqual(book['pattern']['C6'].value, '3,dac[3:0]=0000*')
        book.close()
        with ZipFile(BytesIO(raw)) as new, ZipFile(BytesIO(self.editor.raw)) as old:
            self.assertEqual(new.namelist(), old.namelist())
            self.assertEqual([n for n in new.namelist() if new.read(n) != old.read(n)], [self.editor.sheet_paths['pattern']])

    def test_new_item_and_scalar_edits_coexist_and_stale_or_cross_item_source_rejected(self):
        edits = [{'item': 'Added', 'groups': [{'role': 'verify mode', 'steps': [
            {'sourceRow': None, 'addr': '0000001', 'description': '1,enable=1'}]}]}]
        raw = self.editor.apply(self.editor.sha256, [{'sheet': 'pattern', 'cell': 'C2', 'old': '2,trg=0', 'value': '2,trg=1'}], pattern_edits=edits)
        self.assertEqual(pattern_items(WorkbookEditor(raw, 'new.xlsx'))[-1]['item'], 'Added')
        with self.assertRaisesRegex(ValueError, '版本'):
            self.editor.apply('wrong', [], pattern_edits=edits)
        edits[0]['groups'][0]['steps'][0]['sourceRow'] = 2
        with self.assertRaisesRegex(ValueError, '来源'):
            self.editor.apply(self.editor.sha256, [], pattern_edits=edits)

    def test_incomplete_values_and_duplicate_sources_block_materialization(self):
        entry = {'item': 'First', 'groups': [{'role': 'trim mode', 'steps': [
            {'sourceRow': 2, 'addr': '0100010', 'description': '2,trg='}]}]}
        with self.assertRaisesRegex(ValueError, '尚未填写'):
            self.editor.apply(self.editor.sha256, [], pattern_edits=[entry])
        entry['groups'][0]['steps'][0]['description'] = '2,trg=0'
        entry['groups'][0]['steps'].append(dict(entry['groups'][0]['steps'][0]))
        with self.assertRaisesRegex(ValueError, '来源'):
            self.editor.apply(self.editor.sha256, [], pattern_edits=[entry])

    def test_partial_width_and_formula_dependencies_cannot_be_exported(self):
        entry = {'item': 'Added', 'groups': [{'role': 'trim mode', 'steps': [
            {'sourceRow': None, 'addr': '0100010', 'description': '3,mode[3:0]=01'}]}]}
        with self.assertRaisesRegex(ValueError, '位数'):
            self.editor.apply(self.editor.sha256, [], pattern_edits=[entry])
        book = openpyxl.load_workbook(BytesIO(self.editor.raw))
        book['notes']['A1'] = "=pattern!B2"
        stream = BytesIO()
        book.save(stream)
        book.close()
        editor = WorkbookEditor(stream.getvalue(), 'locked.xlsx')
        entry['groups'][0]['steps'][0]['description'] = '3,mode[3:0]=0001'
        with self.assertRaisesRegex(ValueError, '只读'):
            editor.apply(editor.sha256, [], pattern_edits=[entry])


if __name__ == '__main__':
    unittest.main()
