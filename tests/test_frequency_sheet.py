import unittest

from openpyxl import Workbook

import analog_sweep as A
from program_gen.core import clock_force
from program_gen.core.encoding import calc_limit


def _new_freq_workbook():
    wb = Workbook()
    freq = wb.active
    freq.title = "FREQ"
    freq.append(["FREQ", "source", "mul", "div", "step", "down_limit", "up_limit"])
    freq.append(["Clock_check_Test", "HSI_8M", "/", 128, 1, "0.5%", "0.5%"])
    freq.append([None, "HSI_16M", "/", 128, 2, "-0.625", "+0.625"])
    freq.append(["Hse_Test", "HSE_32M", "/", 128, 3, "0.5%", "0.5%"])
    freq.append(["PLL_Test", "HSI_8M", 6, 128, 4, "0.5%", "0.5%"])
    freq.append([None, "HSE_32M", 4, 128, 5, "0.5%", "0.5%"])

    iomux = wb.create_sheet("IOMUX")
    iomux.append(["pin", "additional"])
    iomux.append([None, "other"])
    iomux.append(["PD00-OSC_IN", "OSC_IN"])
    iomux.append(["PC14-OSC32_IN", "OSC32_IN"])
    return wb


def _legacy_freq_workbook():
    wb = Workbook()
    ws = wb.active
    ws.title = "FREQ"
    ws.append(["FREQ", "mode", "force", "target", "down_limit", "up_limit", "unit"])
    ws.append(["Clock_check_Test", "HSI_8M_Check", "/", 62.5, "0.5%", "0.5%", "KHz"])
    return wb


class FrequencySheetTests(unittest.TestCase):
    def test_new_schema_derives_mode_force_target_and_unit(self):
        rows = A.freq_cases(_new_freq_workbook())

        self.assertEqual(
            rows,
            [
                {
                    "freq": "Clock_check_Test",
                    "source": "HSI_8M",
                    "mul": 1.0,
                    "div": 128.0,
                    "step": "1",
                    "mode": "HSI_8M_Check",
                    "pattern": "HSI_8M",
                    "force": "/",
                    "target": 62.5,
                    "down_limit": "0.5%",
                    "up_limit": "0.5%",
                    "unit": "KHz",
                    "_row": 2,
                },
                {
                    "freq": "Clock_check_Test",
                    "source": "HSI_16M",
                    "mul": 1.0,
                    "div": 128.0,
                    "step": "2",
                    "mode": "HSI_16M_Check",
                    "pattern": "HSI_16M",
                    "force": "/",
                    "target": 125.0,
                    "down_limit": "-0.625",
                    "up_limit": "+0.625",
                    "unit": "KHz",
                    "_row": 3,
                },
                {
                    "freq": "Hse_Test",
                    "source": "HSE_32M",
                    "mul": 1.0,
                    "div": 128.0,
                    "step": "3",
                    "mode": "HSE_Check",
                    "pattern": "HSE",
                    "force": "PD0:32MHz",
                    "target": 250.0,
                    "down_limit": "0.5%",
                    "up_limit": "0.5%",
                    "unit": "KHz",
                    "_row": 4,
                },
                {
                    "freq": "PLL_Test",
                    "source": "HSI_8M",
                    "mul": 6.0,
                    "div": 128.0,
                    "step": "4",
                    "mode": "PLL_HSI8_PLLX6",
                    "pattern": "PLL_HSI8",
                    "force": "/",
                    "target": 375.0,
                    "down_limit": "0.5%",
                    "up_limit": "0.5%",
                    "unit": "KHz",
                    "_row": 5,
                },
                {
                    "freq": "PLL_Test",
                    "source": "HSE_32M",
                    "mul": 4.0,
                    "div": 128.0,
                    "step": "5",
                    "mode": "PLL_HSE32_PLLX4",
                    "pattern": "PLL_HSE32",
                    "force": "PD0:32MHz",
                    "target": 1.0,
                    "down_limit": "0.5%",
                    "up_limit": "0.5%",
                    "unit": "MHz",
                    "_row": 6,
                },
            ],
        )

    def test_group_filter_and_limit_logic_use_derived_target(self):
        rows = A.freq_cases(_new_freq_workbook(), "Clock_check_Test")

        self.assertEqual(len(rows), 2)
        self.assertEqual(calc_limit(rows[0]["target"], rows[0]["down_limit"], rows[0]["up_limit"], True), 62.188)
        self.assertEqual(calc_limit(rows[0]["target"], rows[0]["down_limit"], rows[0]["up_limit"], False), 62.812)
        self.assertEqual(calc_limit(rows[1]["target"], rows[1]["down_limit"], rows[1]["up_limit"], True), 124.375)
        self.assertEqual(calc_limit(rows[1]["target"], rows[1]["down_limit"], rows[1]["up_limit"], False), 125.625)

    def test_legacy_schema_remains_supported(self):
        rows = A.freq_cases(_legacy_freq_workbook())

        self.assertEqual(rows[0]["mode"], "HSI_8M_Check")
        self.assertEqual(rows[0]["pattern"], "HSI_8M")
        self.assertEqual(rows[0]["target"], 62.5)
        self.assertEqual(rows[0]["unit"], "KHz")

    def test_missing_iomux_clock_role_is_rejected(self):
        wb = _new_freq_workbook()
        wb.remove(wb["IOMUX"])

        with self.assertRaisesRegex(ValueError, "IOMUX.*OSC_IN"):
            A.freq_cases(wb)

    def test_iomux_clock_role_ignores_common_separators_and_case(self):
        aliases = ("oscin", "OSC-IN", "Osc In", "osc32in", "OSC32-IN", "Osc32 In")
        for alias in aliases:
            with self.subTest(alias=alias):
                wb = _new_freq_workbook()
                is_lse = "32" in alias
                role = "OSC32_IN" if is_lse else "OSC_IN"
                if is_lse:
                    wb["FREQ"].append(["Lse_Test", "LSE_32K", "/", 1, 6, "0.5%", "0.5%"])
                for row in range(3, wb["IOMUX"].max_row + 1):
                    if wb["IOMUX"].cell(row, 2).value == role:
                        wb["IOMUX"].cell(row, 2).value = alias
                rows = A.freq_cases(wb)
                expected_pin = "PC14" if is_lse else "PD0"
                expected_source = "LSE_32K" if is_lse else "HSE_32M"
                expected = next(row for row in rows if row["source"] == expected_source)
                self.assertTrue(expected["force"].startswith(expected_pin + ":"))

    def test_iomux_clock_role_does_not_confuse_output_with_input(self):
        wb = _new_freq_workbook()
        wb["IOMUX"]["B3"] = "OSC_OUT"

        with self.assertRaisesRegex(ValueError, "OSC_IN"):
            A.freq_cases(wb)

    def test_duplicate_source_and_multiplier_in_one_group_is_rejected(self):
        wb = _new_freq_workbook()
        wb["FREQ"].append(["PLL_Test", "HSI_8M", 6, 128, 7, "0.5%", "0.5%"])

        with self.assertRaisesRegex(ValueError, "PLL_Test.*HSI_8M.*重复"):
            A.freq_cases(wb)

    def test_missing_sheet_keeps_legacy_flow_compatible(self):
        self.assertEqual(A.freq_cases(Workbook()), [])

    def test_s100_external_clock_plan_accepts_derived_force_cell(self):
        self.assertEqual(
            clock_force.s100_cell_plan("PD0:32MHz"),
            {"PD0": {"freq_hz": 32000000.0, "mode": "tmu"}},
        )
        self.assertEqual(
            clock_force.s100_cell_plan("PC14:32768Hz"),
            {"PC14": {"freq_hz": 32768.0, "mode": "pattern"}},
        )


if __name__ == "__main__":
    unittest.main()
