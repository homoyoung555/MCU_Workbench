"""The new-item wizard prepares one atomic draft, without compiling incomplete rows."""
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from gui.server import Session
from gui.workbook_editor import WorkbookEditor


class NewFlowItemTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.folder = tempfile.TemporaryDirectory()
        cls.session = Session(output=Path(cls.folder.name))

    @classmethod
    def tearDownClass(cls):
        cls.folder.cleanup()

    def payload(self, **extra):
        return {'baseSha256': self.session.editor.sha256, 'changes': [],
                'values': {'CP NO': 'CP1', 'item': 'Wizard_New', 'type': 'trim', 'mode': 'fbist'}, **extra}

    def filled(self, row=3, **values):
        sheet = self.session.editor.by_name['FLOW']
        original = {c['label']: '' if cell['value'] is None else str(cell['value'])
                    for c, cell in zip(sheet['columns'], sheet['rows'][row - 1]) if c['label']}
        return {**original, 'item': 'Wizard_New', 'OTP': 'no', 'UID': 'no', 'PUMP': 'no',
                '_storageMode': 'none', '_codeMode': 'full',
                '_enterPattern': 'enter_vref0p8_tc', '_dacPattern': 'set_vref0p8_tc_dac', **values}

    def test_resource_edit_preserves_groups_merges_and_unrelated_draft(self):
        session = self.session
        for name,row,field,value,count in [('ADC',3,'force','0.31',22),('OPA',2,'target','1.66',1),
                                         ('DAC',2,'code','0x003',1),('PVD',5,'limit','31mV',2)]:
            args = dict(sheet=name,editRow=row)
            options = session.new_flow_item(self.payload(action='resource-options',**args))
            self.assertEqual(len(options['templates']),1)
            template = options['templates'][0]
            self.assertEqual(len(template['rows']),count)
            records = [dict(r) for r in template['records']]
            records[template['rows'].index(row)][field] = value
            result = session.new_flow_item(self.payload(action='resource-prepare',**args,
                sourceRows=template['rows'],records=records))
            self.assertTrue(result['ready'],result.get('issues'))
            d = result['draft']
            self.assertEqual(d['appendRows'],[])
            self.assertEqual(len(d['changes']),1)
            reopened = WorkbookEditor(session.editor.apply(session.editor.sha256,d['changes']),'edit.xlsx')
            self.assertEqual(reopened.by_name[name]['maxRow'],session.editor.by_name[name]['maxRow'])
            # A second edit starts from the current draft and preserves the first edit's original value.
            again = session.new_flow_item(self.payload(action='resource-options',**args,**d))
            self.assertEqual(again['templates'][0]['records'][template['rows'].index(row)][field],value)
        session.editor.ensure_source_current()

    def test_resource_edit_rejects_identity_readonly_and_partial_groups(self):
        session = self.session
        options = session.new_flow_item(self.payload(action='resource-options',sheet='ADC',editRow=3))
        template = options['templates'][0]
        for field,value,error in [('ADC','ADC9','关联标识'),('Chanel','21','关联标识'),('step','RP','只读')]:
            records = [dict(r) for r in template['records']]
            records[1][field] = value
            with self.assertRaisesRegex(ValueError,error):
                session.new_flow_item(self.payload(action='resource-prepare',sheet='ADC',editRow=3,
                    sourceRows=template['rows'],records=records))
        with self.assertRaises(ValueError):
            session.new_flow_item(self.payload(action='resource-options',sheet='ADC',editRow=1))

    def test_resource_edit_supports_new_draft_rows(self):
        session = self.session
        template = session.new_flow_item(self.payload(action='resource-options',sheet='DAC'))['templates'][0]
        new = session.new_flow_item(self.payload(action='resource-prepare',sheet='DAC',
            sourceRows=template['rows'],records=template['records']))
        d = new['draft']
        options = session.new_flow_item(self.payload(action='resource-options',sheet='DAC',editRow=new['row'],**d))
        edited = options['templates'][0]
        records = [dict(edited['records'][0],code='0x003')]
        result = session.new_flow_item(self.payload(action='resource-prepare',sheet='DAC',editRow=new['row'],
            sourceRows=edited['rows'],records=records,**d))
        self.assertTrue(result['ready'],result.get('issues'))
        self.assertEqual(result['draft']['appendRows'],d['appendRows'])
        change = next(c for c in result['draft']['changes'] if c['cell']==f"B{new['row']}")
        self.assertIsNone(change['old'])
        self.assertEqual(change['value'],'0x003')

    def test_resource_units_append_and_preserve_flow(self):
        session = self.session
        draft = {}
        for name, index, count in [('OPA',0,1),('DAC',0,1),('PVD',1,2)]:
            options = session.new_flow_item(self.payload(action='resource-options',sheet=name,**draft))
            template = options['templates'][index]
            records = [dict(r) for r in template['records']]
            if name=='PVD':
                for r in records: r['case']='99'
            result = session.new_flow_item(self.payload(action='resource-prepare',sheet=name,
                sourceRows=template['rows'],records=records,**draft))
            self.assertTrue(result['ready'], result.get('issues'))
            self.assertEqual(result['count'],count)
            draft = result['draft']
        raw = session.editor.apply(session.editor.sha256,draft['changes'],draft['appendRows'],draft['flowOrder'])
        reopened = WorkbookEditor(raw,'resources.xlsx')
        for name, count in [('OPA',1),('DAC',1),('PVD',2)]:
            self.assertEqual(reopened.by_name[name]['maxRow'],session.editor.by_name[name]['maxRow']+count)
        self.assertEqual(reopened.by_name['FLOW']['rows'],session.editor.by_name['FLOW']['rows'])
        result = session.new_flow_item(self.payload(action='prepare',before=3,values=self.filled(),**draft))
        self.assertTrue(result['ready'],result.get('issues'))
        self.assertEqual(sum(op['sheet']=='FLOW' for op in result['draft']['appendRows']),1)
        session.editor.ensure_source_current()

    def test_adc_resource_uses_draft_iomux_and_preserves_merges(self):
        import re
        from openpyxl.utils import get_column_letter
        session = self.session
        template = session.new_flow_item(self.payload(action='resource-options',sheet='ADC'))['templates'][0]
        records = [dict(r,ADC='ADC3') for r in template['records']]
        iomux = session.editor.by_name['IOMUX']
        changes = []
        for r,row in enumerate(iomux['rows'],1):
            for c,cell in enumerate(row,1):
                value = cell['value']
                tokens = re.findall(r'ADC(?:12|1)_IN(\d+)',str(value))
                if tokens:
                    changes.append({'sheet':'IOMUX','cell':f'{get_column_letter(c)}{r}',
                                    'old':value,'value':str(value)+''.join('/ADC3_IN'+t for t in tokens)})
        self.assertEqual(sum(c['value'].count('/ADC3_IN') for c in changes),22)
        result = session.new_flow_item(self.payload(action='resource-prepare',sheet='ADC',
            sourceRows=template['rows'],records=records,changes=changes))
        self.assertTrue(result['ready'],result.get('issues'))
        draft = result['draft']
        raw = session.editor.apply(session.editor.sha256,draft['changes'],draft['appendRows'])
        reopened = WorkbookEditor(raw,'adc.xlsx').by_name['ADC']
        self.assertEqual(reopened['maxRow'],session.editor.by_name['ADC']['maxRow']+22)
        self.assertEqual([r[0]['value'] for r in reopened['rows'][-22:]],['ADC3']*22)
        self.assertEqual(reopened['rows'][:45],session.editor.by_name['ADC']['rows'])

    def test_adc_guided_iomux_is_atomic_and_preserves_existing_tokens(self):
        import re
        session = self.session
        options = session.new_flow_item(self.payload(action='resource-options',sheet='ADC'))
        template = options['templates'][0]
        rows = [None]*22
        for pin in options['iomuxPins']:
            for channel in re.findall(r'ADC(?:12|1)_IN(\d+)',pin['analog']):
                rows[int(channel)] = pin['row']
        records = [dict(r,ADC='ADC3') for r in template['records']]
        args = dict(action='resource-prepare',sheet='ADC',sourceRows=template['rows'],records=records)
        with self.assertRaisesRegex(ValueError,'22'):
            session.new_flow_item(self.payload(**args,iomuxRows=rows[:-1]))
        with self.assertRaisesRegex(ValueError,'引脚'):
            session.new_flow_item(self.payload(**args,iomuxRows=[1]*22))
        choice = next(p for p in options['iomuxPins'] if p['row']==rows[0])
        original = session.editor.by_name['IOMUX']['rows'][choice['row']-1][1]['value']
        existing = {'sheet':'IOMUX','cell':choice['cell'],'old':original,
                    'value':choice['analog']+'/CUSTOM_SIGNAL/ADC3_IN0'}
        conflict = list(rows)
        conflict[0] = next(p['row'] for p in options['iomuxPins'] if p['pin']!=choice['pin'])
        with self.assertRaisesRegex(ValueError,'已映射'):
            session.new_flow_item(self.payload(**args,iomuxRows=conflict,changes=[existing]))
        result = session.new_flow_item(self.payload(**args,iomuxRows=rows,changes=[existing]))
        self.assertIn('CUSTOM_SIGNAL',next(c['value'] for c in result['draft']['changes'] if c['sheet']=='IOMUX' and c['cell']==choice['cell']))
        self.assertTrue(result['ready'],result.get('issues'))
        d = result['draft']
        for change in d['changes']:
            if change['sheet']=='IOMUX':
                self.assertTrue(change['value'].startswith(str(change['old'])))
        raw = session.editor.apply(session.editor.sha256,d['changes'],d['appendRows'])
        _,compiled = session.compile(raw)
        self.assertTrue(compiled.is_valid)
        session.editor.ensure_source_current()

    def test_resource_invalid_units_never_produce_drafts(self):
        for name in ('ADC','OPA','DAC','PVD'):
            options = self.session.new_flow_item(self.payload(action='resource-options',sheet=name))
            template = options['templates'][0]
            records = [dict(r) for r in template['records']]
            if name=='ADC':
                for r in records: r['ADC']='ADC99'
            elif name in ('OPA','DAC'): records[0]['mode']=''
            # PVD's unchanged identity already exists and must be rejected.
            with self.assertRaises(ValueError):
                self.session.new_flow_item(self.payload(action='resource-prepare',sheet=name,
                    sourceRows=template['rows'],records=records))

    def test_existing_phase_structure_roundtrips_and_compiles(self):
        session = self.session
        base = session.editor.by_name['FLOW']
        order = list(range(1,base['maxRow']+1))
        swapped = list(order)
        swapped[196:198] = [198,197]
        for current in (swapped,[r for r in order if r!=198]):
            result = session.new_flow_item(self.payload(action='validate-edit',flowOrder=current))
            self.assertTrue(result['ready'],result.get('issues'))
            raw = session.editor.apply(session.editor.sha256,[],[],current)
            reopened = WorkbookEditor(raw,'structure.xlsx')
            self.assertEqual(reopened.by_name['FLOW']['rows'][196][3]['value'],'ADC2' if current==swapped else 'ADC1')
        new = base['maxRow']+1
        added = [{'sheet':'FLOW','sourceRow':87,'copy':True}]
        extended = [*order[:87],new,*order[87:]]
        changes = [{'sheet':'FLOW','cell':f'D{new}','old':base['rows'][86][3]['value'],'value':'Extra_Run'}]
        result = session.new_flow_item(self.payload(action='validate-edit',changes=changes,appendRows=added,flowOrder=extended))
        self.assertTrue(result['ready'],result.get('issues'))
        raw = session.editor.apply(session.editor.sha256,changes,added,extended)
        reopened = WorkbookEditor(raw,'extended.xlsx')
        self.assertEqual(reopened.by_name['FLOW']['rows'][87][3]['value'],'Extra_Run')
        self.assertEqual(reopened.by_name['FLOW']['rows'][88][1]['value'],base['rows'][87][1]['value'])
        # A phase cannot be inserted beyond an unrelated item and split its group.
        split = [*order[:89],new,*order[89:]]
        with self.assertRaisesRegex(ValueError,'拆分'):
            session.editor.apply(session.editor.sha256,changes,added,split)

    def test_existing_phase_groups_use_native_boundaries_and_validate_atomic_edits(self):
        session = self.session
        for selected, expected in ((197,[197,198]),(198,[197,198]),(86,[86,87]),(87,[86,87]),(199,[])):
            result = session.new_flow_item(self.payload(action='describe',editRow=selected,values=self.filled(selected)))
            self.assertEqual(result['editRows'],expected)
        sheet = session.editor.by_name['FLOW']
        from openpyxl.utils import get_column_letter
        def change(row, label, value):
            column = next(i for i,c in enumerate(sheet['columns']) if c['label']==label)
            return {'sheet':'FLOW','cell':get_column_letter(column+1)+str(row),
                    'old':sheet['rows'][row-1][column]['value'],'value':value}
        edits = [change(r,'Hbin','12') for r in (86,87)] + [change(87,'target','LL')]
        result = session.new_flow_item(self.payload(action='validate-edit',changes=edits))
        self.assertTrue(result['ready'],result.get('issues'))
        raw = session.editor.apply(session.editor.sha256,edits)
        reopened = WorkbookEditor(raw,'edited-group.xlsx')
        self.assertEqual(reopened.by_name['FLOW']['maxRow'],sheet['maxRow'])
        bad = session.new_flow_item(self.payload(action='validate-edit',changes=[change(87,'target','HHH')]))
        self.assertFalse(bad['ready']);self.assertNotIn('draft',bad)
        # Validation is read-only, including on failure.
        self.assertEqual(session.editor.sha256,WorkbookEditor(session.editor.apply(session.editor.sha256,[]),'base.xlsx').sha256)

    def test_phase_groups_compile_contiguously_and_reject_invalid_stages(self):
        adc = {'CP NO':'CP3','item':'Wizard_Adc_Group','type':'verify','mode':'swd','FV/FI':'ADC1',
               'Recipe':'SWD_CAPTURE-ADC_JUDGE','Vcc':'VCC_TYP','Hbin':'12','Sbin':'12','_group':'adc'}
        digital = {**adc, 'mode':'keil','Recipe':'RUN_SEQUENCE-JUDGE','step':'PA0:HL-5ms',
                   'target':'HH','meas pin':'PA13+PA14'}
        sram = {**adc,'mode':'keil','Recipe':'RUN_SEQUENCE-SRAM_CAPTURE(0X3000,22x16@2)-ADC_JUDGE','step':'PA0:HL-5ms-222-15'}
        firmware = {**adc,'item':'Wizard_Firmware','_group':'firmware','mode':'hex','FV/FI':'/',
                    'Recipe':'','sub':'download','step':'firmware_download'}
        run = {**firmware,'mode':'keil','Recipe':'RUN_SEQUENCE-JUDGE','sub':'Run_A','step':'PA0:HL-5ms',
               'target':'HH','meas pin':'PA13+PA14','UNIT':'/'}
        for first, second in ((adc,{**adc,'FV/FI':'ADC2'}),(digital,{**digital,'FV/FI':'ADC2'}),
                              (sram,{**sram,'FV/FI':'ADC2'}),(firmware,run)):
            with self.subTest(mode=first['mode'],recipe=first['Recipe']):
                result = self.session.new_flow_item(self.payload(action='prepare',before=199,values=first,phases=[first,second]))
                self.assertTrue(result['ready'],result.get('issues'))
                self.assertEqual(result['rowCount'],2)
                d = result['draft']
                self.assertEqual(len(d['appendRows']),2)
                raw = self.session.editor.apply(self.session.editor.sha256,d['changes'],d['appendRows'],d['flowOrder'],d['trimStorageMappings'],d['verifyStorageMappings'])
                sheet = WorkbookEditor(raw,'phases.xlsx').by_name['FLOW']
                for row in sheet['rows'][result['row']-1:result['row']+1]:
                    record = {c['label']:v['value'] for c,v in zip(sheet['columns'],row)}
                    self.assertEqual(record['item'],first['item'])
                    self.assertEqual(str(record['Hbin']),'12')
        for first, stages in ((adc,[adc,adc]),(adc,[adc,run]),(firmware,[firmware,{**run,'target':'HHH'}])):
            result = self.session.new_flow_item(self.payload(action='prepare',before=199,values=first,phases=stages))
            self.assertFalse(result['ready']); self.assertNotIn('draft',result)

    def test_opa_dac_keil_templates_use_current_sheet_steps(self):
        session = self.session
        from openpyxl.utils import get_column_letter
        for source_row, sheet_name in ((199,'OPA'),(200,'DAC')):
            with self.subTest(sheet=sheet_name):
                changes=[]
                for name, row_number, label, value in [('FLOW',source_row,'mode','keil'),('FLOW',source_row,'step','X')]:
                    sheet=session.editor.by_name[name]
                    col=next(i for i,c in enumerate(sheet['columns']) if c['label']==label)
                    changes.append({'sheet':name,'cell':get_column_letter(col+1)+str(row_number),'old':sheet['rows'][row_number-1][col]['value'],'value':value})
                sheet=session.editor.by_name[sheet_name]
                col=next(i for i,c in enumerate(sheet['columns']) if c['label']=='step')
                for i,row in enumerate(sheet['rows'][1:],2):
                    if row[col]['editable']:
                        changes.append({'sheet':sheet_name,'cell':get_column_letter(col+1)+str(i),'old':row[col]['value'],'value':'1'})
                values=self.filled(source_row,item='Wizard_Keil_'+sheet_name,mode='keil',step='HL-X-LH')
                result=session.new_flow_item(self.payload(action='prepare',before=199,values=values,changes=changes))
                self.assertTrue(result['ready'],result.get('issues'))
                values['step']='HL-1'
                result=session.new_flow_item(self.payload(action='prepare',before=199,values=values,changes=changes))
                self.assertFalse(result['ready']);self.assertNotIn('draft',result)

    def test_analog_sheet_items_roundtrip_and_compile(self):
        session = self.session
        cases = [self.filled(row, item=name) for row,name in ((199,'Wizard_Opa'),(200,'Wizard_Dac'),(201,'Wizard_Pvd'))]
        cases.append({'CP NO':'CP3','item':'Wizard_Adc','type':'verify','mode':'swd','FV/FI':'ADC1',
                      'Recipe':'SWD_CAPTURE-ADC_JUDGE','Vcc':'VCC_TYP','Hbin':'12','Sbin':'12'})
        for values in cases:
            with self.subTest(item=values['item']):
                result = session.new_flow_item(self.payload(action='prepare', before=199, values=values))
                self.assertTrue(result['ready'],result.get('issues'))
                self.assertEqual(result['resourceSheet'],result['kind'].split('_')[0])
                d = result['draft']
                raw = session.editor.apply(session.editor.sha256,d['changes'],d['appendRows'],d['flowOrder'],d['trimStorageMappings'],d['verifyStorageMappings'])
                reopened = WorkbookEditor(raw,'analog.xlsx')
                row = reopened.by_name['FLOW']['rows'][result['row']-1]
                record = {c['label']:v['value'] for c,v in zip(reopened.by_name['FLOW']['columns'],row)}
                self.assertEqual(record['item'],values['item'])
                if result['kind']=='ADC_TEST': self.assertEqual(record['sub'],'ADC1')
                _, compiled = session.compile(raw)
                self.assertTrue(compiled.is_valid,[d.message for d in compiled.diagnostics.errors])
        bad = self.filled(200,item='Wizard_Bad_Dac',target='4096.5')
        result = session.new_flow_item(self.payload(action='prepare',before=199,values=bad))
        self.assertFalse(result['ready']); self.assertNotIn('draft',result)
        missing = {**cases[-1], 'FV/FI':'ADC99'}
        result = session.new_flow_item(self.payload(action='prepare',before=199,values=missing))
        self.assertFalse(result['ready']); self.assertNotIn('draft',result)

    def test_specialized_rows_roundtrip_through_native_compiler(self):
        session = self.session
        for original_row, cp, item, before in ((2, 'CP1', 'os_extra', 3),
                                               (194, 'CP1', 'Clock_check_Test', 3),
                                               (9, 'CP1', 'Pdr_Trim', 10),
                                               (43, 'CP1', 'Pdr_Verify', 44)):
            with self.subTest(item=item):
                values = self.filled(original_row, **{'CP NO': cp, 'item': item})
                extra = {}
                if item.startswith('Pdr_'):
                    extra['flowOrder'] = [r for r in range(1, session.editor.by_name['FLOW']['maxRow'] + 1) if r != original_row]
                payload = self.payload(action='prepare', before=before, values=values, **extra)
                result = session.new_flow_item(payload)
                self.assertTrue(result['ready'], result.get('issues'))
                draft = result['draft']
                self.assertEqual(draft['trimStorageMappings'], session.editor.trim_storage_mappings)
                raw = session.editor.apply(session.editor.sha256, draft['changes'], draft['appendRows'], draft['flowOrder'],
                    draft['trimStorageMappings'], draft['verifyStorageMappings'], draft['patternEdits'])
                reopened = WorkbookEditor(raw, 'specialized.xlsx')
                row = reopened.by_name['FLOW']['rows'][result['row'] - 1]
                record = {c['label']: cell['value'] for c, cell in zip(reopened.by_name['FLOW']['columns'], row)}
                self.assertEqual(record['item'], item)
                self.assertEqual(record['target'], '/')
                self.assertEqual(record['Recipe'] or '', 'FREQ_TEST' if item == 'Clock_check_Test' else '')
                _, compiled = session.compile(raw)
                self.assertTrue(compiled.is_valid, [d.message for d in compiled.diagnostics.errors])

    def test_specialized_incomplete_options_and_missing_resources(self):
        session = self.session
        values = self.filled(194, item='Missing_Freq_Group')
        with patch.object(session, 'compile', side_effect=AssertionError('incomplete options must remain local')):
            view = session.new_flow_item(self.payload(action='options', values=values))
        self.assertTrue(view['specialized'])
        self.assertTrue(view['resources'])
        result = session.new_flow_item(self.payload(action='prepare', before=3, values=values))
        self.assertFalse(result['ready'])
        self.assertNotIn('draft', result)
        self.assertTrue(result['issues'])
        result = session.new_flow_item(self.payload(action='prepare', before=3, values=self.filled(2, item='os_extra', Vcc='')))
        self.assertFalse(result['ready'])
        self.assertIn('Vcc', [i['field'] for i in result['issues']])

    def test_same_name_trim_and_verify_can_be_added_in_one_draft(self):
        session=self.session
        trim=self.filled(item='Wizard_Pair')
        first=session.new_flow_item(self.payload(action='prepare',before=4,values=trim))
        self.assertTrue(first['ready'],first.get('issues'))
        verify=self.filled(4, type='verify',item='Wizard_Pair',_enterPattern='enter_vref0p8_abs_verify')
        options=session.new_flow_item(self.payload(action='options',values=verify,**first['draft']))
        verify['_enterPattern']=options['patterns']['enter'][0]['name']
        second=session.new_flow_item(self.payload(action='prepare',before=5,values=verify,**first['draft']))
        self.assertTrue(second['ready'],second.get('issues'))
        draft=second['draft']
        raw=session.editor.apply(session.editor.sha256,draft['changes'],draft['appendRows'],draft['flowOrder'],draft['trimStorageMappings'],draft['verifyStorageMappings'])
        workbook,compiled=session.compile(raw)
        self.assertTrue(compiled.is_valid,[d.message for d in compiled.diagnostics.errors])
        rows=workbook.dataframe('FLOW');pair=rows[rows['item']=='Wizard_Pair']
        self.assertEqual(list(pair['type']),['trim','verify'])
        for values in (trim,verify):
            with self.assertRaisesRegex(ValueError,'已存在'):
                session.new_flow_item(self.payload(action='prepare',before=6,values=values,**draft))
        session.editor.ensure_source_current()

    def test_prepared_trim_roundtrips_and_preserves_existing_rows(self):
        session = self.session
        payload = self.payload(action='prepare', before=4, values=self.filled())
        with patch.object(session, 'compile', side_effect=AssertionError('prepare is local')):
            result = session.new_flow_item(payload)
        self.assertTrue(result['ready'], result.get('issues'))
        self.assertTrue(result['pendingValidation'])
        draft = result['draft']
        raw = session.editor.apply(session.editor.sha256, draft['changes'], draft['appendRows'], draft['flowOrder'],
                                   draft['trimStorageMappings'], draft['verifyStorageMappings'])
        reopened = WorkbookEditor(raw, 'new-item.xlsx')
        flow = reopened.by_name['FLOW']
        names = [r[1]['value'] for r in flow['rows']]
        self.assertEqual(names[3], 'Wizard_New')
        self.assertEqual(names[:3] + names[4:], [r[1]['value'] for r in session.editor.by_name['FLOW']['rows']])
        _, compiled = session.compile(raw)
        self.assertTrue(compiled.is_valid, [d.message for d in compiled.diagnostics.errors])
        self.assertEqual(compiled.plan.catalogs.trim_specs[2].dac.pattern_name, 'set_vref0p8_tc_dac')

    def test_verify_uses_selected_pattern_and_clears_old_trim_parameters(self):
        session = self.session
        values = self.filled(4, **{'type': 'verify', 'item': 'Wizard_Verify', 'DAC bit': '4',
                                   'Bisection': 'yes', 'step': '0,1', 'LINE': 'no',
                                   '_enterPattern': 'enter_vref0p8_abs_verify'})
        view = session.new_flow_item(self.payload(action='options', values=values))
        values['_enterPattern'] = view['patterns']['enter'][0]['name']
        result = session.new_flow_item(self.payload(action='prepare', before=5, values=values))
        self.assertTrue(result['ready'], result.get('issues'))
        draft = result['draft']
        raw = session.editor.apply(session.editor.sha256, draft['changes'], draft['appendRows'], draft['flowOrder'],
                                   draft['trimStorageMappings'], draft['verifyStorageMappings'])
        _, compiled = session.compile(raw)
        self.assertTrue(compiled.is_valid, [d.message for d in compiled.diagnostics.errors])
        plan = compiled.plan.catalogs.standard_verify_plans[3]
        self.assertEqual(plan.steps[0].module_plan.pattern.name, values['_enterPattern'])
        row = WorkbookEditor(raw, 'verify.xlsx').by_name['FLOW']['rows'][4]
        columns = session.editor.by_name['FLOW']['columns']
        for name in ('DAC bit', 'Bisection', 'step', 'LINE'):
            self.assertEqual(row[next(i for i,c in enumerate(columns) if c['label']==name)]['value'], '/')

    def test_invalid_search_and_unconfirmed_storage_do_not_prepare_rows(self):
        session = self.session
        payload = self.payload(action='prepare', before=4, values=self.filled(**{'DAC bit': '30'}))
        result = session.new_flow_item(payload)
        self.assertFalse(result['ready'])
        self.assertNotIn('draft', result)
        with self.assertRaisesRegex(ValueError, '明确选择'):
            session.new_flow_item(self.payload(action='prepare', before=4, values=self.filled(_storageMode='')))
        with self.assertRaisesRegex(ValueError, '已存在'):
            session.new_flow_item(self.payload(action='prepare', before=4, values=self.filled(item='VREF0P8_TC')))

    def test_storage_flags_require_targets_and_reject_occupied_fields(self):
        session = self.session
        values = self.filled(OTP='yes', _storageMode='map')
        with self.assertRaisesRegex(ValueError, '必须选择 ADDR_MAP'):
            session.new_flow_item(self.payload(action='prepare', before=4, values=values))
        view = session.new_flow_item(self.payload(action='options', values=values))
        occupied = next(p for p in view['storage'] if 'VREF0P8_TC'.casefold() in p['reason'].casefold())
        target = {k: occupied[k] for k in ('page', 'register', 'field')}
        with self.assertRaises(ValueError):
            session.new_flow_item(self.payload(action='prepare', before=4, values=values, targets=[target]))
        with self.assertRaisesRegex(ValueError, '普通 FBIST'):
            session.new_flow_item(self.payload(action='prepare', before=4, values=self.filled(mode='swd')))

    def test_otp_and_uid_mappings_reach_the_real_compiler(self):
        session = self.session
        targets = [
            {'page': 'uid', 'register': 'Hsi_8M_Trim', 'field': 'HSI_TC_TRIMCR_VDDK'},
            {'page': 'trim0', 'register': 'FLASH_OP3', 'field': 'TRIM_LAT'},
        ]
        values = self.filled(OTP='yes', UID='yes', _storageMode='map', **{'DAC bit': '4'})
        options = session.new_flow_item(self.payload(action='options', values=values))
        values['_dacPattern'] = next(p['name'] for p in options['patterns']['dac'] if p['width'] == 4 and not p['reason'])
        result = session.new_flow_item(self.payload(action='prepare', before=4,
            values=values, targets=targets))
        self.assertTrue(result['ready'], result.get('issues'))
        draft = result['draft']
        raw = session.editor.apply(session.editor.sha256, draft['changes'], draft['appendRows'], draft['flowOrder'],
                                   draft['trimStorageMappings'], draft['verifyStorageMappings'])
        reopened = WorkbookEditor(raw, 'mapped-new-item.xlsx')
        self.assertEqual(reopened.trim_storage_mappings[-1], {'item': 'Wizard_New', 'targets': targets})
        _, compiled = session.compile(raw)
        self.assertTrue(compiled.is_valid, [d.message for d in compiled.diagnostics.errors])

    def test_incomplete_item_can_choose_resources_without_full_compile(self):
        before = self.session.editor.raw
        with patch.object(self.session, 'compile', side_effect=AssertionError('must not compile incomplete row')):
            result = self.session.new_flow_item(self.payload(action='options'))
        self.assertEqual(result['kind'], 'STANDARD_TRIM')
        self.assertIn('enter_vref0p8_tc', [p['name'] for p in result['patterns']['enter']])
        self.assertIn('set_vref0p8_tc_dac', [p['name'] for p in result['patterns']['dac']])
        self.assertTrue(result['storage'])
        self.assertTrue(result['positions'])
        self.assertEqual(self.session.editor.raw, before)
        self.assertIsNone(self.session.checked)


if __name__ == '__main__':
    unittest.main()
