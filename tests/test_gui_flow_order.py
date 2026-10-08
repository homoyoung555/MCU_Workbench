from io import BytesIO
from tests.workspace_paths import FORMAL_MASTER
import unittest
from unittest.mock import patch
from zipfile import ZipFile

import openpyxl
import xml.etree.ElementTree as ET

from gui.workbook_editor import NS, WorkbookEditor


def reordered(order, *moves):
    result = list(order)
    for destination, source in moves:
        result[destination - 1] = source
    return result


class FlowOrderTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        path = FORMAL_MASTER
        cls.editor = WorkbookEditor(path.read_bytes(), path.name, path)

    def test_append_and_edit_are_applied_before_real_xml_row_reorder(self):
        editor = self.editor
        appended_row = editor.by_name['FLOW']['maxRow'] + 1
        order = reordered(range(1, appended_row + 1), (82, 84), (83, 85),
                          (84, 82), (85, 83))
        raw = editor.apply(
            editor.sha256,
            [{'sheet': 'FLOW', 'cell': f'S{appended_row}', 'old': '/', 'value': '0.801'},
             {'sheet': 'FLOW', 'cell': 'S84', 'old': '/', 'value': '0.777'}],
            [{'sheet': 'FLOW', 'sourceRow': 207, 'copy': True}],
            order,
        )
        with ZipFile(BytesIO(editor.raw)) as before, ZipFile(BytesIO(raw)) as after:
            flow = editor.sheet_paths['FLOW']
            self.assertEqual([part.filename for part in before.infolist()
                              if before.read(part.filename) != after.read(part.filename)], [flow])
            old_rows = ET.fromstring(before.read(flow)).find(NS + 'sheetData')
            new_rows = ET.fromstring(after.read(flow)).find(NS + 'sheetData')
            self.assertEqual(new_rows[81].attrib['r'], '82')
            self.assertEqual(new_rows[81][0].attrib['r'], 'A82')
            self.assertEqual(new_rows[81][0].attrib.get('s'), old_rows[83][0].attrib.get('s'))
        workbook = openpyxl.load_workbook(BytesIO(raw), data_only=True)
        try:
            self.assertEqual(workbook['FLOW']['B82'].value, 'FLASH_160M_Test')
            self.assertEqual(workbook['FLOW']['S82'].value, '0.777')
            self.assertEqual(workbook['FLOW'][f'S{appended_row}'].value, '0.801')
            self.assertEqual(workbook['FLOW']['B84'].value, 'PWR_CP1_Test')
        finally:
            workbook.close()
        self.assertEqual(editor.apply(editor.sha256, [], [], list(range(1, editor.by_name['FLOW']['maxRow'] + 1))), editor.raw)

    def test_rejects_cross_cp_and_split_item_group_orders(self):
        editor = self.editor
        identity = list(range(1, editor.by_name['FLOW']['maxRow'] + 1))
        cross_cp = reordered(identity, (96, 97), (97, 96))
        split_group = reordered(identity, (82, 84), (84, 82))
        for order in (cross_cp, split_group, identity + [identity[-1]]):
            with self.subTest(order=order[:4]), self.assertRaises(ValueError):
                editor.apply(editor.sha256, [], [], order)

    def test_insert_blank_between_rows_and_delete_preserve_excel_content(self):
        editor = self.editor
        total = editor.by_name['FLOW']['maxRow']
        order = list(range(1, total + 1))
        order.insert(4, total + 1)
        changes = [{'sheet': 'FLOW', 'cell': f'A{total+1}', 'old': None, 'value': 'CP1'}]
        raw = editor.apply(editor.sha256, changes, [{'sheet':'FLOW','sourceRow':4,'copy':False}], order)
        saved = WorkbookEditor(raw, 'insert.xlsx')
        self.assertEqual(saved.by_name['FLOW']['rows'][4][0]['value'], 'CP1')
        self.assertIsNone(saved.by_name['FLOW']['rows'][4][1]['value'])
        self.assertEqual(saved.by_name['FLOW']['rows'][5][1]['value'], editor.by_name['FLOW']['rows'][4][1]['value'])
        appended = editor.apply(editor.sha256, changes, [{'sheet':'FLOW','sourceRow':4,'copy':False}])
        original_book = openpyxl.load_workbook(BytesIO(appended))
        moved_book = openpyxl.load_workbook(BytesIO(raw))
        try:
            def rules(book, cell):
                return sorted((v.type, v.formula1 or '', v.formula2 or '')
                              for v in book['FLOW'].data_validations.dataValidation if cell in v)
            for destination, source in [(5, total+1), (6, 5), (98, 97)]:
                for col in ['A', 'E', 'J', 'U']:
                    self.assertEqual(rules(moved_book, f'{col}{destination}'), rules(original_book, f'{col}{source}'))
        finally:
            original_book.close()
            moved_book.close()
        order.remove(4)
        raw = editor.apply(editor.sha256, changes, [{'sheet':'FLOW','sourceRow':4,'copy':False}], order)
        saved = WorkbookEditor(raw, 'delete.xlsx')
        self.assertEqual(saved.by_name['FLOW']['maxRow'], total)
        self.assertIsNone(saved.by_name['FLOW']['rows'][3][1]['value'])
        self.assertEqual(saved.by_name['FLOW']['rows'][4][1]['value'], editor.by_name['FLOW']['rows'][4][1]['value'])

    def test_rejects_moved_protected_rows(self):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = 'FLOW'
        sheet.append(['CP NO', 'item', 'type', 'Recipe'])
        sheet.append(['CP1', 'A', 'verify', ''])
        sheet.append(['CP1', 'B', 'verify', ''])
        sheet.protection.sheet = True
        sheet['A2'].protection = openpyxl.styles.Protection(locked=True)
        stream = BytesIO()
        workbook.save(stream)
        workbook.close()
        editor = WorkbookEditor(stream.getvalue(), 'protected.xlsx')
        with self.assertRaises(ValueError):
            editor.apply(editor.sha256, [], [], [1, 3, 2])

    def test_fingerprint_and_copyable_rows_follow_flow_order(self):
        from gui.server import effective_copy_editor, fingerprint
        payload = {'baseSha256': 'base', 'changes': [], 'appendRows': [], 'flowOrder': [1, 3, 2]}
        self.assertNotEqual(fingerprint(payload), fingerprint({**payload, 'flowOrder': [1, 2, 3]}))
        source = type('Editor', (), {'by_name': {'FLOW': {'maxRow': 3, 'appendableRows': [2]}}})()
        self.assertEqual(effective_copy_editor(source, [], payload['flowOrder']).by_name['FLOW']['appendableRows'], [3])

    def test_pure_order_changes_real_flow_rows(self):
        editor = self.editor
        order = reordered(range(1, editor.by_name['FLOW']['maxRow'] + 1), (4, 5), (5, 4))
        raw = editor.apply(editor.sha256, [], [], order)
        workbook = openpyxl.load_workbook(BytesIO(raw), data_only=True)
        try:
            self.assertEqual(workbook['FLOW']['B4'].value, 'Iref1u')
            self.assertEqual(workbook['FLOW']['B5'].value, 'Vref0p8_ABS')
        finally:
            workbook.close()

    def test_item_preview_uses_the_displayed_physical_row(self):
        from gui.server import Session, fingerprint
        from types import SimpleNamespace
        import pandas as pd
        session = Session()
        order = reordered(range(1, session.editor.by_name['FLOW']['maxRow'] + 1), (4, 5), (5, 4))
        payload = {'baseSha256': session.editor.sha256, 'changes': [], 'appendRows': [],
                   'flowOrder': order, 'row': 4, 'platform': 'chroma', 'variant': 'lab'}
        session.item_cache = {'key': fingerprint(payload), 'workbook': SimpleNamespace(dataframe=lambda _: pd.DataFrame()), 'compilation': object()}
        with patch('gui.item_preview.render_item_preview', return_value={'valid': True, 'text': ''}) as render:
            session.item_preview(payload)
        self.assertEqual(render.call_args.args[2], 4)


if __name__ == '__main__':
    unittest.main()
