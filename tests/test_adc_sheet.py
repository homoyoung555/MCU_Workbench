import unittest

import pandas as pd
from openpyxl import Workbook

import analog_sweep as A
from program_gen import state as G
from program_gen.platforms.chroma.adc_writer import _adc_channels
from program_gen.platforms.s100.adc_writer import _adc_channels_full


def _workbook_with_iomux():
    wb = Workbook()
    iomux = wb.active
    iomux.title = "IOMUX"
    iomux.append(["pin", "analog", None, None])
    iomux.append([None, "ADC", "ADC FORCE", "ADC LIMIT"])
    iomux.append(["PA01", "ADC1_IN0", 3.0, 99])
    iomux.append(["INT_BUFOUT", "ADC12_IN1", 2.0, 98])
    return wb


def _add_adc_sheet(wb):
    adc = wb.create_sheet("ADC")
    adc.append(["ADC", "Chanel", "step", "force", "limit"])
    adc.append(["ADC1", 0, 11, 1.25, 16])
    adc.append([None, 1, None, 1.2, 32])


class AdcSheetTests(unittest.TestCase):
    def setUp(self):
        self.old_wb = G.wb
        self.old_version = G.VERSION
        G.VERSION = pd.DataFrame({"NRST PIN": ["PD11"]})
        G._col_cache.clear()

    def tearDown(self):
        G.wb = self.old_wb
        G.VERSION = self.old_version
        G._col_cache.clear()

    def test_separate_adc_sheet_takes_priority_on_both_platforms(self):
        G.wb = _workbook_with_iomux()
        _add_adc_sheet(G.wb)

        self.assertEqual(
            _adc_channels_full(1, {"PA1"}),
            [
                (0, "PA1", 1.25, "PA1", True, 16),
                (1, "BUFOUT", 1.2, None, True, 32),
            ],
        )
        self.assertEqual(
            _adc_channels(1, {"PA1"}),
            [
                (0, "PA1", 1.25, "PA1", True, 16),
                (1, None, 1.2, "BUFOUT", True, 32),
            ],
        )
        self.assertEqual(A.adc_step(G.wb, 1), "11")

    def test_adc_step_may_only_be_set_once_per_adc_group(self):
        G.wb = _workbook_with_iomux()
        _add_adc_sheet(G.wb)
        G.wb["ADC"]["C3"] = 12

        with self.assertRaisesRegex(ValueError, "ADC1.*step"):
            A.adc_step(G.wb, 1)

    def test_iomux_is_used_when_adc_sheet_is_absent(self):
        G.wb = _workbook_with_iomux()

        self.assertEqual(
            _adc_channels_full(1, {"PA1"}),
            [
                (0, "PA1", 3.0, "PA1", True, 99),
                (1, "BUFOUT", 2.0, None, True, 98),
            ],
        )
        self.assertEqual(
            _adc_channels(1, {"PA1"}),
            [
                (0, "PA1", 3.0, "PA1", True, 99),
                (1, None, 2.0, "BUFOUT", True, 98),
            ],
        )

    def test_malformed_adc_sheet_does_not_silently_fall_back_to_iomux(self):
        G.wb = _workbook_with_iomux()
        adc = G.wb.create_sheet("ADC")
        adc.append(["ADC", "force", "limit"])
        adc.append(["ADC1", 1.25, 16])

        with self.assertRaisesRegex(ValueError, "ADC sheet.*CH"):
            _adc_channels_full(1, {"PA1"})

    def test_partially_filled_adc_row_reports_missing_value(self):
        G.wb = _workbook_with_iomux()
        adc = G.wb.create_sheet("ADC")
        adc.append(["ADC", "CH", "force", "limit"])
        adc.append(["ADC1", 0, None, 16])

        with self.assertRaisesRegex(ValueError, "第 2 行.*force"):
            _adc_channels_full(1, {"PA1"})

    def test_new_adc_sheet_only_requires_iomux_adc_mapping_column(self):
        G.wb = Workbook()
        iomux = G.wb.active
        iomux.title = "IOMUX"
        iomux.append(["pin", "analog"])
        iomux.append([None, "ADC"])
        iomux.append(["PA01", "ADC1_IN0"])
        iomux.append(["INT_BUFOUT", "ADC12_IN1"])
        _add_adc_sheet(G.wb)

        self.assertEqual(
            _adc_channels_full(1, {"PA1"}),
            [
                (0, "PA1", 1.25, "PA1", True, 16),
                (1, "BUFOUT", 1.2, None, True, 32),
            ],
        )

    def test_adc_channel_without_iomux_mapping_is_rejected(self):
        G.wb = _workbook_with_iomux()
        adc = G.wb.create_sheet("ADC")
        adc.append(["ADC", "CH", "force", "limit"])
        adc.append(["ADC1", 2, 1.0, 16])

        with self.assertRaisesRegex(ValueError, "ADC1 CH2.*IOMUX"):
            _adc_channels_full(1, {"PA1"})

    def test_channels_sharing_one_pin_must_use_same_force(self):
        G.wb = Workbook()
        iomux = G.wb.active
        iomux.title = "IOMUX"
        iomux.append(["pin", "analog"])
        iomux.append([None, "ADC"])
        iomux.append(["PA01", "ADC1_IN0/ADC1_IN1"])
        adc = G.wb.create_sheet("ADC")
        adc.append(["ADC", "CH", "force", "limit"])
        adc.append(["ADC1", 0, 1.0, 16])
        adc.append([None, 1, 1.2, 16])

        with self.assertRaisesRegex(ValueError, "PA01.*force 不一致"):
            _adc_channels_full(1, {"PA1"})


if __name__ == "__main__":
    unittest.main()
