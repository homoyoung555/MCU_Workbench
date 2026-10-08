import hashlib
import io
import unittest
from zipfile import ZipFile
import xml.etree.ElementTree as ET

import openpyxl


def workbook_bytes():
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Sheet1"
    sheet.append(["pattern", "addr", "data", "type", "step", "reset value", "continue", "note", "extra", "formula", None, None])
    sheet.append(["PT104_demo", "0XE000EDF0", "0X22000042", "W", "/", None, "Y", "keep", "extra-value", "=LEN(I2)"])
    sheet.append([None, "WAIT", 200, "us", "/", None, None, "wait", None])
    sheet["I2"].fill = openpyxl.styles.PatternFill("solid", fgColor="00FF00")
    sheet["L1"].fill = openpyxl.styles.PatternFill("solid", fgColor="FFFFFF")
    other = book.create_sheet("unchanged")
    other["A1"] = "must survive"
    raw = io.BytesIO()
    book.save(raw)
    return raw.getvalue()


def six_column_workbook_bytes():
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Sheet1"
    sheet.append(["pattern", "addr", "data", "type", "continue", "description"])
    sheet.append(["PT104_demo", "0X40000000", "0X1", "W", "N", "old row"])
    raw = io.BytesIO()
    book.save(raw)
    return raw.getvalue()


class SwdGeneratorTests(unittest.TestCase):
    def test_save_keeps_linked_comments_on_selected_sheet(self):
        from gui.swd_generator import export_excel, load_excel
        book = openpyxl.load_workbook(io.BytesIO(workbook_bytes()))
        book['Sheet1']['I2'].comment = openpyxl.comments.Comment('source note', 'author')
        stream = io.BytesIO(); book.save(stream)
        raw = stream.getvalue(); model, sheet = load_excel(raw)
        model['rows'][0]['note'] = 'edited'
        result = openpyxl.load_workbook(io.BytesIO(export_excel(raw, sheet, model['rows'])))
        self.assertEqual(result[sheet]['I2'].comment.text, 'source note')
        self.assertEqual(result['unchanged']['A1'].value, 'must survive')

    def test_partial_save_removes_stale_formula_chain_and_requests_recalculation(self):
        from gui.swd_generator import export_excel, load_excel
        raw = workbook_bytes()
        archive = io.BytesIO()
        with ZipFile(io.BytesIO(raw)) as source, ZipFile(archive, 'w') as target:
            for entry in source.infolist():
                data = source.read(entry.filename)
                if entry.filename == 'xl/_rels/workbook.xml.rels':
                    data = data.replace(b'</Relationships>', b'<Relationship Id="rIdChain" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/calcChain" Target="calcChain.xml"/></Relationships>')
                if entry.filename == '[Content_Types].xml':
                    data = data.replace(b'</Types>', b'<Override PartName="/xl/calcChain.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.calcChain+xml"/></Types>')
                target.writestr(entry, data)
            target.writestr('xl/calcChain.xml', '<calcChain xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><c r="J2" i="1"/></calcChain>')
        raw = archive.getvalue()
        model, sheet = load_excel(raw)
        model['rows'][0]['note'] = 'edited'
        output = export_excel(raw, sheet, model['rows'])
        with ZipFile(io.BytesIO(output)) as package:
            self.assertNotIn('xl/calcChain.xml', package.namelist())
            self.assertNotIn(b'calcChain', package.read('xl/_rels/workbook.xml.rels'))
            self.assertNotIn(b'calcChain', package.read('[Content_Types].xml'))
            props = ET.fromstring(package.read('xl/workbook.xml')).find('{http://schemas.openxmlformats.org/spreadsheetml/2006/main}calcPr')
            self.assertEqual(props.attrib['fullCalcOnLoad'], '1')
        book = openpyxl.load_workbook(io.BytesIO(output))
        self.assertEqual(book[sheet]['J2'].value, '=LEN(I2)')
        self.assertEqual(book['unchanged']['A1'].value, 'must survive')

    def test_save_changes_only_selected_sheet_and_styles_in_excel_package(self):
        from gui.swd_generator import export_excel, load_excel
        raw = workbook_bytes()
        original = io.BytesIO()
        with ZipFile(io.BytesIO(raw)) as source, ZipFile(original, 'w') as target:
            for entry in source.infolist():
                data = source.read(entry.filename)
                if entry.filename == 'xl/worksheets/sheet2.xml':
                    data = data.replace(b'</worksheet>', b'<!-- untouched auxiliary data --></worksheet>')
                target.writestr(entry, data)
        raw = original.getvalue()
        model, sheet = load_excel(raw)
        model['rows'][0]['note'] = 'changed'
        output = export_excel(raw, sheet, model['rows'])
        with ZipFile(io.BytesIO(raw)) as before, ZipFile(io.BytesIO(output)) as after:
            self.assertEqual(set(before.namelist()), set(after.namelist()))
            for name in before.namelist():
                if name not in ('xl/worksheets/sheet1.xml', 'xl/styles.xml'):
                    self.assertEqual(before.read(name), after.read(name), name)
        self.assertEqual(load_excel(output)[0]['rows'][0]['note'], 'changed')

    def test_chroma_write_preserves_original_drw_command(self):
        from gui.swd_generator import generate_pattern
        output = generate_pattern('PT104_demo', [{'pattern': 'PT104_demo', 'addr': '0X40021000',
                                                  'data': '0X1', 'type': 'W', 'continue': 'N'}], 'chroma')
        original_drw = b' *111* RPT 2;\n *101*;\n *111* RPT 3;\n *101*;\n *111*;\n *1H1*;\n *1L1* RPT 2;\n *1X1*;\n *101*;\n'
        self.assertIn(original_drw + b' *111*;//1', output)

    def test_s100_repeat_instruction_is_separated_from_vector(self):
        from gui.swd_generator import generate_pattern
        output = generate_pattern('PT104_demo', [{'pattern': 'PT104_demo', 'addr': '0X40021000',
                                                  'data': '0X1', 'type': 'W', 'continue': 'Y'}], 's100').decode()
        for line in output.splitlines():
            if 'RPT ' in line:
                self.assertIn(';', line.split('RPT ', 1)[0], line)

    def test_excel_round_trip_keeps_extra_sheet_fields_styles_and_formula(self):
        from gui.swd_generator import export_excel, load_excel

        raw = workbook_bytes()
        model, sheet = load_excel(raw)
        self.assertEqual(sheet, "Sheet1")
        self.assertEqual(model["rows"][0]["row"], 2)
        self.assertEqual(model["rows"][0]["extra"], "extra-value")
        self.assertIn("__blank_11", model["headers"])
        self.assertEqual(export_excel(raw, sheet, model["rows"]), raw)

        edited = [dict(row) for row in model["rows"]]
        edited[0]["note"] = "edited"
        output = export_excel(raw, sheet, edited)
        book = openpyxl.load_workbook(io.BytesIO(output), data_only=False)
        self.assertEqual(book["Sheet1"]["H2"].value, "edited")
        self.assertEqual(book["Sheet1"]["I2"].value, "extra-value")
        self.assertEqual(book["Sheet1"]["I2"].fill.fgColor.rgb, "0000FF00")
        self.assertEqual(book["Sheet1"]["J2"].value, "=LEN(I2)")
        self.assertEqual(book["unchanged"]["A1"].value, "must survive")

    def test_excel_export_orders_existing_rows_and_appends_new_row(self):
        from gui.swd_generator import export_excel, load_excel

        raw = workbook_bytes()
        model, sheet = load_excel(raw)
        new = dict(model["rows"][0], row=None, addr="WAIT", data=1, type="ms", note="new")
        output = export_excel(raw, sheet, [model["rows"][1], new])
        book = openpyxl.load_workbook(io.BytesIO(output), data_only=False)
        self.assertEqual(book[sheet]["B2"].value, "WAIT")
        self.assertEqual(book[sheet]["H3"].value, "new")
        self.assertEqual(book[sheet]["I3"].value, "extra-value")

    def test_optional_columns_are_added_and_new_rows_do_not_inherit_operations(self):
        from gui.swd_generator import export_excel, load_excel

        raw = six_column_workbook_bytes()
        model, sheet = load_excel(raw)
        changed = [dict(model["rows"][0], step="3", **{"reset value": "0X00000000"}),
                   {"row": None, "pattern": "PT104_new", "description": "new row"}]
        output = export_excel(raw, sheet, changed)
        book = openpyxl.load_workbook(io.BytesIO(output), data_only=False)
        headers = [cell.value for cell in book[sheet][1]]
        self.assertEqual(book[sheet].cell(2, headers.index("step") + 1).value, "3")
        self.assertEqual(book[sheet].cell(2, headers.index("reset value") + 1).value, "0X00000000")
        self.assertIsNone(book[sheet]["B3"].value)
        self.assertIsNone(book[sheet]["D3"].value)

    def test_generates_legacy_chroma_and_s100_for_all_operations(self):
        from gui.swd_generator import generate_pattern

        rows = [
            {"row": 2, "pattern": "PT104_demo", "addr": "0XE000EDF0", "data": "0X22000042", "type": "W", "step": "/", "continue": "Y"},
            {"row": 3, "pattern": "PT104_demo", "addr": "0X40000000", "data": "0X00000001", "type": "R", "step": "/"},
            {"row": 4, "pattern": "PT104_demo", "addr": "0X40000004", "data": "0X00000001", "type": "C_H", "step": "/"},
            {"row": 5, "pattern": "PT104_demo", "addr": "0X40000008", "data": "0X00000000", "type": "C_L", "step": "/"},
            {"row": 6, "pattern": "PT104_demo", "addr": "0X4000000C", "data": "0X80000000", "type": "CAP", "step": "/"},
            {"row": 7, "pattern": "PT104_demo", "addr": "0X20000000", "data": "0X20000004", "type": "SRAM", "step": "/"},
            {"row": 8, "pattern": "PT104_demo", "addr": "WAIT", "data": 200, "type": "us", "step": "/"},
            {"row": 9, "pattern": "PT104_demo", "addr": "WAIT", "data": 1, "type": "ms", "step": "/"},
        ]
        chroma = generate_pattern("PT104_demo", rows, "chroma", chip="PT104")
        s100 = generate_pattern("PT104_demo", rows, "s100", {"swc": "PA14", "swd": "PA13", "reset": "NRST", "trig": "TRIGPIN"}, "PCD009")
        self.assertIsInstance(chroma, bytes)
        self.assertIn(b"SPM_PATTERN(PT104_demo)", chroma)
        self.assertIn(b"//read0X40000008", chroma)
        self.assertIn(b"RPT 44;//wait 200us", chroma)
        self.assertIn(b"RPT 1000;//wait 1ms", chroma)
        self.assertIn(b" *1L1*;//00000000", chroma)  # C_L keeps a zero-bit check.
        self.assertIn(b"@@PATTERN_DEFINE", s100)
        self.assertNotIn(b"SPM_PATTERN", s100)
        self.assertIn(b"PA14,PA13,NRST,TRIGPIN", s100)
        self.assertIn(b"PTN_TRIG", s100)
        self.assertTrue(s100.endswith(b"@@END_PATTERN_DEFINE\n"))
        self.assertEqual(hashlib.sha256(chroma).hexdigest(), "7f4f456dc9e78b673013d8a5716f91947172a30f174b53f1f5effb41d3188475")
        # Includes the original S100 chip_db commands and vector separators.
        self.assertEqual(hashlib.sha256(s100).hexdigest(), "f5bfa0938c220e25ee46477cc9aa8726cbd28b2ded446044ed06c220e77b9b98")

    def test_unknown_operation_fails_with_source_row(self):
        from gui.swd_generator import generate_pattern

        with self.assertRaisesRegex(ValueError, "第7行.*未知"):
            generate_pattern("PT104_demo", [{"row": 7, "pattern": "PT104_demo", "addr": "0X0", "data": "0X0", "type": "ERASE"}], "chroma")

    def test_pattern_title_without_operation_is_legal(self):
        from gui.swd_generator import generate_pattern

        output = generate_pattern("PT104_demo", [
            {"row": 2, "pattern": "PT104_demo", "addr": None, "type": None, "continue": "N"},
            {"row": 3, "pattern": None, "addr": "WAIT", "data": 1, "type": "ms"},
        ], "s100")
        self.assertIn(b"RPT 1000;//wait 1ms", output)


if __name__ == "__main__":
    unittest.main()
