import hashlib
import io
import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from gui.swd_materials import MaterialLibrary


SVD = b'''<device><name>PY32F420</name><peripherals>
  <peripheral><name>RCC</name><baseAddress>0x40021000</baseAddress><registers>
    <register><name>CSR</name><addressOffset>0x34</addressOffset><resetValue>0x00000000</resetValue></register>
  </registers></peripheral>
  <peripheral derivedFrom="RCC"><name>RCC_COPY</name><baseAddress>0x40021400</baseAddress></peripheral>
</peripherals></device>'''

SVD_CACHE = json.dumps({
    "version": "1.0",
    "data": {
        "device": {"name": "PY32F420xx"},
        "peripherals": [{
            "name": "RCC",
            "baseAddress": 0x40021000,
            "description": "Reset and clock control",
            "registers": [{
                "name": "CSR",
                "address": 0x40021034,
                "addressOffset": 0x34,
                "description": "Clock control/status register",
                "resetValue": "32'b00000000000000000000000000000010",
                "access": "RW",
                "fields": [
                    {"name": "LSION", "bitRange": "[0:0]", "bitOffset": 0, "bitWidth": 1,
                     "access": "RW", "description": "Internal low-speed clock enable", "resetValue": "1'b1"},
                    {"name": "UNKNOWN", "bitRange": "[2:1]", "bitOffset": 1, "bitWidth": 2,
                     "access": "RO", "description": "Unknown reset", "resetValue": "2'bx1"},
                    {"name": "INVALID", "bitRange": "[3:3]", "bitOffset": "3", "bitWidth": 1},
                ],
            }, {"name": "INVALID_ADDRESS", "address": "0x40021038", "fields": []}],
        }],
    },
}, ensure_ascii=False).encode("utf-8")

REAL_SVD_CACHE = Path(__file__).resolve().parents[1] / "gui" / "data" / "py32f420_svd_cache.json"


def register_book():
    book = Workbook()
    sheet = book.active
    sheet.title = "RCC"
    sheet.append(["RegName", "BaseAddress", "OffsetAddress", "ResetValue"])
    sheet.append(["CSR", "0x40021000; 0x40021400", "0x34", "0x00000001"])
    sheet.append(["CFGR", None, "0x08", "0x00000000"])
    raw = io.BytesIO()
    book.save(raw)
    return raw.getvalue()


class MaterialLibraryTests(unittest.TestCase):
    def test_svd_cache_preserves_register_and_field_facts_after_reload(self):
        with tempfile.TemporaryDirectory() as temp:
            library = MaterialLibrary(temp)
            first = library.import_material(SVD_CACHE, "py32f420_svd_cache.json", "svd_cache")
            second = library.import_material(SVD_CACHE, "renamed.json", "svd_cache")
            state = library.state()

            self.assertEqual(first["kind"], "svd_cache")
            self.assertEqual(second, first)
            self.assertEqual(len(state["materials"]), 1)
            source = next(item for item in state["registers"] if item["address"] == "0x40021034")["sources"][0]
            self.assertEqual(source["description"], "Clock control/status register")
            self.assertEqual(source["resetValue"], "0x2")
            self.assertEqual(source["fields"], [
                {"name": "LSION", "bits": "[0:0]", "bitOffset": 0, "bitWidth": 1,
                 "access": "RW", "description": "Internal low-speed clock enable", "resetValue": "0x1"},
                {"name": "UNKNOWN", "bits": "[2:1]", "bitOffset": 1, "bitWidth": 2,
                 "access": "RO", "description": "Unknown reset", "resetValue": "2'bx1"},
            ])
            self.assertEqual(MaterialLibrary(temp).state(), state)

    def test_svd_cache_rejects_invalid_data_without_material(self):
        invalid = json.dumps({"data": {"device": {"name": "OTHER"}, "peripherals": []}}).encode("utf-8")
        with tempfile.TemporaryDirectory() as temp:
            library = MaterialLibrary(temp)

            with self.assertRaisesRegex(ValueError, "PY32F420"):
                library.import_material(invalid, "invalid.json", "svd_cache")

            self.assertEqual(library.state()["materials"], [])

    def test_real_py32f420_svd_cache_import_has_verified_counts(self):
        raw = REAL_SVD_CACHE.read_bytes()
        with tempfile.TemporaryDirectory() as temp:
            material = MaterialLibrary(temp).import_material(raw, REAL_SVD_CACHE.name, "svd_cache")

        self.assertEqual(material["sha256"], hashlib.sha256(raw).hexdigest())
        self.assertEqual(material["analysis"], {
            "device": "PY32F420xx", "peripheralCount": 47, "registerCount": 525, "fieldCount": 3717,
        })

    def test_svd_is_persisted_and_derives_register_addresses(self):
        with tempfile.TemporaryDirectory() as temp:
            library = MaterialLibrary(temp)
            material = library.import_material(SVD, "PY32F420.svd")
            state = library.state()

            self.assertEqual(material["sha256"], hashlib.sha256(SVD).hexdigest())
            self.assertEqual(state["materials"][0]["originalName"], "PY32F420.svd")
            registers = {item["address"]: item for item in state["registers"]}
            self.assertIn("0x40021034", registers)
            self.assertIn("0x40021434", registers)
            self.assertEqual(registers["0x40021434"]["sources"][0]["name"], "CSR")
            self.assertTrue((Path(temp) / material["storedAs"]).is_file())
            self.assertEqual(MaterialLibrary(temp).state(), state)

    def test_excel_inherits_base_and_expands_each_base_address(self):
        with tempfile.TemporaryDirectory() as temp:
            library = MaterialLibrary(temp)
            library.import_material(register_book(), "registers.xlsx", "excel")
            registers = {item["address"]: item for item in library.state()["registers"]}

            self.assertEqual(registers["0x40021034"]["sources"][0]["name"], "CSR")
            self.assertIn("0x40021434", registers)
            self.assertIn("0x40021008", registers)
            self.assertEqual(registers["0x40021008"]["sources"][0]["name"], "CFGR")

    def test_same_register_name_at_two_addresses_is_a_conflict(self):
        with tempfile.TemporaryDirectory() as temp:
            library = MaterialLibrary(temp)
            library.import_material(SVD, "chip.svd")
            other = SVD.replace(b"0x40021000", b"0x40022000")
            library.import_material(other, "other.svd")

            conflict = next(item for item in library.state()["conflicts"] if item["kind"] == "registerAddress")
            self.assertIn("0x40021034", conflict["values"])
            self.assertIn("0x40022034", conflict["values"])

    def test_excel_merged_field_rows_do_not_duplicate_a_register_source(self):
        book = Workbook()
        sheet = book.active
        sheet.append(["RegName", "BaseAddress", "OffsetAddress", "FieldName"])
        sheet.append(["CSR", "0x40021000", "0x34", "READY"])
        sheet.append([None, None, None, "LSION"])
        raw = io.BytesIO()
        book.save(raw)
        with tempfile.TemporaryDirectory() as temp:
            library = MaterialLibrary(temp)
            library.import_material(raw.getvalue(), "registers.xlsx")
            register = next(item for item in library.state()["registers"] if item["address"] == "0x40021034")

            self.assertEqual(len(register["sources"]), 1)

    def test_reimporting_the_same_raw_material_is_idempotent(self):
        with tempfile.TemporaryDirectory() as temp:
            library = MaterialLibrary(temp)
            first = library.import_material(SVD, "chip.svd")
            second = library.import_material(SVD, "renamed.svd")
            state = library.state()

            self.assertEqual(second, first)
            self.assertEqual(len(state["materials"]), 1)
            self.assertEqual(len(state["registers"]), 2)

    def test_c_only_accepts_a_single_literal_pointer_address_as_a_register(self):
        source = b'''#define PLAIN_RESET 0x40021034UL\n#define RCC_CSR ((volatile unsigned int *)0x40021034UL)\n#define RCC_COMPLEX ((volatile unsigned int *)(0x40021000UL + 0x34UL))\n'''
        with tempfile.TemporaryDirectory() as temp:
            library = MaterialLibrary(temp)
            library.import_material(source, "chip.h")
            state = library.state()

            self.assertEqual(len(state["registers"]), 1)
            self.assertEqual(state["registers"][0]["sources"][0]["name"], "RCC_CSR")
            unresolved = state["materials"][0]["analysis"]["unresolvedMacros"]
            self.assertIn("RCC_COMPLEX", [item["name"] for item in unresolved])

    def test_c_pointer_expression_with_decimal_arithmetic_is_not_a_register(self):
        source = b"#define REG (*(volatile uint32_t *)(0x40000000UL + 4UL))\n"
        with tempfile.TemporaryDirectory() as temp:
            library = MaterialLibrary(temp)
            library.import_material(source, "chip.h")
            state = library.state()

            self.assertEqual(state["registers"], [])
            self.assertEqual(state["materials"][0]["analysis"]["unresolvedMacros"][0]["name"], "REG")

    def test_excel_numeric_zero_offset_and_field_reset_do_not_inherit_previous_values(self):
        book = Workbook()
        sheet = book.active
        sheet.append(["RegName", "BaseAddress", "OffsetAddress", "FieldName", "Bits", "ResetValue"])
        sheet.append(["HIGH", "0x40021000", 16, "A", "[0]", 1])
        sheet.append(["ZERO", None, 0, "B", "[1]", 0])
        raw = io.BytesIO()
        book.save(raw)
        with tempfile.TemporaryDirectory() as temp:
            library = MaterialLibrary(temp)
            library.import_material(raw.getvalue(), "registers.xlsx")
            registers = {item["address"]: item for item in library.state()["registers"]}

            self.assertIn("0x40021000", registers)
            zero = registers["0x40021000"]["sources"][0]
            self.assertEqual(zero["name"], "ZERO")
            self.assertEqual(zero["fields"][0]["resetValue"], "0")

    def test_excel_keeps_instance_identity_and_field_reset_out_of_register_conflicts(self):
        book = Workbook()
        sheet = book.active
        sheet.title = "ADC"
        sheet.append(["RegName", "BaseAddress", "OffsetAddress", "FieldName", "Bits", "ResetValue"])
        sheet.append(["CR", "ADC1：0x4001 2400\nADC2：0x4001 2800", "0x00", "EN", "[0]", "1"])
        raw = io.BytesIO()
        book.save(raw)
        with tempfile.TemporaryDirectory() as temp:
            library = MaterialLibrary(temp)
            library.import_material(raw.getvalue(), "registers.xlsx")
            state = library.state()
            registers = {item["address"]: item for item in state["registers"]}

            self.assertEqual(registers["0x40012400"]["sources"][0]["peripheral"], "ADC1")
            self.assertEqual(registers["0x40012800"]["sources"][0]["peripheral"], "ADC2")
            self.assertEqual(registers["0x40012400"]["sources"][0]["resetValue"], "")
            self.assertEqual(registers["0x40012400"]["sources"][0]["fields"][0]["resetValue"], "1")
            self.assertFalse(any(item["kind"] == "registerAddress" for item in state["conflicts"]))

    def test_normalized_excel_name_matches_svd_and_field_resets_compare_by_field(self):
        def excel(reset):
            book = Workbook()
            sheet = book.active
            sheet.title = "RCC"
            sheet.append(["RegName", "BaseAddress", "OffsetAddress", "FieldName", "Bits", "ResetValue"])
            sheet.append(["RCC_CSR", "RCC：0x4002 1000", "0x34", "LSION", "[0]", reset])
            raw = io.BytesIO()
            book.save(raw)
            return raw.getvalue()

        with tempfile.TemporaryDirectory() as temp:
            library = MaterialLibrary(temp)
            library.import_material(SVD, "chip.svd")
            library.import_material(excel("0"), "registers.xlsx")
            self.assertFalse(any(item["kind"] == "addressIdentity" for item in library.state()["conflicts"]))
            library.import_material(excel("1"), "updated-registers.xlsx")
            conflicts = library.state()["conflicts"]

            self.assertFalse(any(item["kind"] == "resetValue" for item in conflicts))
            self.assertTrue(any(item["kind"] == "fieldReset" for item in conflicts))

    def test_header_keeps_source_extracts_only_confirmable_macros_and_keeps_conflict(self):
        header = b'''#define RCC_CSR ((volatile unsigned int *)0x40021034UL)\n#define RCC_CSR_RESET 0x00000002UL\n#define RCC_UNRESOLVED (RCC_BASE + 0x34UL)\n'''
        with tempfile.TemporaryDirectory() as temp:
            library = MaterialLibrary(temp)
            library.import_material(SVD, "chip.svd", "svd")
            material = library.import_material(header, "../unsafe\\chip.h", "h")
            state = library.state()

            self.assertNotIn("..", material["storedAs"])
            self.assertEqual(state["materials"][1]["analysis"]["unresolvedMacros"][0]["name"], "RCC_UNRESOLVED")
            self.assertIn("RCC_CSR", [item["name"] for item in state["materials"][1]["analysis"]["macros"]])
            conflict = next(item for item in state["conflicts"] if item["kind"] == "resetValue")
            library.confirm(conflict["id"], source_id=conflict["sourceIds"][0], note="以芯片手册复核")
            confirmation = MaterialLibrary(temp).state()["confirmations"][conflict["id"]]
            self.assertEqual(confirmation["note"], "以芯片手册复核")


if __name__ == "__main__":
    unittest.main()
