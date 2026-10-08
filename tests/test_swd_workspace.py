"""Public workspace lifecycle checks; no server, vendor compiler or hardware."""
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
import unittest
from unittest.mock import patch
from zipfile import ZipFile
import openpyxl


def source():
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = 'Sheet1'
    sheet.append(['pattern', 'addr', 'data', 'step', 'reset value', 'type', 'continue', 'description'])
    sheet.append(['PT104_demo', '0X40021000', '0X00000001', '/', '/', 'W', 'Y', 'demo'])
    sheet.append([None, '0X40021004', '0X00000002', '/', '/', 'W', 'N', 'second'])
    raw = BytesIO(); book.save(raw)
    return raw.getvalue()


def create_saved_pattern(workspace, payload):
    draft = workspace.create_pattern(payload)
    doc = next(item for item in draft['documents'] if item['id'] == draft['createdDocumentId'])
    saved = workspace.save({'documentId': doc['id'], 'rows': doc['rows']})
    saved.update(createdDocumentId=doc['id'], createdPattern=draft['createdPattern'])
    return saved


class SwdWorkspaceTests(unittest.TestCase):
    def test_complex_excel_fallback_reports_missing_when_it_cannot_preserve_ai_metadata(self):
        from gui.assistant_swd import export_candidate, read_provenance
        from gui.swd_workspace import SwdWorkspace
        result = SimpleNamespace(name='Assistant_Config', warnings=[], rows=[
            SimpleNamespace(type='W', addr=0x40021008, data=0x02000000,
                            mask=0x0F000000, phase=6, sources=['fixture'], desc='assistant')])
        candidate = export_candidate(result, chip_id='py32f420', product_profile_id='PCD009',
                                     platform='chroma', initial_state='reset', knowledge_version='fixture-v1')
        source, complex_book = BytesIO(candidate['excel']), BytesIO()
        with ZipFile(source) as package, ZipFile(complex_book, 'w') as output:
            for entry in package.infolist():
                output.writestr(entry, package.read(entry.filename))
            output.writestr('xl/worksheets/_rels/sheet1.xml.rels',
                            b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                            b'<Relationship Id="rId99" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" '
                            b'Target="https://example.invalid" TargetMode="External"/></Relationships>')
        with TemporaryDirectory() as temp:
            workspace = SwdWorkspace(SimpleNamespace(output=Path(temp), checked=None, editor=None))
            doc = workspace.import_excel(complex_book.getvalue(), 'complex.xlsx')['documents'][0]
            doc['rows'][0]['data'] = '0x03000000'
            saved = workspace.save({'documentId': doc['id'], 'rows': doc['rows']})
            self.assertEqual(saved['documents'][0]['provenance']['status'], 'missing')
            self.assertEqual(read_provenance(workspace.export_excel(doc['id']))['status'], 'missing')

    def test_ai_provenance_survives_stage_save_export_reload_and_downgrades_after_edit(self):
        from gui.assistant_swd import export_candidate, read_provenance
        from gui.swd_workspace import SwdWorkspace
        result = SimpleNamespace(name='Assistant_Config', warnings=[], rows=[
            SimpleNamespace(type='W', addr=0x40021008, data=0x02000000,
                            mask=0x0F000000, phase=6, sources=['fixture'], desc='assistant')])
        candidate = export_candidate(result, chip_id='py32f420', product_profile_id='PCD009',
                                     platform='chroma', initial_state='reset', knowledge_version='fixture-v1')
        with TemporaryDirectory() as temp:
            session = SimpleNamespace(output=Path(temp), checked=None, editor=None)
            workspace = SwdWorkspace(session)
            plain = workspace.import_excel(source(), 'manual.xlsx')['documents'][0]
            self.assertEqual(plain['provenance']['status'], 'missing')
            self.assertFalse(plain['provenance']['trusted'])
            staged = workspace.stage_excel(candidate['excel'], 'assistant.xlsx', 'a' * 32)
            self.assertEqual(staged['provenance']['status'], 'consistent')
            self.assertFalse(staged['provenance']['trusted'])
            self.assertEqual(staged['provenance']['productProfileId'], 'PCD009')
            saved = workspace.save({'documentId': staged['id'], 'rows': staged['rows']})
            document = next(doc for doc in saved['documents'] if doc['id'] == staged['id'])
            self.assertEqual(document['provenance']['status'], 'consistent')
            exported = workspace.export_excel(staged['id'])
            self.assertEqual(read_provenance(exported)['record']['productProfileId'], 'PCD009')
            restored = SwdWorkspace(session).state()['documents']
            document = next(doc for doc in restored if doc['id'] == staged['id'])
            self.assertEqual(document['provenance']['status'], 'consistent')
            document['rows'][0]['data'] = '0x03000000'
            changed = workspace.save({'documentId': document['id'], 'rows': document['rows']})
            document = next(doc for doc in changed['documents'] if doc['id'] == staged['id'])
            self.assertEqual(document['provenance']['status'], 'modified')
            self.assertNotIn('validation', document['provenance'])
            self.assertEqual(read_provenance(workspace.export_excel(staged['id']))['status'], 'modified')

    def test_create_is_only_a_draft_and_save_does_not_export_excel(self):
        from gui.swd_workspace import SwdWorkspace
        with TemporaryDirectory() as temp:
            session = SimpleNamespace(output=Path(temp), checked=None, editor=None)
            workspace = SwdWorkspace(session)
            doc = workspace.import_excel(source(), 'steps.xlsx')['documents'][0]
            original_source = Path(workspace.data['documents'][0]['source'])
            original_bytes = original_source.read_bytes()
            state_bytes = workspace.file.read_bytes()
            with patch('gui.swd_generator.export_excel', side_effect=AssertionError('save must not export Excel')):
                created = workspace.create_pattern({'documentId': doc['id'], 'name': 'PT104_draft'})
                self.assertEqual(workspace.file.read_bytes(), state_bytes)
                self.assertNotIn('PT104_draft', workspace.state()['documents'][0]['patterns'])
                rows = created['documents'][0]['rows']
                rows[-1].update(addr='0X40021008', data='0X1', type='W')
                saved = workspace.save({'documentId': doc['id'], 'rows': rows})
                self.assertIn('PT104_draft', saved['documents'][0]['patterns'])
                self.assertEqual(original_source.read_bytes(), original_bytes)
                self.assertEqual(workspace.data['documents'][0]['source'], str(original_source))
                restored = SwdWorkspace(session)
                restored_rows = restored.state()['documents'][0]['rows']
                self.assertEqual(restored_rows[-1]['data'], '0X1')
                self.assertEqual(len({row['row'] for row in restored_rows}), len(restored_rows))
                restored_rows[-1]['data'] = '0X2'
                restored.save({'documentId': doc['id'], 'rows': restored_rows})
            exported = openpyxl.load_workbook(BytesIO(restored.export_excel(doc['id'])))['Sheet1']
            self.assertEqual(exported['A4'].value, 'PT104_draft')
            self.assertEqual(exported['C4'].value, '0X2')

    def test_new_document_draft_does_not_persist_until_saved(self):
        from gui.swd_workspace import SwdWorkspace
        with TemporaryDirectory() as temp:
            session = SimpleNamespace(output=Path(temp), checked=None, editor=None)
            workspace = SwdWorkspace(session)
            created = workspace.create_pattern({'name': 'PT104_new'})
            workspace.configure(workspace.state()['config'])
            self.assertEqual(SwdWorkspace(session).state()['documents'], [])
            self.assertEqual(list(Path(temp).rglob('*.xlsx')), [])
            workspace.save({'documentId': created['createdDocumentId'], 'rows': created['documents'][0]['rows']})
            self.assertEqual(SwdWorkspace(session).state()['documents'][0]['patterns'], ['PT104_new'])

    def test_create_pattern_preserves_blank_rows_inside_merged_regions(self):
        from gui.swd_workspace import SwdWorkspace
        book = openpyxl.load_workbook(BytesIO(source()))
        sheet = book['Sheet1']
        sheet.insert_rows(3, 2)
        sheet.merge_cells('A2:A5')
        sheet.merge_cells('I2:J6')
        sheet['I2'] = 'note'
        raw = BytesIO(); book.save(raw)
        with TemporaryDirectory() as temp:
            workspace = SwdWorkspace(SimpleNamespace(output=Path(temp), checked=None, editor=None))
            doc = workspace.import_excel(raw.getvalue(), 'sparse.xlsx')['documents'][0]
            create_saved_pattern(workspace, {'documentId': doc['id'], 'name': 'PT104_TEST'})
            saved = openpyxl.load_workbook(BytesIO(workspace.export_excel(doc['id'])))['Sheet1']
            self.assertEqual(set(map(str, saved.merged_cells.ranges)), {'A2:A5', 'I2:J6'})
            self.assertEqual(saved['B5'].value, '0X40021004')
            self.assertEqual(saved['A7'].value, 'PT104_TEST')

    def test_create_pattern_in_merged_workbook_preserves_existing_steps(self):
        from gui.swd_workspace import SwdWorkspace
        book = openpyxl.load_workbook(BytesIO(source()))
        sheet = book['Sheet1']
        sheet.merge_cells('A2:A3')
        sheet.merge_cells('I2:J3')
        sheet['I2'] = 'merged note'
        sheet['B3'].fill = openpyxl.styles.PatternFill('solid', fgColor='00FF00')
        raw = BytesIO(); book.save(raw)
        with TemporaryDirectory() as temp:
            workspace = SwdWorkspace(SimpleNamespace(output=Path(temp), checked=None, editor=None))
            doc = workspace.import_excel(raw.getvalue(), 'merged.xlsx')['documents'][0]
            state = create_saved_pattern(workspace, {'documentId': doc['id'], 'name': 'PT104_TEST'})
            saved = openpyxl.load_workbook(BytesIO(workspace.export_excel(doc['id'])))['Sheet1']
            self.assertEqual(set(map(str, saved.merged_cells.ranges)), {'A2:A3', 'I2:J3'})
            self.assertEqual(saved['B3'].value, '0X40021004')
            self.assertEqual(saved['B3'].fill.fgColor.rgb, '0000FF00')
            self.assertEqual(saved['I2'].value, 'merged note')
            self.assertEqual(saved['A4'].value, 'PT104_TEST')
            self.assertEqual(state['documents'][0]['patterns'], ['PT104_demo', 'PT104_TEST'])
            rows = state['documents'][0]['rows']
            rows[1]['__blank_10'] = 'edited covered cell'
            rows[1]['data'] = '0X00000003'
            workspace.save({'documentId': doc['id'], 'rows': rows})
            saved = openpyxl.load_workbook(BytesIO(workspace.export_excel(doc['id'])))['Sheet1']
            self.assertEqual(set(map(str, saved.merged_cells.ranges)), {'A2:A3'})
            self.assertEqual(saved['J3'].value, 'edited covered cell')
            self.assertEqual(saved['I2'].value, 'merged note')
            self.assertEqual(saved['C3'].value, '0X00000003')
            rows = workspace.state()['documents'][0]['rows']
            rows.insert(0, rows.pop())
            workspace.save({'documentId': doc['id'], 'rows': rows})
            saved = openpyxl.load_workbook(BytesIO(workspace.export_excel(doc['id'])))['Sheet1']
            self.assertEqual(set(map(str, saved.merged_cells.ranges)), {'A3:A4'})
            self.assertEqual(saved['A2'].value, 'PT104_TEST')
            self.assertEqual(saved['C4'].value, '0X00000003')

    def test_create_first_pattern_can_be_saved_reloaded_and_generated_for_both_platforms(self):
        from gui.swd_workspace import SwdWorkspace
        with TemporaryDirectory() as temp:
            session = SimpleNamespace(output=Path(temp), checked=None, editor=None)
            workspace = SwdWorkspace(session)
            state = create_saved_pattern(workspace, {'name': ' PT 104_first '})
            self.assertEqual(state['createdDocumentId'], state['documents'][0]['id'])
            self.assertEqual(state['createdPattern'], 'PT104_first')
            doc = state['documents'][0]
            self.assertEqual(doc['patterns'], ['PT104_first'])
            rows = doc['rows'] + [
                {'row': None, 'pattern': None, 'addr': '0X40021000', 'data': '0X00000001',
                 'step': '/', 'reset value': '/', 'type': 'W', 'continue': 'Y', 'description': 'write'},
                {'row': None, 'pattern': None, 'addr': '0X40021000', 'data': '0X00000001',
                 'step': '/', 'reset value': '/', 'type': 'R', 'continue': '', 'description': 'read'},
            ]
            workspace.save({'documentId': doc['id'], 'rows': rows})
            restored = SwdWorkspace(session)
            book = openpyxl.load_workbook(BytesIO(restored.export_excel(doc['id'])))
            self.assertEqual([cell.value for cell in book['Sheet1'][1]],
                             ['pattern', 'addr', 'data', 'step', 'reset value', 'type', 'continue', 'description'])
            self.assertEqual(book['Sheet1']['A2'].value, 'PT104_first')
            state = restored.generate({'documentId': doc['id'], 'pattern': 'PT104_first'})
            self.assertEqual({entry['platform'] for entry in state['generations'][0]['files']}, {'chroma', 's100'})

    def test_create_pattern_appends_to_existing_document_without_changing_original_or_style(self):
        from gui.swd_workspace import SwdWorkspace
        with TemporaryDirectory() as temp:
            styled = openpyxl.load_workbook(BytesIO(source()))
            styled['Sheet1']['A3'] = 'PT104_other'
            styled['Sheet1']['B3'].fill = openpyxl.styles.PatternFill('solid', fgColor='00FF00')
            raw = BytesIO(); styled.save(raw)
            workspace = SwdWorkspace(SimpleNamespace(output=Path(temp), checked=None, editor=None))
            doc = workspace.import_excel(raw.getvalue(), 'steps.xlsx')['documents'][0]
            source_path = Path(workspace.data['documents'][0]['source'])
            original = source_path.read_bytes()
            state = create_saved_pattern(workspace, {'documentId': doc['id'], 'name': 'PT104_second'})
            self.assertEqual(state['createdDocumentId'], doc['id'])
            self.assertEqual(len(state['documents']), 1)
            self.assertEqual(state['documents'][0]['patterns'], ['PT104_demo', 'PT104_other', 'PT104_second'])
            self.assertEqual(source_path.read_bytes(), original)
            book = openpyxl.load_workbook(BytesIO(workspace.export_excel(doc['id'])))
            self.assertEqual(book['Sheet1']['A3'].value, 'PT104_other')
            self.assertEqual(book['Sheet1']['B3'].value, '0X40021004')
            self.assertEqual(book['Sheet1']['B3'].fill.fgColor.rgb, '0000FF00')
            self.assertEqual(book['Sheet1']['A4'].value, 'PT104_second')
            self.assertIsNone(book['Sheet1']['B4'].value)

    def test_create_pattern_rejects_invalid_or_duplicate_name_without_writing_state(self):
        from gui.swd_workspace import SwdWorkspace
        with TemporaryDirectory() as temp:
            workspace = SwdWorkspace(SimpleNamespace(output=Path(temp), checked=None, editor=None))
            state = create_saved_pattern(workspace, {'name': 'PT104_unique'})
            doc_id = state['createdDocumentId']
            before = workspace.state()
            for payload in (
                {'documentId': doc_id, 'name': ' PT104_unique '},
                {'documentId': doc_id, 'name': '9invalid'},
                {'documentId': doc_id, 'name': None},
            ):
                with self.assertRaises(ValueError):
                    create_saved_pattern(workspace, payload)
                self.assertEqual(workspace.state(), before)

    def test_import_material_passes_an_explicit_kind_to_the_library(self):
        from gui.swd_workspace import SwdWorkspace
        with TemporaryDirectory() as temp:
            workspace = SwdWorkspace(SimpleNamespace(output=Path(temp), checked=None, editor=None))
            state = workspace.import_material(b'#define RCC_CSR 0x40021034UL\n', 'cached-source.bin', kind='h')
            self.assertEqual(state['materials']['materials'][0]['kind'], 'h')

    def test_repeated_save_preserves_new_physical_row_identity_and_exports(self):
        from gui.swd_workspace import SwdWorkspace
        with TemporaryDirectory() as temp:
            session = SimpleNamespace(output=Path(temp), checked=None, editor=None)
            workspace = SwdWorkspace(session)
            state = workspace.import_excel(source(), 'steps.xlsx')
            doc = state['documents'][0]
            rows = doc['rows'][1:]
            rows[0]['pattern'] = 'PT104_demo'
            state = workspace.save({'documentId': doc['id'], 'rows': rows})
            rows = state['documents'][0]['rows']
            rows[0]['data'] = '0X00000003'
            workspace.save({'documentId': doc['id'], 'rows': rows})
            restored = SwdWorkspace(session)
            self.assertEqual(restored.state()['documents'][0]['rows'][0]['data'], '0X00000003')
            book = openpyxl.load_workbook(BytesIO(restored.export_excel(doc['id'])))
            self.assertEqual(book['Sheet1']['C2'].value, '0X00000003')

    def test_pin_change_marks_previous_generation_stale_without_deleting_it(self):
        from gui.swd_workspace import SwdWorkspace
        with TemporaryDirectory() as temp:
            workspace = SwdWorkspace(SimpleNamespace(output=Path(temp), checked=None, editor=None))
            doc = workspace.import_excel(source(), 'steps.xlsx')['documents'][0]
            state = workspace.generate({'documentId': doc['id'], 'pattern': 'PT104_demo'})
            gen = state['generations'][0]
            self.assertFalse(gen['stale'])
            self.assertEqual({f['platform'] for f in gen['files']}, {'chroma', 's100'})
            config = state['config']; config['pins']['swc'] = 'PA11'
            state = workspace.configure(config)
            self.assertTrue(state['generations'][0]['stale'])
            self.assertTrue(workspace.output(gen['id'], gen['files'][0]['path']).is_file())
            with self.assertRaises(ValueError):
                workspace.output(gen['id'], '../../workspace.json')

    def test_invalid_document_and_duplicate_pins_are_rejected(self):
        from gui.swd_workspace import SwdWorkspace
        with TemporaryDirectory() as temp:
            workspace = SwdWorkspace(SimpleNamespace(output=Path(temp), checked=None, editor=None))
            with self.assertRaisesRegex(ValueError, '不存在'):
                workspace.export_excel('missing')
            with self.assertRaisesRegex(ValueError, '重复'):
                workspace.configure({'chip': 'PT104', 'pins': {'swc': 'PA13', 'swd': 'PA13', 'reset': 'NRST', 'trig': 'TRIGPIN'}})

    def test_identical_material_reimport_keeps_generation_current(self):
        from gui.swd_workspace import SwdWorkspace
        with TemporaryDirectory() as temp:
            workspace = SwdWorkspace(SimpleNamespace(output=Path(temp), checked=None, editor=None))
            raw = b'#define ENABLE 0x01\n'
            workspace.import_material(raw, 'chip.h')
            doc = workspace.import_excel(source(), 'steps.xlsx')['documents'][0]
            workspace.generate({'documentId': doc['id'], 'pattern': 'PT104_demo'})
            self.assertFalse(workspace.import_material(raw, 'copy.h')['generations'][0]['stale'])
            self.assertTrue(workspace.import_material(raw + b'#define NEW 2\n', 'chip.h')['generations'][0]['stale'])

    def test_link_freezes_both_platforms_and_requires_explicit_confirmation_after_change(self):
        from gui.swd_workspace import SwdWorkspace
        with TemporaryDirectory() as temp:
            editor = SimpleNamespace(sha256='source-sha', raw=b'program', name='program.xlsx')
            slot = {'phase': 'CONFIG', 'role': 'config', 'status': 'referenced',
                    'names': {'shared': ['PT104_SWD_VOL']}}
            session = SimpleNamespace(output=Path(temp), checked=None, editor=editor,
                scan=SimpleNamespace(candidate=lambda draft: (None, b'program')),
                links=[{'flowRow': 10, 'cp': 'CP3', 'item': 'VOL', 'resourceSlots': [slot]}])
            workspace = SwdWorkspace(session)
            doc = workspace.import_excel(source(), 'steps.xlsx')['documents'][0]
            gen = workspace.generate({'documentId': doc['id'], 'pattern': 'PT104_demo'})['generations'][0]
            target = workspace.targets()['targets'][0]
            workspace.link({'targetId': target['id'], 'generationId': gen['id']})
            frozen = workspace.freeze(cps={'CP3'})
            self.assertEqual(set(frozen[0]['contents']), {'chroma', 's100'})
            config = workspace.state()['config']; config['pins']['swc'] = 'PA11'
            workspace.configure(config)
            with self.assertRaisesRegex(ValueError, '重新生成'):
                workspace.freeze(cps={'CP3'})
            new = workspace.generate({'documentId': doc['id'], 'pattern': 'PT104_demo'})['generations'][-1]
            with self.assertRaises(ValueError):
                workspace.freeze(cps={'CP3'})
            workspace.link({'targetId': target['id'], 'generationId': new['id']})
            self.assertNotEqual(frozen[0]['identity'], workspace.freeze()[0]['identity'])
            workspace.persist('saved-sha')
            session.editor.sha256 = 'saved-sha'
            self.assertEqual(len(SwdWorkspace(session).freeze()), 1)
            session.links.clear()
            with self.assertRaisesRegex(ValueError, '引用已变化'):
                workspace.freeze()


if __name__ == '__main__':
    unittest.main()
