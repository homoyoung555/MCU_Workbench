from __future__ import annotations

import tempfile
from time import perf_counter
from pathlib import Path
import unittest


WGL = '''{ Begin_async_clocks }
{ PAD_PA[1] = period: 100ns }
{ PAD_PB[0] = period: 200ns }
{ End_async_clocks }
waveform chip_top
signal
    "PAD_PA"[1..0] : input initialp[N];
    "PAD_PB"[0] : output initialp[N];
end
timeplate "_WFT_" period 40ns
    "PAD_PA"[1] := input [0ns:D, 10ns:S, 20ns:D];
    "PAD_PA"[0] := input [0ns:D, 10ns:S, 20ns:D];
    "PAD_PA"[1] := output [20ns:Q, 30ns:D];
    "PAD_PA"[0] := output [20ns:Q, 30ns:D];
    "PAD_PB"[0] := output [16ns:Q, 32ns:D];
end
pattern scan_demo("PAD_PA"[1] :I, "PAD_PA"[0] :I, "PAD_PB"[0] :O)
    vector(+, "_WFT_") := [010];
end
'''


class _Workbook:
    def __init__(self, program_name="ignored"):
        self.program_name = program_name

    def dataframe(self, name):
        import pandas as pd
        if name == "PIN DEC":
            return pd.DataFrame([
                {"E/C(Pad Name)": "PAD_PA[0]"},
                {"E/C(Pad Name)": "PAD_PA[1]"},
            ])
        if name == "VERSION":
            return pd.DataFrame([{"NRST PIN": "PAD_PF[2]", "Program Name": self.program_name}])
        raise KeyError(name)


class _MultiCpWorkbook(_Workbook):
    def dataframe(self, name):
        import pandas as pd
        if name == "PIN DEC":
            return pd.DataFrame([
                {"CP NO": "CP1", "E/C(Pad Name)": "PAD_PA[1]"},
                {"CP NO": "CP2", "E/C(Pad Name)": "PAD_PA[0]"},
            ])
        if name == "VERSION":
            return pd.DataFrame([{"CP num": "CP1", "NRST PIN": "PAD_PF[2]"}])
        raise KeyError(name)


def _reader_workbook(path: Path):
    from openpyxl import Workbook
    from program_gen.workbook import WorkbookReader

    book = Workbook()
    pins = book.active
    pins.title = "PIN DEC"
    pins.append(["E/C(Pad Name)"])
    pins.append(["PAD_PA[0]"])
    pins.append(["PAD_PA[1]"])
    version = book.create_sheet("VERSION")
    version.append(["NRST PIN", "Program Name"])
    version.append(["PAD_PF[2]", "ignored"])
    book.save(path)
    return WorkbookReader().read(path)


class ScanConversionTests(unittest.TestCase):
    def test_labels_reuse_original_markers_with_pin_filtering(self):
        import re
        from gui.scan_conversion import convert_scan
        marker = '{ Measure Iddq All }\n'
        content = WGL.replace('    vector(+, "_WFT_") := [010];', marker + '    vector(+, "_WFT_") := [010];\n' + marker + '    vector(+, "_WFT_") := [110];')
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root/'labels.wgl'
            source.write_text(content, encoding='utf-8')
            for mode, options, expected in [('manual', {'manual_pins':['PA0']}, ['1', '1']), ('excel', {'workbook':_Workbook()}, ['10','11'])]:
                out = root/mode
                convert_scan([source], out, ['chroma', 's100'], label_marker='Measure Iddq All', label_base='M_Iddq', **options)
                pat = (out/'chroma/labels.pat').read_text(encoding='utf-8')
                sd = (out/'s100/labels.sd').read_text(encoding='utf-8')
                for index, data in enumerate(expected, 1):
                    self.assertRegex(pat, rf'M_Iddq_{index}:\n\s*\*{data}\*')
                    self.assertRegex(sd, rf'M_Iddq_{index}:\s*{data};')
                self.assertIn('TIMING_SET G2_1;', sd)
                self.assertNotIn('TIMING_SET WFT;', sd)
            for marker, base in [('text', None), (None, 'name'), ('text', 'bad-name'), (1, 'name')]:
                with self.subTest(marker=marker, base=base), self.assertRaises(ValueError):
                    convert_scan([source], root/'bad', ['chroma'], label_marker=marker, label_base=base)

    def test_excel_reset_alias_keeps_vector_timing_and_clock(self):
        from gui.scan_conversion import convert_scan
        class ResetWorkbook(_Workbook):
            def dataframe(self, name):
                frame = super().dataframe(name)
                if name == 'VERSION':
                    frame['NRST PIN'] = 'PAD_PA[1]'
                return frame
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root/'reset.wgl'
            source.write_text(WGL, encoding='utf-8')
            convert_scan([source], root/'out', ['chroma', 's100'], workbook=ResetWorkbook())
            pat = (root/'out/chroma/reset.pat').read_text(encoding='utf-8')
            self.assertIn('PA0, NRST;', pat)
            self.assertIn('*10*', pat)
            sd = (root/'out/s100/reset.sd').read_text(encoding='utf-8')
            self.assertIn('\tPA0,\n\tNRST;', sd)
            self.assertIn('START:10;', sd)
            for name in ('chroma/reset.scan_timing.inc', 'chroma/reset.scan.inc', 's100/reset.scan.timing', 's100/reset.scan.prg'):
                self.assertIn('NRST', (root/'out'/name).read_text(encoding='utf-8'))

    def test_manual_pins_preserve_source_order_and_filter_all_artifacts(self):
        from gui.scan_conversion import convert_scan, read_wgl_pins
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            first, second = root/'first.wgl', root/'second.wgl'
            first.write_text(WGL, encoding='utf-8')
            second.write_text(WGL.replace('[1..0] : input', '[0..1] : input'), encoding='utf-8')
            self.assertEqual(read_wgl_pins(first), ['PA1', 'PA0', 'PB0'])
            result = convert_scan([first, second], root/'out', ['chroma', 's100'], manual_pins=['PA0', 'PA1'])
            self.assertEqual(result['mode'], 'manual-pins')
            for stem, header, vector in [('first', 'PA1, PA0;', '01'), ('second', 'PA0, PA1;', '10')]:
                pat = (root/f'out/chroma/{stem}.pat').read_text(encoding='utf-8')
                self.assertIn(header, pat)
                self.assertIn('*' + vector + '*', pat)
                sd = (root/f'out/s100/{stem}.sd').read_text(encoding='utf-8')
                self.assertIn('START:' + vector + ';', sd)
            convert_scan([first], root/'single', ['chroma', 's100'], manual_pins=['PA0'])
            for path in (root/'single').rglob('*'):
                if path.is_file():
                    text = path.read_text(encoding='utf-8')
                    self.assertNotIn('PA1', text)
                    self.assertNotIn('PB0', text)
            self.assertEqual(first.read_text(encoding='utf-8'), WGL)

    def test_manual_pin_selection_validation(self):
        from gui.scan_conversion import convert_scan
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root/'first.wgl'
            source.write_text(WGL, encoding='utf-8')
            for selection in ([], ['missing'], ['PA0', 'pa0'], 'PA0', [1]):
                with self.subTest(selection=selection), self.assertRaises(ValueError):
                    convert_scan([source], root/'out', ['chroma'], manual_pins=selection)
            with self.assertRaises(ValueError):
                convert_scan([source], root/'out', ['chroma'], workbook=_Workbook(), manual_pins=['PA0'])
            other = root/'other.wgl'
            other.write_text(WGL.replace('PAD_PA', 'PAD_PC').replace('PAD_PB', 'PAD_PD'), encoding='utf-8')
            with self.assertRaisesRegex(ValueError, 'other.wgl'):
                convert_scan([source, other], root/'out', ['chroma'], manual_pins=['PA0'])

    def test_chroma_uses_ts8_through_ts15_and_rejects_a_ninth_group(self):
        import re
        from gui.scan_conversion import convert_scan
        from program_gen.platforms.chroma.scan_writer import _GROUP_TS, _GROUP_TS_VAR
        template = WGL[WGL.index('timeplate '):WGL.index('pattern scan_demo')]
        for count in (3, 8, 9):
            with self.subTest(count=count), tempfile.TemporaryDirectory() as temp:
                root = Path(temp)
                source = root/'scan_demo.wgl'
                content = WGL[:WGL.index('timeplate ')]
                content += ''.join(template.replace('_WFT_', f'WFT{i}').replace('period 40ns', f'period {40+i*10}ns') for i in range(count))
                content += 'pattern scan_demo("PAD_PA"[1] :I, "PAD_PA"[0] :I, "PAD_PB"[0] :O)\n'
                content += ''.join(f'    vector(+, "WFT{i}") := [010];\n' for i in range(count)) + 'end\n'
                source.write_text(content, encoding='utf-8')
                if count > 8:
                    with self.assertRaisesRegex(ValueError, 'CHROMA_TIMESET_LIMIT.*9'):
                        convert_scan([str(source)], str(root/'out'), ['chroma'])
                else:
                    convert_scan([str(source)], str(root/'out'), ['chroma'])
                    expected = {f'TS{i}' for i in range(8, 8+count)}
                    pattern = (root/'out/chroma/scan_demo.pat').read_text(encoding='utf-8')
                    timing = (root/'out/chroma/scan_demo.scan_timing.inc').read_text(encoding='utf-8')
                    self.assertEqual(set(re.findall(r'\bTS\d+\b', pattern)), expected)
                    self.assertEqual(set(re.findall(r'RATEN\((TS\d+)', timing)), expected)
                    for name in expected:
                        self.assertIn(name, _GROUP_TS)
                        self.assertIn(_GROUP_TS_VAR[name], timing)

    def test_converts_full_pin_for_both_platforms(self):
        from gui.scan_conversion import convert_scan

        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "scan_demo.wgl"
            output = root / "out"
            source.write_text(WGL, encoding="utf-8")

            result = convert_scan([str(source.resolve())], str(output.resolve()), ["chroma", "s100"])

            self.assertEqual({item["platform"] for item in result["files"]}, {"chroma", "s100"})
            self.assertTrue(all(not Path(item["path"]).is_absolute() for item in result["files"]))
            self.assertTrue(all(item["sourceName"] == "scan_demo.wgl" for item in result["files"]))
            self.assertIn("PA1, PA0, PB0", (output / "chroma" / "scan_demo.pat").read_text(encoding="utf-8"))
            self.assertIn("TIMING_SET G2_1", (output / "s100" / "scan_demo.sd").read_text(encoding="utf-8"))
            self.assertIn("RATEN(TS8", (output / "chroma" / "scan_demo.scan_timing.inc").read_text(encoding="utf-8"))
            self.assertIn("@@AC_SET_1", (output / "s100" / "scan_demo.scan.timing").read_text(encoding="utf-8"))
            self.assertIn("PB0", (output / "chroma" / "scan_demo.scan_timing.inc").read_text(encoding="utf-8"))
            self.assertIn("fsclk_scan8*0.400", (output / "chroma" / "scan_demo.scan_timing.inc").read_text(encoding="utf-8"))
            self.assertIn("PB0: {STB_RS=16 }", (output / "s100" / "scan_demo.scan.timing").read_text(encoding="utf-8"))

    def test_adapts_pins_in_pin_dec_order_and_fingerprints_only_relevant_configuration(self):
        from gui.scan_conversion import configuration_fingerprint, convert_scan

        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "scan_demo.wgl"
            output = root / "out"
            source.write_text(WGL, encoding="utf-8")
            workbook = _reader_workbook(root / "program.xlsx")

            result = convert_scan([str(source.resolve())], str(output.resolve()), ["chroma", "s100"], workbook=workbook)

            self.assertEqual(result["configurationFingerprint"], configuration_fingerprint(workbook))
            self.assertEqual(configuration_fingerprint(workbook), configuration_fingerprint(_Workbook("changed elsewhere")))
            chroma = (output / "chroma" / "scan_demo.pat").read_text(encoding="utf-8")
            self.assertIn("PA0, PA1;", chroma)
            self.assertNotIn("PB0", chroma)
            self.assertIn("*10*", chroma)
            s100 = (output / "s100" / "scan_demo.sd").read_text(encoding="utf-8")
            self.assertIn("\tPA0,\n\tPA1;", s100)
            self.assertIn("START:10;TIMING_SET G2_1", s100)

    def test_requires_cp_for_multi_cp_pin_dec(self):
        from gui.scan_conversion import convert_scan

        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "scan_demo.wgl"
            source.write_text(WGL, encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "SCAN.PIN_DEC_CP_REQUIRED"):
                convert_scan([str(source.resolve())], str((root / "out").resolve()), ["chroma"], _MultiCpWorkbook())
            result = convert_scan([str(source.resolve())], str((root / "out").resolve()), ["chroma"], _MultiCpWorkbook(), cp="CP2")
            self.assertEqual(result["configurationCp"], "CP2")
            self.assertIn("PA0;", (root / "out" / "chroma" / "scan_demo.pat").read_text(encoding="utf-8"))

    def test_s100_uses_batch_wide_timeset_numbers_and_rejects_illegal_logical_names(self):
        from gui.scan_conversion import convert_scan

        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            first = root / "first.wgl"
            second = root / "second.wgl"
            first.write_text(WGL.replace("_WFT_", "_FIRST_WFT_"), encoding="utf-8")
            second.write_text(WGL.replace("_WFT_", "_SECOND_WFT_"), encoding="utf-8")
            output = root / "out"
            convert_scan([str(first.resolve()), str(second.resolve())], str(output.resolve()), ["s100"])
            self.assertIn("TIMING_SET G2_1", (output / "s100" / "first.sd").read_text(encoding="utf-8"))
            self.assertIn("TIMING_SET G2_2", (output / "s100" / "second.sd").read_text(encoding="utf-8"))
            self.assertIn("@@AC_SET_1", (output / "s100" / "first.scan.timing").read_text(encoding="utf-8"))
            self.assertIn("@@AC_SET_2", (output / "s100" / "second.scan.timing").read_text(encoding="utf-8"))
            self.assertNotIn("@@TIMING_DEFINE", (output / "s100" / "first.scan.timing").read_text(encoding="utf-8"))

            invalid = root / "bad name.wgl"
            invalid.write_text(WGL, encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "SCAN.LOGICAL_NAME"):
                convert_scan([str(invalid.resolve())], str((root / "bad").resolve()), ["chroma"])

    def test_async_clocks_are_scoped_to_each_source_and_adapted_pin_set(self):
        from gui.scan_conversion import convert_scan

        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "scan_demo.wgl"
            source.write_text(WGL, encoding="utf-8")
            full = root / "full"
            convert_scan([str(source.resolve())], str(full.resolve()), ["chroma", "s100"])

            chroma_code = (full / "chroma" / "scan_demo.scan.inc").read_text(encoding="utf-8")
            self.assertIn("FRUN_CLK(PA1,100.00nS,0nS,50.00nS);", chroma_code)
            self.assertIn("FRUN_CLK_MODE(PB0,1,0,ON);", chroma_code)
            self.assertLess(chroma_code.index("FRUN_CLK_MODE(PA1,1,0,ON);"), chroma_code.index("JUDGE_PAT"))
            self.assertLess(chroma_code.index("JUDGE_PAT"), chroma_code.index("FRUN_CLK_MODE(PA1,1,0,OFF);"))
            s100_code = (full / "s100" / "scan_demo.scan.prg").read_text(encoding="utf-8")
            self.assertNotIn("@@PLAN", s100_code)
            self.assertIn("LOGIC_SET_FREE_RUN_CLOCK_PIN(PA1,10000000);", s100_code)
            self.assertIn("LOGIC_SET_FREE_RUN_CLOCK_PIN(NULL,5000000);", s100_code)

            adapted = root / "adapted"
            convert_scan([str(source.resolve())], str(adapted.resolve()), ["chroma", "s100"], workbook=_Workbook())
            self.assertNotIn("PB0", (adapted / "chroma" / "scan_demo.scan.inc").read_text(encoding="utf-8"))
            self.assertNotIn("PB0", (adapted / "s100" / "scan_demo.scan.prg").read_text(encoding="utf-8"))

    def test_fingerprint_reads_formal_pin_dec_once_within_polling_budget(self):
        from gui.scan_conversion import configuration_fingerprint
        from program_gen.workbook import WorkbookReader
        from tests.workspace_paths import FORMAL_MASTER

        if not FORMAL_MASTER.is_file():
            self.skipTest("external formal master unavailable")
        model = WorkbookReader().read(FORMAL_MASTER)

        class CountingWorkbook:
            def __init__(self, inner):
                self.inner = inner
                self.calls = []

            def dataframe(self, name):
                self.calls.append(name)
                return self.inner.dataframe(name)

        workbook = CountingWorkbook(model)
        started = perf_counter()
        fingerprint = configuration_fingerprint(workbook, "CP1")
        elapsed = perf_counter() - started

        self.assertEqual(len(fingerprint), 64)
        self.assertEqual(workbook.calls.count("PIN DEC"), 1)
        self.assertEqual(workbook.calls.count("VERSION"), 1)
        self.assertLess(elapsed, 0.25, f"fingerprint took {elapsed:.3f}s")


if __name__ == "__main__":
    unittest.main()
