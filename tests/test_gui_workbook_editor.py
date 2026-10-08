from io import BytesIO
from tests.workspace_paths import FORMAL_MASTER
import unittest
from zipfile import ZipFile

import openpyxl
from gui.workbook_editor import WorkbookEditor, _patch_sheet, formula_dependencies, NS
import xml.etree.ElementTree as ET


class MergeLayoutTests(unittest.TestCase):
    def test_merge_bounds_and_anchor_survive_edit_roundtrip(self):
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = 'FB MAP'
        sheet['B2'] = 'field'
        sheet.merge_cells('B2:D3')
        stream = BytesIO()
        book.save(stream)
        book.close()
        editor = WorkbookEditor(stream.getvalue(), 'merged.xlsx')
        info = editor.by_name['FB MAP']
        self.assertEqual(info['merges'], [{'minRow': 2, 'maxRow': 3, 'minCol': 2, 'maxCol': 4}])
        self.assertTrue(info['rows'][1][1]['editable'])
        self.assertFalse(info['rows'][2][2]['editable'])
        raw = editor.apply(editor.sha256, [{'sheet': 'FB MAP', 'cell': 'B2', 'old': 'field', 'value': 'updated'}])
        reopened = openpyxl.load_workbook(BytesIO(raw))
        self.assertEqual(str(next(iter(reopened['FB MAP'].merged_cells.ranges))), 'B2:D3')
        self.assertEqual(reopened['FB MAP']['B2'].value, 'updated')
        reopened.close()


class WorkbookEditorTests(unittest.TestCase):
    def test_choices_follow_contract_not_misaligned_excel_validation(self):
        from preflight.schema import TYPED, ENUM
        e = self.editor
        for (name, label), spec in TYPED.items():
            if spec['type'] == ENUM:
                col = next(c for c in e.by_name[name]['columns'] if c['label'] == label)
                self.assertEqual(col['choices'], [''] + spec['enum'] + (['/'] if spec.get('slash_ok') else []))
        flow = {c['label']: c for c in e.by_name['FLOW']['columns']}
        self.assertEqual(flow['PUMP']['choices'], ['', 'no', 'yes', '/'])
        for label in ('LINE', 'Vcc', 'Recipe', 'Sbin', 'item'):
            self.assertNotIn('choices', flow[label])
        for name, label in [('DAC', 'mode'), ('OPA', 'mode'), ('PVD', 'expect'), ('ADC', 'ADC'), ('FREQ', 'source')]:
            self.assertNotIn('choices', next(c for c in e.by_name[name]['columns'] if c['label'] == label))
        self.assertEqual(next(c for c in e.by_name['PVD']['columns'] if c['label'] == 'up/down')['choices'], ['', 'up', 'down'])

    @classmethod
    def setUpClass(cls):
        cls.path = FORMAL_MASTER
        cls.editor = WorkbookEditor(cls.path.read_bytes(), cls.path.name, cls.path)

    def test_roundtrip_keeps_formula_caches_and_only_patches_requested_cells(self):
        editor = self.editor
        changes = [{'sheet': 'FLOW', 'cell': 'S4', 'old': 0.8, 'value': '0.801'},
                   {'sheet': 'FLOW', 'cell': 'F4', 'old': None, 'value': 'ENTER-TRIM-JUDGE'},
                   {'sheet': 'DAC', 'cell': 'D2', 'old': editor.by_name['DAC']['rows'][1][3]['value'], 'value': '+0.02'}]
        result = editor.apply(editor.sha256, changes)
        with ZipFile(BytesIO(editor.raw)) as a, ZipFile(BytesIO(result)) as b:
            self.assertEqual(a.namelist(), b.namelist())
            self.assertEqual({p for p in a.namelist() if a.read(p) != b.read(p)},
                             {editor.sheet_paths['FLOW'], editor.sheet_paths['DAC']})
        wb = openpyxl.load_workbook(BytesIO(result), data_only=True)
        try:
            self.assertEqual(wb['FLOW']['S4'].value, .801)
            self.assertEqual(wb['FLOW']['F4'].value, 'ENTER-TRIM-JUDGE')
            self.assertEqual(wb['DAC']['D2'].value, '+0.02')
            self.assertEqual(wb['PIN DEC']['C21'].value, 2214.915)
        finally: wb.close()
        self.assertEqual(self.path.read_bytes(), editor.raw)

    def test_rejects_protected_and_stale_changes_atomically(self):
        e = self.editor
        for change in [{'sheet':'PIN DEC','cell':'C2','old':420.615,'value':'1'},
                       {'sheet':'FLOW','cell':'B1','old':'item','value':'x'},
                       {'sheet':'FLOW','cell':'S4','old':.9,'value':'1'},
                       {'sheet':'FLOW','cell':'S4','old':.8,'value':'=1+1'},
                       {'sheet':'FLOW','cell':'S4','old':.8,'value':'9'*600},
                       {'sheet':'ADC','cell':'A3','old':None,'value':'ADC2'}]:
            with self.subTest(change=change), self.assertRaises(ValueError): e.apply(e.sha256,[change])
        with self.assertRaises(ValueError): e.apply('stale',[])
        self.assertEqual(e.apply(e.sha256,[]),e.raw)

    def test_clear_and_duplicate_coordinate(self):
        e = self.editor
        change={'sheet':'FLOW','cell':'S4','old':.8,'value':''}
        with self.assertRaises(ValueError): e.apply(e.sha256,[change,change])
        wb=openpyxl.load_workbook(BytesIO(e.apply(e.sha256,[change])))
        try: self.assertIsNone(wb['FLOW']['S4'].value)
        finally: wb.close()

    def test_missing_cells_rows_and_default_namespace(self):
        xml='<worksheet xmlns="'+NS[1:-1]+'"><sheetData><row r="2"><c r="A2" s="4" t="n"><v>1</v></c><c r="C2" t="n"><v>3</v></c></row><row r="4"/></sheetData></worksheet>'
        updated=_patch_sheet(xml,{'B2':('inlineStr','a & b'), 'A3':('n','2'), 'A4':('n','4')})
        root=ET.fromstring(updated)
        self.assertEqual([r.attrib['r'] for r in root.find(NS+'sheetData')],['2','3','4'])
        self.assertEqual([c.attrib['r'] for c in root.find(NS+'sheetData')[0]],['A2','B2','C2'])
        self.assertIn('<c r="A2" s="4" t="n"><v>1</v></c>',updated)
        self.assertEqual(root.find(NS+'sheetData')[0][1].find(NS+'is')[0].text,'a & b')

    def test_formula_references_protect_input_sheets(self):
        wb=openpyxl.Workbook(); wb.active.title='Output'; wb.create_sheet('Inputs'); wb.create_sheet('Independent')
        wb['Output']['A1']="='Inputs'!B2+1"
        self.assertEqual(formula_dependencies(wb),{'Output','Inputs'})
        wb['Output']['A2']='=INDIRECT("Independent!A1")'
        self.assertEqual(formula_dependencies(wb),set(wb.sheetnames))
        wb.close()

    def test_append_flow_row_preserves_existing_xml_and_style(self):
        e=self.editor; r=e.by_name['FLOW']['maxRow']+1
        append=[{'sheet':'FLOW','sourceRow':4,'copy':True}]
        changes=[{'sheet':'FLOW','cell':f'B{r}','old':e.by_name['FLOW']['rows'][3][1]['value'],'value':'GUI_Row_Verify'}]
        raw=e.apply(e.sha256,changes,append)
        with ZipFile(BytesIO(e.raw)) as a, ZipFile(BytesIO(raw)) as b:
            flow=e.sheet_paths['FLOW']
            self.assertEqual([n for n in a.namelist() if a.read(n)!=b.read(n)],[flow])
            before=ET.fromstring(a.read(flow)); after=ET.fromstring(b.read(flow))
            old_rows=before.find(NS+'sheetData'); new_rows=after.find(NS+'sheetData')
            self.assertEqual([ET.tostring(x) for x in old_rows],[ET.tostring(x) for x in list(new_rows)[:-1]])
        wb=openpyxl.load_workbook(BytesIO(raw),data_only=True)
        try:
            self.assertEqual(wb['FLOW'].max_row,r)
            self.assertEqual(wb['FLOW'][f'B{r}'].value,'GUI_Row_Verify')
            self.assertEqual(wb['FLOW'][f'S{r}'].value,wb['FLOW']['S4'].value)
            self.assertEqual(wb['FLOW'][f'B{r}'].style_id,wb['FLOW']['B4'].style_id)
            self.assertEqual(wb['PIN DEC']['C21'].value,2214.915)
        finally:wb.close()
        self.assertEqual(e.apply(e.sha256,[],[]),e.raw)

    def test_blank_append_and_structural_rejections(self):
        e=self.editor; r=e.by_name['FLOW']['maxRow']+1
        raw=e.apply(e.sha256,[{'sheet':'FLOW','cell':f'B{r}','old':None,'value':'NEW'}],
                    [{'sheet':'FLOW','sourceRow':4,'copy':False}])
        wb=openpyxl.load_workbook(BytesIO(raw),data_only=True)
        try:
            self.assertEqual(wb['FLOW'][f'B{r}'].value,'NEW');self.assertIsNone(wb['FLOW'][f'S{r}'].value)
        finally:wb.close()
        for operations in [[{'sheet':'PIN DEC','sourceRow':4,'copy':True}],
                           [{'sheet':'FLOW','sourceRow':1,'copy':True}],
                           [{'sheet':'FLOW','sourceRow':r,'copy':True}],
                           [{'sheet':'FLOW','sourceRow':4,'copy':'false'}],
                           [{'sheet':'FLOW','sourceRow':4,'copy':True}]*101]:
            with self.subTest(operations=operations[:1]),self.assertRaises(ValueError):e.apply(e.sha256,[],operations)


if __name__ == '__main__': unittest.main()
