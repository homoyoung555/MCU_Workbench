"""Canonical, immutable workbook input model for the V1.3 generator.

The source file is read exactly once.  WorkbookReader builds one snapshot that
contains cached values, formula text, cell provenance, formatting markers,
merged ranges, row/column grouping metadata, and pandas-compatible table data.
Preflight and the temporary legacy-state adapter consume this same model; they
must not reopen the workbook.

The worksheet facade intentionally implements only the read-only subset used by
the existing generator.  It is a migration adapter, not the public Product
Compiler interface.
"""

from __future__ import annotations

from dataclasses import asdict, dataclass
from datetime import date, datetime, time
from decimal import Decimal
from io import BytesIO
import hashlib
import json
import math
from pathlib import Path
from typing import Callable, Iterable, Iterator

import openpyxl
from openpyxl.utils import column_index_from_string, coordinate_to_tuple, get_column_letter
import pandas as pd

from program_gen.core.project_mappings import read_storage_mappings, read_verify_storage_mappings
from program_gen.core.pattern_links import read_pattern_links


@dataclass(frozen=True)
class ColorModel:
    type: str | None = None
    rgb: str | None = None
    indexed: int | None = None
    theme: int | None = None
    tint: float = 0.0


@dataclass(frozen=True)
class FillModel:
    fill_type: str | None
    fgColor: ColorModel
    bgColor: ColorModel

    @property
    def fg_color(self) -> ColorModel:
        return self.fgColor

    @property
    def bg_color(self) -> ColorModel:
        return self.bgColor


@dataclass(frozen=True)
class FontModel:
    name: str | None
    size: float | None
    bold: bool
    italic: bool
    underline: str | None
    strike: bool
    color: ColorModel


@dataclass(frozen=True)
class AlignmentModel:
    horizontal: str | None
    vertical: str | None
    wrap_text: bool | None
    shrink_to_fit: bool | None
    indent: float
    text_rotation: int


@dataclass(frozen=True)
class CellModel:
    sheet: str
    row: int
    column: int
    coordinate: str
    value: object
    formula: str | None
    cached_value: object
    data_type: str | None
    number_format: str | None
    style_id: int
    fill: FillModel
    font: FontModel
    alignment: AlignmentModel

    @property
    def col_idx(self) -> int:
        return self.column

    @property
    def column_letter(self) -> str:
        return get_column_letter(self.column)

    @property
    def text(self) -> str:
        return "" if self.value is None else str(self.value)

    @property
    def is_blank(self) -> bool:
        return self.value is None or (
            isinstance(self.value, str) and self.value.strip() == ""
        )


@dataclass(frozen=True)
class MergedRangeModel:
    coordinate: str
    min_row: int
    max_row: int
    min_col: int
    max_col: int

    def __str__(self) -> str:
        return self.coordinate


@dataclass(frozen=True)
class MergedCellsModel:
    ranges: tuple[MergedRangeModel, ...]


@dataclass(frozen=True)
class RowDimensionModel:
    index: int
    hidden: bool
    outline_level: int
    collapsed: bool
    height: float | None


@dataclass(frozen=True)
class ColumnDimensionModel:
    key: str
    min_column: int
    max_column: int
    hidden: bool
    outline_level: int
    collapsed: bool
    width: float | None


@dataclass(frozen=True)
class SheetModel:
    title: str
    state: str
    max_row: int
    max_column: int
    rows: tuple[tuple[CellModel, ...], ...]
    merged_ranges: tuple[MergedRangeModel, ...]
    row_dimensions: tuple[RowDimensionModel, ...]
    column_dimensions: tuple[ColumnDimensionModel, ...]

    @property
    def merged_cells(self) -> MergedCellsModel:
        return MergedCellsModel(self.merged_ranges)

    def cell(self, row: int, column: int) -> CellModel:
        if row < 1 or column < 1:
            raise ValueError("row and column must be 1-based positive integers")
        if row <= self.max_row and column <= self.max_column:
            return self.rows[row - 1][column - 1]
        return _blank_cell(self.title, row, column)

    def iter_rows(
        self,
        min_row: int = 1,
        max_row: int | None = None,
        min_col: int = 1,
        max_col: int | None = None,
        values_only: bool = False,
    ) -> Iterator[tuple[object, ...]]:
        last_row = self.max_row if max_row is None else max_row
        last_col = self.max_column if max_col is None else max_col
        for row in range(min_row, last_row + 1):
            cells = tuple(self.cell(row, column) for column in range(min_col, last_col + 1))
            yield tuple(cell.value for cell in cells) if values_only else cells

    def __getitem__(self, key: int | str):
        if isinstance(key, int):
            return tuple(self.cell(key, column) for column in range(1, self.max_column + 1))
        if not isinstance(key, str):
            raise TypeError(f"unsupported worksheet key: {type(key).__name__}")
        token = key.strip().upper()
        if token.isalpha():
            column = column_index_from_string(token)
            return tuple(self.cell(row, column) for row in range(1, self.max_row + 1))
        row, column = coordinate_to_tuple(token)
        return self.cell(row, column)


@dataclass(frozen=True)
class TableModel:
    title: str
    columns: tuple[object, ...]
    rows: tuple[tuple[object, ...], ...]

    @classmethod
    def from_dataframe(cls, title: str, frame: pd.DataFrame) -> "TableModel":
        return cls(
            title=title,
            columns=tuple(frame.columns.tolist()),
            rows=tuple(tuple(row) for row in frame.itertuples(index=False, name=None)),
        )

    def dataframe(self) -> pd.DataFrame:
        return pd.DataFrame(
            [list(row) for row in self.rows],
            columns=list(self.columns),
            dtype=object,
        )


@dataclass(frozen=True)
class WorkbookModel:
    source_path: str
    source_size: int
    source_sha256: str
    sheets: tuple[SheetModel, ...]
    tables: tuple[TableModel, ...]
    identity: str
    trim_storage_mappings: tuple = ()
    verify_storage_mappings: tuple = ()
    pattern_links: tuple[dict, ...] = ()

    @property
    def sheetnames(self) -> list[str]:
        return [sheet.title for sheet in self.sheets]

    @property
    def worksheets(self) -> tuple[SheetModel, ...]:
        return self.sheets

    def has_sheet(self, title: str) -> bool:
        return any(sheet.title == title for sheet in self.sheets)

    def sheet(self, title: str) -> SheetModel:
        for sheet in self.sheets:
            if sheet.title == title:
                return sheet
        raise KeyError(f"缺少 sheet: {title}")

    def dataframe(self, title: str) -> pd.DataFrame:
        for table in self.tables:
            if table.title == title:
                return table.dataframe()
        raise KeyError(f"缺少 pandas table: {title}")

    def __getitem__(self, title: str) -> SheetModel:
        return self.sheet(title)

    def close(self) -> None:
        """Compatibility no-op: the immutable model owns no open file handle."""


class WorkbookReader:
    """Read one workbook file into one canonical immutable snapshot."""

    def __init__(self, byte_loader: Callable[[Path], bytes] | None = None):
        self._byte_loader = byte_loader or (lambda path: path.read_bytes())

    def read(self, source: str | Path) -> WorkbookModel:
        path = Path(source).resolve()
        raw = self._byte_loader(path)
        source_sha = hashlib.sha256(raw).hexdigest().upper()
        keep_vba = path.suffix.casefold() == ".xlsm"
        cached_wb = openpyxl.load_workbook(
            BytesIO(raw), read_only=False, data_only=True, keep_vba=keep_vba
        )
        formula_wb = openpyxl.load_workbook(
            BytesIO(raw), read_only=False, data_only=False, keep_vba=keep_vba
        )
        try:
            if cached_wb.sheetnames != formula_wb.sheetnames:
                raise ValueError("formula/cache workbook sheet order mismatch")
            frames = pd.read_excel(BytesIO(raw), sheet_name=None, dtype=object)
            if list(frames) != cached_wb.sheetnames:
                raise ValueError("pandas/openpyxl workbook sheet order mismatch")
            sheets = tuple(
                _snapshot_sheet(cached_wb[title], formula_wb[title])
                for title in cached_wb.sheetnames
            )
            tables = tuple(
                TableModel.from_dataframe(title, frames[title])
                for title in cached_wb.sheetnames
            )
            trim_storage_mappings = read_storage_mappings(raw)
            verify_storage_mappings = read_verify_storage_mappings(raw)
            pattern_links = tuple(read_pattern_links(raw))
            identity = _workbook_identity(
                source_sha, sheets, tables, trim_storage_mappings, verify_storage_mappings, pattern_links
            )
            return WorkbookModel(
                source_path=str(path),
                source_size=len(raw),
                source_sha256=source_sha,
                sheets=sheets,
                tables=tables,
                identity=identity,
                trim_storage_mappings=trim_storage_mappings,
                verify_storage_mappings=verify_storage_mappings,
                pattern_links=pattern_links,
            )
        finally:
            cached_wb.close()
            formula_wb.close()


def _color_model(color) -> ColorModel:
    if color is None:
        return ColorModel()
    color_type = getattr(color, "type", None)
    return ColorModel(
        type=color_type,
        rgb=getattr(color, "rgb", None) if color_type == "rgb" else None,
        indexed=getattr(color, "indexed", None) if color_type == "indexed" else None,
        theme=getattr(color, "theme", None) if color_type == "theme" else None,
        tint=float(getattr(color, "tint", 0.0) or 0.0),
    )


def _cell_model(sheet: str, cached_cell, formula_cell) -> CellModel:
    formula = formula_cell.value if formula_cell.data_type == "f" else None
    font = cached_cell.font
    alignment = cached_cell.alignment
    return CellModel(
        sheet=sheet,
        row=cached_cell.row,
        column=cached_cell.column,
        coordinate=cached_cell.coordinate,
        value=cached_cell.value,
        formula=formula,
        cached_value=cached_cell.value,
        data_type=cached_cell.data_type,
        number_format=cached_cell.number_format,
        style_id=cached_cell.style_id,
        fill=FillModel(
            fill_type=cached_cell.fill.fill_type,
            fgColor=_color_model(cached_cell.fill.fgColor),
            bgColor=_color_model(cached_cell.fill.bgColor),
        ),
        font=FontModel(
            name=font.name,
            size=float(font.sz) if font.sz is not None else None,
            bold=bool(font.bold),
            italic=bool(font.italic),
            underline=font.underline,
            strike=bool(font.strike),
            color=_color_model(font.color),
        ),
        alignment=AlignmentModel(
            horizontal=alignment.horizontal,
            vertical=alignment.vertical,
            wrap_text=alignment.wrap_text,
            shrink_to_fit=alignment.shrink_to_fit,
            indent=float(alignment.indent or 0.0),
            text_rotation=int(alignment.text_rotation or 0),
        ),
    )


def _blank_cell(sheet: str, row: int, column: int) -> CellModel:
    return CellModel(
        sheet=sheet,
        row=row,
        column=column,
        coordinate=f"{get_column_letter(column)}{row}",
        value=None,
        formula=None,
        cached_value=None,
        data_type="n",
        number_format="General",
        style_id=0,
        fill=FillModel(None, ColorModel(type="rgb", rgb="00000000"), ColorModel(type="rgb", rgb="00000000")),
        font=FontModel(None, None, False, False, None, False, ColorModel()),
        alignment=AlignmentModel(None, None, None, None, 0.0, 0),
    )


def _snapshot_sheet(cached_sheet, formula_sheet) -> SheetModel:
    max_row = max(cached_sheet.max_row, formula_sheet.max_row)
    max_column = max(cached_sheet.max_column, formula_sheet.max_column)
    rows = tuple(
        tuple(
            _cell_model(
                cached_sheet.title,
                cached_sheet.cell(row=row, column=column),
                formula_sheet.cell(row=row, column=column),
            )
            for column in range(1, max_column + 1)
        )
        for row in range(1, max_row + 1)
    )
    merged = tuple(
        MergedRangeModel(
            coordinate=str(cell_range),
            min_row=cell_range.min_row,
            max_row=cell_range.max_row,
            min_col=cell_range.min_col,
            max_col=cell_range.max_col,
        )
        for cell_range in cached_sheet.merged_cells.ranges
    )
    row_dimensions = tuple(
        RowDimensionModel(
            index=int(index),
            hidden=bool(dimension.hidden),
            outline_level=int(dimension.outlineLevel or 0),
            collapsed=bool(dimension.collapsed),
            height=float(dimension.height) if dimension.height is not None else None,
        )
        for index, dimension in sorted(cached_sheet.row_dimensions.items())
    )
    column_dimensions = tuple(
        ColumnDimensionModel(
            key=str(key),
            min_column=int(dimension.min or 0),
            max_column=int(dimension.max or 0),
            hidden=bool(dimension.hidden),
            outline_level=int(dimension.outlineLevel or 0),
            collapsed=bool(dimension.collapsed),
            width=float(dimension.width) if dimension.width is not None else None,
        )
        for key, dimension in sorted(cached_sheet.column_dimensions.items())
    )
    return SheetModel(
        title=cached_sheet.title,
        state=cached_sheet.sheet_state,
        max_row=max_row,
        max_column=max_column,
        rows=rows,
        merged_ranges=merged,
        row_dimensions=row_dimensions,
        column_dimensions=column_dimensions,
    )


def _stable_value(value: object):
    if value is None:
        return ["none", None]
    if isinstance(value, float) and math.isnan(value):
        return ["nan", None]
    if value is pd.NA:
        return ["pd.NA", None]
    if isinstance(value, (datetime, date, time)):
        return [type(value).__name__, value.isoformat()]
    if isinstance(value, Decimal):
        return ["Decimal", str(value)]
    if isinstance(value, bytes):
        return ["bytes", value.hex()]
    return [f"{type(value).__module__}.{type(value).__qualname__}", repr(value)]


def _digest_record(digest, record: object) -> None:
    digest.update(
        json.dumps(record, ensure_ascii=True, separators=(",", ":"), sort_keys=True).encode("utf-8")
    )
    digest.update(b"\n")


def _workbook_identity(
    source_sha: str,
    sheets: Iterable[SheetModel],
    tables: Iterable[TableModel],
    trim_storage_mappings: Iterable[object] = (),
    verify_storage_mappings: Iterable[object] = (),
    pattern_links: Iterable[object] = (),
) -> str:
    digest = hashlib.sha256()
    _digest_record(digest, ["source", source_sha])
    for sheet in sheets:
        _digest_record(
            digest,
            [
                "sheet",
                sheet.title,
                sheet.state,
                sheet.max_row,
                sheet.max_column,
                [str(item) for item in sheet.merged_ranges],
                [dimension.__dict__ for dimension in sheet.row_dimensions],
                [dimension.__dict__ for dimension in sheet.column_dimensions],
            ],
        )
        for row in sheet.rows:
            for cell in row:
                _digest_record(
                    digest,
                    [
                        cell.coordinate,
                        _stable_value(cell.value),
                        cell.formula,
                        cell.data_type,
                        cell.number_format,
                        cell.style_id,
                        asdict(cell.fill),
                        asdict(cell.font),
                        asdict(cell.alignment),
                    ],
                )
    for table in tables:
        _digest_record(digest, ["table", table.title, [_stable_value(value) for value in table.columns]])
        for row in table.rows:
            _digest_record(digest, [_stable_value(value) for value in row])
    for mapping in trim_storage_mappings:
        _digest_record(digest, ["trimStorage", asdict(mapping)])
    for mapping in verify_storage_mappings:
        _digest_record(digest, ["verifyStorage", asdict(mapping)])
    for record in pattern_links:
        _digest_record(digest, ["patternLink", record])
    return digest.hexdigest().upper()


__all__ = [
    "CellModel",
    "ColorModel",
    "FillModel",
    "MergedRangeModel",
    "SheetModel",
    "TableModel",
    "WorkbookModel",
    "WorkbookReader",
]
